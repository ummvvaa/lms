"""Состав подгрупп английского из файла школы (решение владельца, 02.10.2026).

Книга расписания говорит, из каких групп собрана подгруппа потока
(EEP-8-1 — из четырёх групп восьмой параллели), но не кто в ней учится.
Без состава урок такой подгруппы стоит у всех групп потока, а ученик
своего английского не видит. Файл школы — список «подгруппа — ученик —
группа»; по нему у каждой подгруппы появляется состав с даты.

Правила:

- подгруппа — по названию, только подгруппа внутри потока; кириллица,
  похожая на латиницу («ЕЕР-8-1», «ГЕ-10.1-2»), и пробелы не мешают;
- ученик — по ФИО среди учеников своей группы (колонка «Группа») или,
  если её нет, среди групп потока подгруппы; неуверенное совпадение —
  ошибка строки с кандидатами, молча не угадывается;
- у ученика одна подгруппа английского — EEP или GE: второй раз в файле —
  ошибка; подгруппы потоков из одного набора групп (EEP-8 и GE-8)
  взаимоисключающие, и при применении ученик уходит из прежней, даже если
  её нет в файле;
- любая ошибка строки останавливает применение: состав либо верный
  целиком, либо не пишется вовсе;
- ученики групп потока, которых нет в файле, и подгруппы, которых нет
  в файле, — предупреждения: их состав не трогается.

Раскладка файла: строка заголовков в первых десяти строках листа —
«Подгруппа», «ФИО» (или «Фамилия» и «Имя»), необязательная «Группа».
Лист без колонки «Подгруппа», названный как подгруппа («EEP-8-1»), — это
состав этой подгруппы. Формат для людей — `guides/SUBGROUP_MEMBERS.md`.
"""

from __future__ import annotations

import datetime as dt
import io
import re
from dataclasses import dataclass, field

from django.db import transaction
from django.utils.translation import gettext as _

from academics import cache
from academics.calendar import today
from academics.cohorts import group_ids_of, set_members
from academics.models import Cohort, CohortKind, CohortMembership
from students.models import Student, StudyGroup

#: где искать строку заголовков
HEADER_ROWS = 10

#: кириллица, которую пишут вместо латиницы в названиях подгрупп («ЕЕР», «ГЕ»)
LOOKALIKE = str.maketrans("АВГЕКМНОРСТХ", "ABGEKMHOPCTX")  # i18n-skip: буквы для сравнения названий


def subgroup_key(name: str) -> str:
    """«ЕЕР 8 – 1» → «EEP8-1»: регистр, пробелы, тире и кириллица не мешают."""
    text = re.sub(r"[‐‑‒–—―]", "-", str(name or "").upper().translate(LOOKALIKE))
    return re.sub(r"\s+", "", text)


@dataclass
class Line:
    """Строка файла: что написано и чем это оказалось."""

    sheet: str
    row: int
    subgroup: str
    name: str
    group: str
    cohort: Cohort | None = None
    student: Student | None = None
    error: str = ""

    @property
    def where(self) -> str:
        return _("«{sheet}», строка {row}").format(sheet=self.sheet, row=self.row)


@dataclass
class Plan:
    """Что сделает загрузка: строки, итог по подгруппам, ошибки и предупреждения."""

    lines: list[Line] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    #: подгруппа → ученики по файлу
    wanted: dict[int, list[int]] = field(default_factory=dict)
    cohorts: dict[int, Cohort] = field(default_factory=dict)

    @property
    def ok(self) -> bool:
        return not self.errors and bool(self.wanted)


# --- Чтение файла --------------------------------------------------------------


#: заголовок колонки → что в ней: слово внутри заголовка или его начало.
#: Порядок важен: «Подгруппа» проверяется раньше «Группы»
HEADER_WORDS: tuple[tuple[str, str, tuple[str, ...]], ...] = (
    ("subgroup", "in", ("подгрупп", "англ", "english")),  # i18n-skip: заголовки колонок файла школы
    ("name", "in", ("фио", "ф.и.о", "ученик", "студент")),  # i18n-skip: заголовки колонок файла школы
    ("last", "start", ("фамилия",)),  # i18n-skip: заголовок колонки файла школы
    ("first", "start", ("имя",)),  # i18n-skip: заголовок колонки файла школы
    ("group", "start", ("группа", "класс")),  # i18n-skip: заголовки колонок файла школы
)


def _header_kind(title: str) -> str:
    text = str(title or "").strip().lower()
    for kind, how, words in HEADER_WORDS:
        if any((word in text) if how == "in" else (text == word or text.startswith(word + " ")) for word in words):
            return kind
    return ""


def _cell(kinds: dict[str, int], row: list, kind: str) -> str:
    column = kinds.get(kind)
    value = row[column] if column is not None and column < len(row) else None
    return "" if value is None else str(value).strip()


def read(content: bytes, known: set[str] | None = None) -> tuple[list[Line], list[str]]:
    """Строки файла по листам. Ошибка — когда читать нечего.

    `known` — ключи подгрупп: лист без колонки «Подгруппа» берётся как состав
    подгруппы, только если называется как она.
    """
    from openpyxl import load_workbook

    try:
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:  # битый файл — не повод для 500
        return [], [_("Файл не читается: нужен Excel (.xlsx)")]
    lines: list[Line] = []
    problems: list[str] = []
    for sheet in book.worksheets:
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        found = None
        for index, row in enumerate(rows[:HEADER_ROWS]):
            kinds = {_header_kind(cell): column for column, cell in reversed(list(enumerate(row))) if cell}
            if "name" in kinds or "last" in kinds:
                found = index, kinds
                break
        if found is None:
            continue
        header, kinds = found
        if "subgroup" not in kinds and known is not None and subgroup_key(sheet.title) not in known:
            problems.append(
                _("Лист «{sheet}»: нет колонки «Подгруппа», а название листа — не подгруппа").format(sheet=sheet.title)
            )
            continue

        for offset, row in enumerate(rows[header + 1 :], start=header + 2):
            name = _cell(kinds, row, "name") or " ".join(
                x for x in (_cell(kinds, row, "last"), _cell(kinds, row, "first")) if x
            )
            subgroup = _cell(kinds, row, "subgroup") if "subgroup" in kinds else sheet.title
            if not name and not _cell(kinds, row, "subgroup"):
                continue
            lines.append(
                Line(sheet=sheet.title, row=offset, subgroup=subgroup, name=name, group=_cell(kinds, row, "group"))
            )
    if not lines:
        return [], problems or [_("В файле нет строк с колонками «Подгруппа» и «ФИО» — образец в описании формата")]
    return lines, problems


# --- Сопоставление -------------------------------------------------------------


def _stream_subgroups() -> dict[str, Cohort]:
    rows = Cohort.objects.filter(kind=CohortKind.SUBGROUP, stream__isnull=False).select_related("stream", "subject")
    return {subgroup_key(c.name): c for c in rows}


def _groups() -> dict[str, StudyGroup]:
    return {subgroup_key(g.code): g for g in StudyGroup.objects.all()}


def _group_of(text: str, groups: dict[str, StudyGroup]) -> StudyGroup | None:
    """Группа по ячейке: код целиком или словом в ячейке («8 KIOTO»)."""
    found = groups.get(subgroup_key(text))
    if found is not None:
        return found
    for word in re.split(r"[\s,;/]+", text):
        found = groups.get(subgroup_key(word))
        if found is not None:
            return found
    return None


def plan(content: bytes) -> Plan:
    """Прочитать файл и сопоставить: ничего не пишет."""
    from suggestions.name_matching import find

    out = Plan()
    subgroups = _stream_subgroups()
    lines, problems = read(content, set(subgroups))
    out.errors.extend(problems)
    out.lines = lines
    groups = _groups()
    pupils_of: dict[tuple, list[Student]] = {}
    seen: dict[int, Line] = {}
    for line in lines:
        cohort = subgroups.get(subgroup_key(line.subgroup))
        if cohort is None:
            line.error = _("{where}: подгруппы «{subgroup}» нет среди подгрупп потоков").format(
                where=line.where, subgroup=line.subgroup or _("пусто")
            )
            continue
        line.cohort = cohort
        if not line.name:
            line.error = _("{where}: не указан ученик").format(where=line.where)
            continue
        if line.group:
            group = _group_of(line.group, groups)
            if group is None:
                line.error = _("{where}: группы «{group}» нет").format(where=line.where, group=line.group)
                continue
            scope: tuple = (group.pk,)
        else:
            scope = tuple(group_ids_of(cohort))
        if scope not in pupils_of:
            pupils_of[scope] = list(Student.objects.filter(group_id__in=scope, is_active=True).select_related("group"))
        outcome = find(line.name, students=pupils_of[scope])
        if not (outcome.is_confident and outcome.best):
            if outcome.is_ambiguous:
                names = ", ".join(c.full_name for c in outcome.candidates[:3])
                line.error = _("{where}: «{name}» — похожих учеников несколько: {names}").format(
                    where=line.where, name=line.name, names=names
                )
            else:
                line.error = _("{where}: ученика «{name}» нет в группах подгруппы").format(
                    where=line.where, name=line.name
                )
            continue
        student = next(s for s in pupils_of[scope] if s.pk == outcome.best.student_id)
        line.student = student
        earlier = seen.get(student.pk)
        if earlier is not None:
            # у ученика одна подгруппа английского — EEP или GE, не обе
            line.error = _("{where}: {name} уже стоит в {subgroup} ({earlier}) — у ученика одна подгруппа").format(
                where=line.where, name=student.full_name, subgroup=earlier.cohort.name, earlier=earlier.where
            )
            continue
        seen[student.pk] = line
        out.cohorts[cohort.pk] = cohort
        out.wanted.setdefault(cohort.pk, []).append(student.pk)
    out.errors.extend(line.error for line in lines if line.error)
    _warn_about_the_rest(out, subgroups)
    return out


def alternatives(cohorts) -> list[Cohort]:
    """Подгруппы потоков, взаимоисключающие с этими: потоки из того же набора групп.

    EEP-8 и GE-8 — разные потоки и предметы, но из одних групп и в одно
    время: ученик ходит в одну подгруппу из всех их подгрупп.
    """
    sets = {frozenset(group_ids_of(c.stream)) for c in cohorts if c.stream_id}
    return [
        c
        for c in Cohort.objects.filter(kind=CohortKind.SUBGROUP, stream__isnull=False).select_related("stream")
        if frozenset(group_ids_of(c.stream)) in sets
    ]


def _warn_about_the_rest(out: Plan, subgroups: dict[str, Cohort]) -> None:
    """Кого и что файл не назвал: эти составы загрузка не трогает."""
    missing = sorted(c.name for c in alternatives(out.cohorts.values()) if c.pk not in out.wanted)
    if missing:
        out.warnings.append(_("Подгрупп нет в файле, их состав не меняется: {names}").format(names=", ".join(missing)))
    listed = {sid for ids in out.wanted.values() for sid in ids}
    group_ids: set[int] = set()
    for cohort in out.cohorts.values():
        group_ids.update(group_ids_of(cohort))
    for group in StudyGroup.objects.filter(pk__in=group_ids).order_by("code"):
        left = [
            s.full_name
            for s in Student.objects.filter(group=group, is_active=True).order_by("last_name", "first_name")
            if s.pk not in listed
        ]
        if left:
            out.warnings.append(
                _("{group}: нет в файле, английского в расписании у них не будет — {names}").format(
                    group=group.code, names=", ".join(left)
                )
            )


# --- Отчёт и применение --------------------------------------------------------


def year_start() -> dt.date:
    """Начало учебного года — дата состава по умолчанию: журналы прошлых уроков получают своих учеников."""
    from academics import calendar as school_calendar

    try:
        return school_calendar.load().year.starts
    except Exception:  # учебного года ещё нет — с сегодняшнего дня
        return today()


def report(out: Plan, since: dt.date) -> dict:
    """Что покажет экран: по подгруппам — сколько учеников, кто придёт и кто уйдёт."""
    rows = []
    for pk, ids in out.wanted.items():
        cohort = out.cohorts[pk]
        current = set(
            CohortMembership.objects.filter(cohort=cohort, until__isnull=True).values_list("student_id", flat=True)
        )
        rows.append(
            {
                "id": pk,
                "name": cohort.name,
                "stream": cohort.stream.name if cohort.stream_id else "",
                "subject": cohort.subject.title if cohort.subject_id else "",
                "students": len(ids),
                "joined": len(set(ids) - current),
                "left": len(current - set(ids)),
            }
        )
    rows.sort(key=lambda row: row["name"])
    return {
        "since": since,
        "lines": len(out.lines),
        "matched": sum(len(ids) for ids in out.wanted.values()),
        "subgroups": rows,
        "errors": out.errors[:200],
        "errors_total": len(out.errors),
        "warnings": out.warnings,
        "ok": out.ok,
    }


@transaction.atomic
def apply(out: Plan, since: dt.date, *, actor=None) -> None:
    """Записать составы с даты. Ученик файла уходит из прежней подгруппы, которой нет в файле."""
    from academics import schedule
    from core import stored_text
    from core.domains import Source

    if not out.ok:
        raise ValueError("план с ошибками не применяется")  # i18n-skip: страж кода, экран сюда не доходит
    listed = {sid for ids in out.wanted.values() for sid in ids}
    # прежняя подгруппа из тех же групп (EEP ↔ GE), которой нет в файле: ученик переехал — закрываем
    others = [c.pk for c in alternatives(out.cohorts.values()) if c.pk not in out.wanted]
    stale = CohortMembership.objects.filter(student_id__in=listed, until__isnull=True, cohort_id__in=others)
    stale.update(until=since)
    for pk, ids in out.wanted.items():
        set_members(out.cohorts[pk], ids, since)
        schedule.log_change(
            stored_text.store(
                stored_text.COHORT_CHANGED, name=out.cohorts[pk].name, count=len(ids), date=f"{since:%d.%m.%Y}"
            ),
            actor=actor,
            source=Source.IMPORT,
        )
    cache.invalidate()
