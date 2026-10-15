"""Импорт расписания и сотрудников из книги школы (решение владельца, 29.09.2026).

Книга — шесть листов с постоянными заголовками, читаются по названию:
«Сотрудники», «Группы», «Предметы», «Звонки», «Подгруппы», «Уроки».
«Инструкция» и «Допущения» — для людей и пропускаются.

Предпросмотр и применение — один и тот же код: предпросмотр пишет всё
в транзакцию и откатывает её, поэтому цифры предпросмотра совпадают
с тем, что сделает «Применить». Разница одна: при предпросмотре не
выпускаются пароли и не заводятся строки уроков — это дорого, а счёт
уроков считается и без записи.

Повторный импорт того же файла ничего не удваивает: человек находится
по логину, ФИО или телефону, группа — по коду, предмет — по названию,
подгруппа и поток — по названию, журнал — по предмету, учителю и составу,
еженедельный урок — по журналу, дню и номеру.

Замена расписания с даты (решение владельца, 15.10.2026): файл заменяет
прежнее расписание у пар «группа + предмет», которые в нём есть. Еженедельный
урок такой пары, которого нет в файле, заканчивается накануне даты замены;
его уроки с этой даты без данных удаляются, с данными остаются и называются
в отчёте. Пары, которых в файле нет совсем, не трогаются. Учитель пары
сменился — журнал передаётся с даты замены, как «Сменить учителя»: журнал
общий, итоги четверти и раздел отчёта остаются при нём.

Что правило, а что предупреждение:

- ошибка останавливает весь импорт: нет листа или колонки, ссылка
  на несуществующего сотрудника, группу или подгруппу, два человека
  в базе с одним ФИО;
- накладки кабинетов и учителей, время мимо звонков, роль, которую
  импорт не меняет, — предупреждения: файл школы бывает таким на деле.

Пароли открытым текстом возвращаются ровно один раз — в ответе
«Применить»; в журнал и в логи они не попадают.
"""

from __future__ import annotations

import datetime as dt
import hashlib
import io
import re
from collections import defaultdict
from dataclasses import dataclass, field

from django.db import transaction
from django.db.models import Exists, OuterRef, Q
from django.utils.translation import gettext as _
from django.utils.translation import gettext_lazy, gettext_noop

from academics import cache
from academics import calendar as school_calendar
from academics.calendar import today
from academics.cohorts import group_cohort, make_stream
from academics.models import (
    AcademicYear,
    Attendance,
    Bell,
    BellSchedule,
    Cohort,
    CohortKind,
    Course,
    Grade,
    Lesson,
    LessonKind,
    LessonSeries,
    LessonStatus,
    ReportRole,
    Scheme,
    StreamPart,
    Subject,
    TeacherProfile,
)
from academics.payloads import user_name
from academics.schedule import SCHEDULE_EVENT, materialize, reassign, series_dates
from accounts.models import Role, User
from core.audit import record_change
from core.domains import Source
from core.models import AuditLog
from students.models import GroupLanguage, Parallel, StudyGroup
from students.phones import normalize_kz

#: Еженедельный урок действует с начала года; строки уроков — со дня импорта
SERIES_START = dt.date(2026, 9, 1)

#: Листы книги школы. Названия листов, колонок и значений в ячейках — как
#: в файле школы, по-русски: по ним книга читается, и они не переводятся
#: (шаблон школы один, на русском). Переводятся ошибки и предупреждения
STAFF, GROUPS, SUBJECTS, BELLS, SUBGROUPS, LESSONS = (  # i18n-skip: названия листов книги школы
    "Сотрудники",
    "Группы",
    "Предметы",
    "Звонки",
    "Подгруппы",
    "Уроки",
)

#: Лист → колонки, без которых импорт не идёт. Лишние колонки не мешают
SHEETS: dict[str, tuple[str, ...]] = {  # i18n-skip: заголовки колонок книги школы
    STAFF: ("ID", "ФИО", "Фамилия", "Имя", "Отчество", "Роль", "Предметы", "Телефон", "Почта", "Логин"),
    GROUPS: ("Группа", "Параллель", "Литера", "Язык обучения", "ID куратора", "Домашний кабинет"),
    SUBJECTS: ("Предмет",),
    BELLS: ("Параллели", "Урок", "Начало", "Конец"),
    SUBGROUPS: ("Подгруппа", "Параллель", "Предмет", "Из групп", "ID учителя", "Основной кабинет"),
    LESSONS: (
        "Параллель",
        "Группа",
        "Подгруппа",
        "День",
        "Урок",
        "Начало",
        "Конец",
        "Предмет",
        "ID учителя",
        "Кабинет",
        "Совместный урок",
    ),
}

#: Колонки, которых в старых книгах нет: пустое значение — умолчание
OPTIONAL: dict[str, tuple[str, ...]] = {  # i18n-skip: заголовки колонок книги школы
    SUBJECTS: ("Оценивание",),
    LESSONS: ("Учитель не менять",),
}

#: «Учитель не менять» = «да»: учитель — тот, что сейчас у журнала в LMS
KEEP_TEACHER = "да"  # i18n-skip: значение ячейки книги школы

#: «Оценивание» на листе «Предметы» → схема предмета. «Без оценок» — существующая
#: схема «Только ФО», в табель не идёт (решение владельца, 29.09.2026); пусто — ФО, СОР, СОЧ
SCHEMES: dict[str, str] = {"фо, сор, соч": Scheme.KZ, "без оценок": Scheme.FO}  # i18n-skip: значения ячейки книги

#: Роль в файле → роль учётки и нужен ли профиль учителя. `None` — администрация:
#: роли под неё нет, учётка заводится выключенной, роль назначает администратор
ROLES: dict[str, tuple[str | None, bool]] = {  # i18n-skip: значения ячейки книги школы
    "учитель": (Role.TEACHER, True),
    "куратор": (Role.CURATOR, False),
    "учитель + куратор": (Role.CURATOR, True),
    "администрация": (None, False),
}

#: Роли, которые импорт меняет друг на друга. Остальные только называет в отчёте
SWITCHABLE = (Role.TEACHER, Role.CURATOR)

LANGUAGES = {"казахский": GroupLanguage.KK, "русский": GroupLanguage.RU}  # i18n-skip: значения ячейки книги школы

WEEKDAYS = {"пн": 1, "вт": 2, "ср": 3, "чт": 4, "пт": 5, "сб": 6}  # i18n-skip: день недели в ячейке книги
#: День недели в предупреждениях — на языке того, кто загружает
WEEKDAY_WORDS = {
    1: gettext_lazy("Пн"),
    2: gettext_lazy("Вт"),
    3: gettext_lazy("Ср"),
    4: gettext_lazy("Чт"),
    5: gettext_lazy("Пт"),
    6: gettext_lazy("Сб"),
}

#: Предел размера файла: книга школы весит около 100 КБ
MAX_BYTES = 5 * 1024 * 1024


class ImportRefused(ValueError):
    """Импорт невозможен — текст объясняет почему."""


class _Rollback(Exception):
    """Откатить транзакцию предпросмотра или импорта с ошибками."""


# --- Чтение ячеек ----------------------------------------------------------------


def _text(value) -> str:
    """Ячейка строкой: число без «.0», пробелы схлопнуты."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value).split())


def _int(value) -> int | None:
    text = _text(value)
    return int(text) if text.isdigit() else None


def _time(value) -> dt.time | None:
    if isinstance(value, dt.datetime):
        return value.time().replace(second=0, microsecond=0)
    if isinstance(value, dt.time):
        return value.replace(second=0, microsecond=0)
    match = re.fullmatch(r"(\d{1,2})[:.](\d{2})", _text(value))
    if not match:
        return None
    hours, minutes = int(match.group(1)), int(match.group(2))
    return dt.time(hours, minutes) if hours < 24 and minutes < 60 else None


def _where(sheet: str, row: int) -> str:
    """«Лист», строка N — где в книге ошибка; лист назван, как в файле."""
    return _("«{sheet}», строка {row}").format(sheet=sheet, row=row)


def _key(text: str) -> str:
    """Ключ сравнения: без регистра и лишних пробелов."""
    return " ".join((text or "").split()).casefold()


def _list(text: str) -> list[str]:
    return [part.strip() for part in re.split(r"[,;]", text or "") if part.strip() and part.strip() != "—"]


def _parallels(text: str) -> list[int]:
    """«8–9» → [8, 9]; «10» → [10]; «8, 9» → [8, 9]."""
    found = [int(n) for n in re.findall(r"\d+", text or "")]
    if len(found) == 2 and re.search(r"[–—-]", text) and found[0] < found[1]:
        return list(range(found[0], found[1] + 1))
    return found


# --- Разбор книги -----------------------------------------------------------------


@dataclass
class Staff:
    row: int
    code: str
    full_name: str
    last_name: str
    first_name: str
    role_text: str
    subjects: list[str]
    phone: str
    email: str
    login: str


@dataclass
class GroupRow:
    row: int
    code: str
    parallel: int | None
    letter: str
    language: str
    curator: str
    home_room: str


@dataclass
class SubjectRow:
    row: int
    title: str
    #: «Оценивание» как в файле, без регистра; пусто — колонки нет или ячейка пустая
    grading: str


@dataclass
class BellRow:
    row: int
    label: str
    parallels: list[int]
    number: int | None
    starts: dt.time | None
    ends: dt.time | None


@dataclass
class SubgroupRow:
    row: int
    name: str
    parallel: int | None
    subject: str
    groups: list[str]
    teacher: str
    room: str


@dataclass
class LessonRow:
    row: int
    parallel: int | None
    group: str
    subgroup: str
    weekday: int | None
    slot: int | None
    starts: dt.time | None
    ends: dt.time | None
    subject: str
    teacher: str
    room: str
    joint: str
    #: «Учитель не менять» как в файле, без регистра; пусто — колонки нет
    keep_teacher: str = ""


@dataclass
class Book:
    staff: list[Staff] = field(default_factory=list)
    groups: list[GroupRow] = field(default_factory=list)
    subjects: list[SubjectRow] = field(default_factory=list)
    bells: list[BellRow] = field(default_factory=list)
    subgroups: list[SubgroupRow] = field(default_factory=list)
    lessons: list[LessonRow] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)


def fingerprint(data: bytes, replace_from: dt.date | None = None) -> str:
    """Отпечаток файла и режима: «Применить» принимает только то, что видели в предпросмотре.

    Дата замены входит в отпечаток: тот же файл с другой датой — другой импорт.
    """
    digest = hashlib.sha256(data)
    if replace_from is not None:
        digest.update(f"|replace:{replace_from.isoformat()}".encode())
    return digest.hexdigest()


def _rows(sheet, title: str, errors: list[str]):
    """Строки листа словарями «колонка → значение» с номером строки файла."""
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None) or ()
    names = [_text(cell) for cell in header]
    missing = [column for column in SHEETS[title] if column not in names]
    if missing:
        columns = ", ".join(f"«{c}»" for c in missing)
        errors.append(_("На листе «{sheet}» нет колонок: {columns}").format(sheet=title, columns=columns))
        return
    index = {name: names.index(name) for name in (*SHEETS[title], *OPTIONAL.get(title, ())) if name in names}
    for number, values in enumerate(rows, start=2):
        if not any(_text(v) for v in values):
            continue
        cells = {name: (values[i] if i < len(values) else None) for name, i in index.items()}
        yield number, {name: None for name in OPTIONAL.get(title, ())} | cells


def parse(data: bytes) -> Book:
    """Прочитать книгу. Ошибки формата — в `Book.errors`, до базы дело не доходит."""
    import zipfile

    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    book = Book()
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (InvalidFileException, KeyError, OSError, ValueError, zipfile.BadZipFile):
        book.errors.append(_("Файл не читается как книга Excel (.xlsx)"))
        return book
    missing = [title for title in SHEETS if title not in workbook.sheetnames]
    if missing:
        sheets = ", ".join(f"«{t}»" for t in missing)
        book.errors.append(_("В файле нет листов: {sheets}").format(sheets=sheets))
        return book

    for number, cells in _rows(workbook[STAFF], STAFF, book.errors):  # i18n-skip: колонки книги школы
        book.staff.append(
            Staff(
                row=number,
                code=_text(cells["ID"]),
                full_name=_text(cells["ФИО"]),
                last_name=_text(cells["Фамилия"]),
                first_name=_text(cells["Имя"]),
                role_text=_key(cells["Роль"] and str(cells["Роль"]).replace("+", " + ")),
                subjects=_list(_text(cells["Предметы"])),
                phone=_text(cells["Телефон"]),
                email=_text(cells["Почта"]).lower(),
                login=_text(cells["Логин"]).lower(),
            )
        )
    for number, cells in _rows(workbook[GROUPS], GROUPS, book.errors):  # i18n-skip: колонки книги школы
        book.groups.append(
            GroupRow(
                row=number,
                code=_text(cells["Группа"]).upper(),
                parallel=_int(cells["Параллель"]),
                letter=_text(cells["Литера"]),
                language=_key(_text(cells["Язык обучения"])),
                curator=_text(cells["ID куратора"]),
                home_room=_text(cells["Домашний кабинет"]),
            )
        )
    for number, cells in _rows(workbook[SUBJECTS], SUBJECTS, book.errors):  # i18n-skip: колонки книги школы
        book.subjects.append(
            SubjectRow(
                row=number,
                title=_text(cells["Предмет"]),
                grading=re.sub(r"\s*,\s*", ", ", _key(_text(cells["Оценивание"]))),
            )
        )
    for number, cells in _rows(workbook[BELLS], BELLS, book.errors):  # i18n-skip: колонки книги школы
        label = _text(cells["Параллели"])
        book.bells.append(
            BellRow(
                row=number,
                label=label,
                parallels=_parallels(label),
                number=_int(cells["Урок"]),
                starts=_time(cells["Начало"]),
                ends=_time(cells["Конец"]),
            )
        )
    for number, cells in _rows(workbook[SUBGROUPS], SUBGROUPS, book.errors):  # i18n-skip: колонки книги школы
        book.subgroups.append(
            SubgroupRow(
                row=number,
                name=_text(cells["Подгруппа"]),
                parallel=_int(cells["Параллель"]),
                subject=_text(cells["Предмет"]),
                groups=[code.upper() for code in _list(_text(cells["Из групп"]))],
                teacher=_text(cells["ID учителя"]),
                room=_text(cells["Основной кабинет"]),
            )
        )
    for number, cells in _rows(workbook[LESSONS], LESSONS, book.errors):  # i18n-skip: колонки книги школы
        book.lessons.append(
            LessonRow(
                row=number,
                parallel=_int(cells["Параллель"]),
                group=_text(cells["Группа"]).upper(),
                subgroup=_text(cells["Подгруппа"]),
                weekday=WEEKDAYS.get(_key(_text(cells["День"]))[:2]),
                slot=_int(cells["Урок"]),
                starts=_time(cells["Начало"]),
                ends=_time(cells["Конец"]),
                subject=_text(cells["Предмет"]),
                teacher=_text(cells["ID учителя"]),
                room=_text(cells["Кабинет"]),
                joint=_text(cells["Совместный урок"]),
                keep_teacher=_key(_text(cells["Учитель не менять"])),
            )
        )
    workbook.close()
    return book


# --- Отчёт ------------------------------------------------------------------------


@dataclass
class Section:
    code: str
    title: str
    created: int = 0
    updated: int = 0
    unchanged: int = 0

    def as_dict(self) -> dict:
        return {
            "code": self.code,
            "title": self.title,
            "created": self.created,
            "updated": self.updated,
            "unchanged": self.unchanged,
        }


@dataclass
class Report:
    fingerprint: str
    applied: bool = False
    errors: list[str] = field(default_factory=list)
    warnings: list[dict] = field(default_factory=list)
    sections: dict[str, Section] = field(default_factory=dict)
    curator_changes: list[dict] = field(default_factory=list)
    needs_role: list[dict] = field(default_factory=list)
    roles_kept: list[dict] = field(default_factory=list)
    stale_series: list[str] = field(default_factory=list)
    groups_outside: list[str] = field(default_factory=list)
    skipped_bells: int = 0
    lessons_from: dt.date | None = None
    lessons_to_create: int = 0
    series_without_teacher: int = 0
    #: уроки в неделю без учителя: состав и предмет
    without_teacher: list[dict] = field(default_factory=list)
    credentials: list[dict] = field(default_factory=list)
    # --- замена расписания с даты ---
    replace_from: dt.date | None = None
    series_ended: int = 0
    lessons_deleted: int = 0
    lessons_kept: list[dict] = field(default_factory=list)
    untouched: list[dict] = field(default_factory=list)
    journals_passed: list[dict] = field(default_factory=list)
    subgroups_unused: list[str] = field(default_factory=list)

    def section(self, code: str, title: str) -> Section:
        return self.sections.setdefault(code, Section(code, title))

    def warn(self, kind: str, text: str) -> None:
        self.warnings.append({"kind": kind, "text": text})

    def as_dict(self) -> dict:
        return {
            "fingerprint": self.fingerprint,
            "applied": self.applied,
            "errors": self.errors,
            "warnings": self.warnings,
            "sections": [s.as_dict() for s in self.sections.values()],
            "curator_changes": self.curator_changes,
            "needs_role": self.needs_role,
            "roles_kept": self.roles_kept,
            "stale_series": self.stale_series,
            "groups_outside": self.groups_outside,
            "skipped_bells": self.skipped_bells,
            "lessons_from": self.lessons_from,
            "lessons_to_create": self.lessons_to_create,
            "series_without_teacher": self.series_without_teacher,
            "without_teacher": self.without_teacher,
            "credentials": self.credentials,
            "replace_from": self.replace_from,
            "series_ended": self.series_ended,
            "lessons_deleted": self.lessons_deleted,
            "lessons_kept": self.lessons_kept,
            "untouched": self.untouched,
            "journals_passed": self.journals_passed,
            "subgroups_unused": self.subgroups_unused,
            "detail": self.detail(),
        }

    def detail(self) -> str:
        if self.errors:
            return _("Импорт не выполнен: ошибок {count}").format(count=len(self.errors))
        made = sum(s.created for s in self.sections.values())
        changed = sum(s.updated for s in self.sections.values())
        numbers = {"created": made, "updated": changed, "warnings": len(self.warnings)}
        if self.applied:
            return _("Импорт выполнен: создаётся {created}, обновляется {updated}, предупреждений {warnings}").format(
                **numbers
            )
        return _("Предпросмотр: создаётся {created}, обновляется {updated}, предупреждений {warnings}").format(
            **numbers
        )


# --- Запись -----------------------------------------------------------------------


def preview(data: bytes, *, replace_from: dt.date | None = None) -> dict:
    """Что сделает импорт — без единой записи в базе."""
    return _run(data, actor=None, commit=False, replace_from=replace_from)


def apply(data: bytes, *, actor, expected: str, replace_from: dt.date | None = None) -> dict:
    """Применить файл, который видели в предпросмотре. Всё или ничего."""
    if expected != fingerprint(data, replace_from):
        raise ImportRefused(_("Файл или дата замены не те, что в предпросмотре: загрузите файл заново и проверьте"))
    return _run(data, actor=actor, commit=True, replace_from=replace_from)


def _run(data: bytes, *, actor, commit: bool, replace_from: dt.date | None = None) -> dict:
    if len(data) > MAX_BYTES:
        raise ImportRefused(_("Файл больше 5 МБ — это не книга расписания"))
    if replace_from is not None and replace_from < today():
        # прошедшие уроки не меняются: иначе за прошедшие дни заведутся неотмеченные уроки
        raise ImportRefused(_("Дата замены — не раньше сегодняшнего дня"))
    report = Report(fingerprint=fingerprint(data, replace_from), replace_from=replace_from)
    book = parse(data)
    report.errors.extend(book.errors)
    if report.errors:
        return report.as_dict()
    try:
        with transaction.atomic():
            _Writer(book, report, actor=actor, commit=commit, replace_from=replace_from).run()
            if report.errors or not commit:
                raise _Rollback
    except _Rollback:
        report.credentials = []
    else:
        report.applied = True
    cache.invalidate()
    return report.as_dict()


class _Writer:
    """Один проход по книге: листы по порядку, всё в одной транзакции."""

    def __init__(self, book: Book, report: Report, *, actor, commit: bool, replace_from: dt.date | None = None) -> None:
        self.book = book
        self.report = report
        self.actor = actor if getattr(actor, "pk", None) else None
        self.commit = commit
        self.day = today()
        self.replace_from = replace_from
        #: с какого дня импорт меняет уроки: дата замены или сегодня
        self.cut = replace_from or self.day
        self.year: AcademicYear | None = None
        self.subjects: dict[str, Subject] = {}
        self.staff: dict[str, User] = {}
        self.groups: dict[str, StudyGroup] = {}
        self.subgroups: dict[str, Cohort] = {}
        self.grids: dict[int, dict[int, tuple[dt.time, dt.time]]] = {}
        self.touched_series: set[int] = set()
        self.missing_bells: dict[tuple, list[str]] = defaultdict(list)
        self.new_series: list[LessonSeries] = []
        self.units: list[dict] = []
        self.cohort_groups: dict[int, set[int]] = {}
        self.current_teachers: dict[tuple[int, int, int, int], User | None] = {}
        self.keep_warned: set[tuple[int, int]] = set()

    def error(self, text: str) -> None:
        self.report.errors.append(text)

    def run(self) -> None:
        self.year = school_calendar.current_year()
        if self.year is None:
            self.error(_("Нет текущего учебного года: заведите его на странице «Учебный год»"))
            return
        if self.replace_from is not None and self.replace_from > self.year.ends:
            self.error(_("Дата замены позже конца учебного года"))
            return
        self._subjects()
        self._staff()
        self._groups()
        self._bells()
        self._subgroups()
        if self.report.errors:
            return
        self._lessons()
        self._missing_bells()
        if self.replace_from is not None:
            self._replace()
        else:
            self._stale()
        self._conflicts()
        if self.report.errors:
            return
        self._materialize()
        if self.commit:
            self._log()

    # --- Предметы ---

    def _subjects(self) -> None:
        section = self.report.section("subjects", _("Предметы"))
        known = {_key(s.title): s for s in Subject.objects.all()}
        for row in self.book.subjects:
            title = row.title
            if not title:
                continue
            if row.grading and row.grading not in SCHEMES:
                self.error(
                    _("{where}: оценивание «{value}» — нужно «ФО, СОР, СОЧ» или «без оценок»").format(
                        where=_where(SUBJECTS, row.row), value=row.grading
                    )
                )
                continue
            scheme = SCHEMES.get(row.grading, Scheme.KZ)
            if not self._sat_mode(title, _where(SUBJECTS, row.row)):
                continue
            found = known.get(_key(title))
            if found is None and self.report_role_for(title):
                # «SAT Math» и «SAT Verbal» в книге при одном предмете — это SAT
                from academics import sat_subjects

                found = sat_subjects._find(sat_subjects.MERGED_NAMES)
                if found is not None:
                    known[_key(title)] = found
            if found is None:
                from academics.keep_subjects import is_kept_title

                # в LMS ведутся только предметы, которых нет в Kundelik; новый предмет
                # не из них — «только расписание», включают его в «Учебном году»
                kept = is_kept_title(title)
                found = Subject.objects.create(
                    code=self._subject_code(title),
                    title=title[:100],
                    short_title=title[:32],
                    # шесть предметов школы — только ФО из 10 (решение владельца, 07.10.2026)
                    scheme=Scheme.FO if kept else scheme,
                    order=100,
                    in_lms=kept,
                )
                known[_key(title)] = found
                section.created += 1
                if not row.grading and not kept:
                    self.report.warn(
                        "subject",
                        _("Новый предмет «{subject}»: оценивание в файле не указано — ФО, СОР и СОЧ, проверьте").format(
                            subject=title
                        ),
                    )
            else:
                section.unchanged += 1
                # схему заведённого предмета импорт не меняет — только называет расхождение
                if row.grading and found.scheme != scheme:
                    self.report.warn(
                        "scheme",
                        _("Предмет «{subject}»: в LMS — «{current}», в файле — «{wanted}»; схема не меняется").format(
                            subject=found.title, current=found.get_scheme_display(), wanted=Scheme(scheme).label
                        ),
                    )
        self.subjects = known

    def _subject_code(self, title: str) -> str:
        from accounts.logins import latin

        base = re.sub(r"[^a-z0-9]+", "-", latin(title).lower()).strip("-")[:28] or "subject"
        code, number = base, 1
        while Subject.objects.filter(code=code).exists():
            number += 1
            code = f"{base}-{number}"
        return code

    def subject(self, title: str, where: str) -> Subject | None:
        found = self.subjects.get(_key(title))
        if found is None:
            self.error(_("{where}: предмета «{subject}» нет ни в файле, ни в LMS").format(where=where, subject=title))
        return found

    def _sat_mode(self, title: str, where: str) -> bool:
        """SAT в книге — по режиму школы: при двух предметах «SAT» в книге — ошибка.

        «Укажите SAT Math или SAT Verbal» (решение владельца, 11.10.2026); при одном
        предмете строки «SAT Math» и «SAT Verbal» в книге ведут на SAT, журнал получает
        раздел отчёта по названию.
        """
        from academics import sat_subjects

        key = _key(title)
        if key in sat_subjects.MERGED_NAMES and sat_subjects.two_subjects():
            self.error(_("{where}: SAT ведётся двумя предметами — укажите SAT Math или SAT Verbal").format(where=where))
            return False
        return True

    def report_role_for(self, title: str) -> str:
        """Раздел отчёта у журнала SAT из названия в книге при одном предмете."""
        from academics import sat_subjects

        key = _key(title)
        if sat_subjects.two_subjects():
            return ""
        if key in sat_subjects.MATH_NAMES:
            return ReportRole.SAT_MATH
        if key in sat_subjects.VERBAL_NAMES:
            return ReportRole.SAT_VERBAL
        return ""

    # --- Сотрудники ---

    def _staff(self) -> None:
        section = self.report.section("staff", _("Сотрудники"))
        people = list(User.objects.exclude(role=Role.STUDENT))
        by_email = {p.email.lower(): p for p in people if p.email}
        by_login = {p.login.lower(): p for p in people if p.login}
        by_name: dict[str, list[User]] = defaultdict(list)
        by_phone: dict[str, list[User]] = defaultdict(list)
        for person in people:
            if person.full_name:
                by_name[_key(person.full_name)].append(person)
            if person.phone:
                by_phone[person.phone].append(person)
        taken_logins: set[str] = set()
        codes: set[str] = set()
        for row in self.book.staff:
            where = _where(STAFF, row.row)
            if not row.code or not row.full_name:
                self.error(_("{where}: нет ID или ФИО").format(where=where))
                continue
            if row.code in codes:
                self.error(_("{where}: ID {code} встречается дважды").format(where=where, code=row.code))
                continue
            codes.add(row.code)
            if row.role_text not in ROLES:
                self.error(
                    _("{where}: роль «{role}» не из списка: {roles}").format(
                        where=where, role=row.role_text, roles=", ".join(ROLES)
                    )
                )
                continue
            phone = normalize_kz(row.phone)
            if phone and not re.fullmatch(r"\+7\d{10}", phone):
                self.report.warn(
                    "phone",
                    _("{name}: телефон «{phone}» не похож на номер — не записан").format(
                        name=row.full_name, phone=row.phone
                    ),
                )
                phone = ""
            user, why = self._match(row, phone, by_email, by_login, by_name, by_phone)
            if why:
                self.error(f"{where}: {why}")
                continue
            if user is None:
                user = self._create(row, phone, taken_logins)
                if user is None:
                    continue
                section.created += 1
            elif self._update(user, row, phone):
                section.updated += 1
            else:
                section.unchanged += 1
            self.staff[row.code] = user

    def _match(self, row: Staff, phone: str, by_email, by_login, by_name, by_phone) -> tuple[User | None, str]:
        """Уже заведённый человек: по почте, по логину, по ФИО, по телефону. Не угадываем.

        Почта — первой (решение владельца, 29.09.2026): у сотрудников на проде
        вход по почте, логина нет, а ФИО короткое («Нурманова Айдана») — по ФИО
        из файла они не находились, и импорт заводил вторые учётки.
        """
        if row.email:
            if row.email in by_email:
                return by_email[row.email], ""
            if User.objects.filter(email__iexact=row.email).exists():
                return None, _("почта {email} — у ученика, не у сотрудника").format(email=row.email)
        if row.login and row.login in by_login:
            return by_login[row.login], ""
        named = by_name.get(_key(row.full_name), [])
        if len(named) > 1:
            return None, _("в LMS несколько сотрудников с ФИО «{name}» — оставьте одного").format(name=row.full_name)
        if named:
            return named[0], ""
        phoned = by_phone.get(phone, []) if phone else []
        if len(phoned) > 1:
            return None, _("в LMS несколько сотрудников с телефоном {phone}").format(phone=phone)
        return (phoned[0] if phoned else None), ""

    def _create(self, row: Staff, phone: str, taken: set[str]) -> User | None:
        from accounts.logins import LOGIN_RE, make_login

        role, _profile = ROLES[row.role_text]
        email = row.email or None
        if email and User.objects.filter(email__iexact=email).exists():
            self.error(
                _("{where}: почта {email} уже занята другой учётной записью").format(
                    where=_where(STAFF, row.row), email=email
                )
            )
            return None
        login = None
        if not email:
            login = row.login if row.login and LOGIN_RE.match(row.login) else ""
            if not login or login in taken or User.objects.filter(login__iexact=login).exists():
                wanted = login or _("без логина")
                login = make_login(row.first_name, row.last_name, taken=taken)
                self.report.warn(
                    "login",
                    _("{name}: логин «{wanted}» занят или неверен — выдан «{login}»").format(
                        name=row.full_name, wanted=wanted, login=login
                    ),
                )
            taken.add(login)
        admin_staff = role is None
        user = User.objects.create_user(
            email=email,
            login=login,
            password=None,
            full_name=row.full_name[:200],
            phone=phone,
            # администрации роли нет: учётка выключена, роль и доступ даёт администратор
            role=role or Role.TEACHER,
            is_active=not admin_staff,
        )
        record_change(
            instance=user, field_name="role", old_value="", new_value=user.role, actor=self.actor, source=Source.IMPORT
        )
        self._profile(user, row)
        if admin_staff:
            self.report.needs_role.append(
                {"full_name": user.full_name, "login": user.handle, "position": ", ".join(row.subjects)}
            )
        elif self.commit:
            from accounts import temporary

            password = temporary.issue(user)
            self.report.credentials.append(
                {
                    "full_name": user.full_name,
                    "email": user.email,
                    "login": user.handle,
                    "password": password,
                    "expires_at": user.temp_password_expires_at,
                    "group": "",
                }
            )
        return user

    def _update(self, user: User, row: Staff, phone: str) -> bool:
        """ФИО, телефон, предметы и роль учитель ↔ куратор. Логин, почту и вход не трогаем."""
        changed = False
        wanted_role, _profile = ROLES[row.role_text]
        fields: dict[str, str] = {}
        if row.full_name and user.full_name != row.full_name:
            fields["full_name"] = row.full_name[:200]
        if phone and user.phone != phone:
            fields["phone"] = phone
        if wanted_role is not None and wanted_role != user.role:
            if user.role in SWITCHABLE:
                fields["role"] = wanted_role
            else:
                self.report.roles_kept.append(
                    {"full_name": user_name(user), "role": user.get_role_display(), "file": row.role_text}
                )
        elif wanted_role is None and user.role in SWITCHABLE:
            self.report.roles_kept.append(
                {"full_name": user_name(user), "role": user.get_role_display(), "file": row.role_text}
            )
        for name, value in fields.items():
            record_change(
                instance=user,
                field_name=name,
                old_value=getattr(user, name),
                new_value=value,
                actor=self.actor,
                source=Source.IMPORT,
            )
            setattr(user, name, value)
        if fields:
            user.save(update_fields=list(fields))
            changed = True
        if self._profile(user, row):
            changed = True
        return changed

    def _profile(self, user: User, row: Staff) -> bool:
        """Профиль учителя с предметами из файла. Возвращает, поменялось ли что-то."""
        _role, needs_profile = ROLES[row.role_text]
        if not needs_profile:
            return False
        wanted = []
        for title in row.subjects:
            found = self.subjects.get(_key(title))
            if found is None:
                self.report.warn(
                    "subject",
                    _("{name}: предмета «{subject}» нет на листе «{sheet}» — пропущен").format(
                        name=row.full_name, subject=title, sheet=SUBJECTS
                    ),
                )
            else:
                wanted.append(found)
        profile, created = TeacherProfile.objects.get_or_create(user=user)
        now = set(profile.subjects.values_list("pk", flat=True))
        if {s.pk for s in wanted} == now:
            return created
        profile.subjects.set(wanted)
        return True

    def person(self, code: str, where: str) -> User | None:
        if not code:
            return None
        found = self.staff.get(code)
        if found is None:
            self.error(
                _("{where}: сотрудника с ID {code} нет на листе «{sheet}»").format(where=where, code=code, sheet=STAFF)
            )
        return found

    # --- Группы ---

    def _groups(self) -> None:
        from accounts.curators import AssignmentRefused, assign, curator_of

        section = self.report.section("groups", _("Группы"))
        seen: set[str] = set()
        for row in self.book.groups:
            where = _where(GROUPS, row.row)
            if not row.code or row.code in seen:
                self.error(_("{where}: нет кода группы или он повторяется").format(where=where))
                continue
            seen.add(row.code)
            if row.parallel not in Parallel.values:
                self.error(_("{where}: параллель должна быть от 8 до 11").format(where=where))
                continue
            language = LANGUAGES.get(row.language)
            if language is None:
                self.error(
                    _("{where}: язык «{language}» — нужен «казахский» или «русский»").format(
                        where=where, language=row.language
                    )
                )
                continue
            wanted = {
                "parallel": row.parallel,
                "letter": row.letter[:4],
                "language": language,
                "home_room": row.home_room[:40],
            }
            group = StudyGroup.objects.filter(code__iexact=row.code).first()
            is_new = group is None
            if is_new:
                group = StudyGroup.objects.create(code=row.code[:16], **wanted)
                section.created += 1
            else:
                diff = {name: value for name, value in wanted.items() if getattr(group, name) != value}
                if "parallel" in diff:
                    self.report.warn(
                        "parallel",
                        _("Группа {group}: параллель меняется {was} → {will}").format(
                            group=group.code, was=group.parallel, will=row.parallel
                        ),
                    )
                for name, value in diff.items():
                    record_change(
                        instance=group,
                        field_name=name,
                        old_value=getattr(group, name),
                        new_value=value,
                        actor=self.actor,
                        source=Source.IMPORT,
                    )
                    setattr(group, name, value)
                if diff:
                    group.save(update_fields=list(diff))
                    section.updated += 1
                else:
                    section.unchanged += 1
            group_cohort(group)
            self.groups[row.code] = group

            curator = self.person(row.curator, where)
            if curator is None:
                continue
            if curator.role != Role.CURATOR:
                self.report.warn(
                    "curator",
                    _("Группа {group}: {name} — не куратор в LMS, куратором не назначен").format(
                        group=group.code, name=user_name(curator)
                    ),
                )
                continue
            current = curator_of(group, self.day)
            if current is not None and current.curator_id == curator.pk:
                continue
            since = self.day if current is not None else self.year.starts
            if current is not None and since <= current.since:
                since = current.since + dt.timedelta(days=1)
            if not is_new:
                self.report.curator_changes.append(
                    {
                        "group": group.code,
                        "parallel": group.parallel,
                        "was": user_name(current.curator) if current else "",
                        "will": user_name(curator),
                        "since": since,
                    }
                )
            try:
                assign(group=group, curator=curator, since=since, actor=self.actor)
            except AssignmentRefused as error:
                self.report.warn("curator", _("Группа {group}: {error}").format(group=group.code, error=error))
        in_file = set(self.groups)
        self.report.groups_outside = sorted(
            g.code for g in StudyGroup.objects.filter(is_active=True) if g.code.upper() not in in_file
        )

    def group(self, code: str, where: str) -> StudyGroup | None:
        found = self.groups.get(code.upper())
        if found is None:
            self.error(_("{where}: группы {code} нет на листе «{sheet}»").format(where=where, code=code, sheet=GROUPS))
        return found

    # --- Звонки ---

    def _bells(self) -> None:
        section = self.report.section("bells", _("Звонки"))
        grids: dict[str, list[BellRow]] = defaultdict(list)
        for row in self.book.bells:
            if not row.parallels:
                self.error(
                    _("{where}: не понятно, каких параллелей звонки «{label}»").format(
                        where=_where(BELLS, row.row), label=row.label
                    )
                )
                continue
            if row.number is None:
                # «обед» и другие перемены — не урок, в звонки не идут
                self.report.skipped_bells += 1
                continue
            if row.starts is None or row.ends is None or row.starts >= row.ends:
                self.error(_("{where}: неверное время урока").format(where=_where(BELLS, row.row)))
                continue
            grids[row.label].append(row)
        for rows in grids.values():
            parallels = rows[0].parallels
            span = f"{parallels[0]}–{parallels[-1]}" if len(parallels) > 1 else str(parallels[0])
            # имя расписания звонков — ключ повторного импорта, не подпись: не переводится
            title = f"Звонки {span}"  # i18n-skip: ключ записи расписания звонков
            schedule, created = BellSchedule.objects.get_or_create(year=self.year, title=title)
            changed = created
            for row in rows:
                bell = Bell.objects.filter(schedule=schedule, number=row.number).first()
                if bell is None:
                    Bell.objects.create(
                        year=self.year, schedule=schedule, number=row.number, starts=row.starts, ends=row.ends
                    )
                    changed = True
                elif (bell.starts, bell.ends) != (row.starts, row.ends):
                    bell.starts, bell.ends = row.starts, row.ends
                    bell.save(update_fields=["starts", "ends"])
                    changed = True
                for parallel in parallels:
                    self.grids.setdefault(parallel, {})[row.number] = (row.starts, row.ends)
            extra = Bell.objects.filter(schedule=schedule).exclude(number__in=[r.number for r in rows])
            if extra.exists():
                numbers = ", ".join(str(b.number) for b in extra)
                self.report.warn(
                    "bells",
                    _("{schedule}: в LMS есть уроки сверх файла — {numbers}").format(schedule=title, numbers=numbers),
                )
            groups = [g for g in self.groups.values() if g.parallel in parallels]
            others = BellSchedule.objects.filter(year=self.year).exclude(pk=schedule.pk)
            for other in others:
                other.groups.remove(*groups)
            if set(schedule.groups.values_list("pk", flat=True)) != {g.pk for g in groups}:
                schedule.groups.set(groups)
                changed = changed or not created
            if created:
                section.created += 1
            elif changed:
                section.updated += 1
            else:
                section.unchanged += 1
        for parallel in sorted({g.parallel for g in self.groups.values()} - set(self.grids)):
            self.report.warn(
                "bells",
                _("Для {parallel} параллели звонков в файле нет — уроки пойдут по общим звонкам").format(
                    parallel=parallel
                ),
            )

    # --- Подгруппы ---

    def _subgroups(self) -> None:
        section = self.report.section("subgroups", _("Подгруппы и потоки"))
        for row in self.book.subgroups:
            where = _where(SUBGROUPS, row.row)
            subject = self.subject(row.subject, where)
            groups = [self.group(code, where) for code in row.groups]
            self.person(row.teacher, where)
            if subject is None or not row.name or not groups or any(g is None for g in groups):
                if not row.name or not groups:
                    self.error(_("{where}: нужны название подгруппы и группы потока").format(where=where))
                continue
            stream, stream_state = self._stream(_stream_name(row.name), groups)
            if stream_state == "created":
                section.created += 1
            elif stream_state == "updated":
                section.updated += 1
            number_match = re.search(r"(\d+)$", row.name)
            wanted = {
                "stream": stream,
                "subject": subject,
                "number": int(number_match.group(1)) if number_match else None,
                "room": row.room[:40],
            }
            cohort = Cohort.objects.filter(kind=CohortKind.SUBGROUP, name=row.name[:120]).first()
            if cohort is None:
                cohort = Cohort.objects.create(
                    kind=CohortKind.SUBGROUP,
                    name=row.name[:120],
                    short_name=row.name[:60],
                    # правило хранится русским исходником, экран переводит его при показе
                    rule=gettext_noop("по файлу школы"),
                    **wanted,
                )
                section.created += 1
            else:
                diff = {name: value for name, value in wanted.items() if getattr(cohort, name) != value}
                for name, value in diff.items():
                    setattr(cohort, name, value)
                if diff:
                    cohort.save(update_fields=list(diff))
                    section.updated += 1
                else:
                    section.unchanged += 1
            self.subgroups[_key(row.name)] = cohort

    def _stream(self, name: str, groups: list[StudyGroup]) -> tuple[Cohort, str]:
        """Поток из групп по названию: есть — сверить части, нет — завести."""
        parts = [group_cohort(g) for g in groups]
        stream = Cohort.objects.filter(kind=CohortKind.STREAM, name=name[:120]).first()
        if stream is None:
            return make_stream(name=name, parts=parts), "created"
        now = set(StreamPart.objects.filter(stream=stream).values_list("part_id", flat=True))
        if now == {p.pk for p in parts}:
            return stream, "same"
        return make_stream(name=name, parts=parts, stream=stream), "updated"

    # --- Уроки ---

    def _lessons(self) -> None:
        section = self.report.section("lessons", _("Уроки в неделю"))
        units: list[dict] = []
        joint: dict[str, list[LessonRow]] = defaultdict(list)
        for row in self.book.lessons:
            where = _where(LESSONS, row.row)
            if bool(row.group) == bool(row.subgroup):
                self.error(_("{where}: заполните либо «Группа», либо «Подгруппа»").format(where=where))
                continue
            if row.weekday is None or row.slot is None:
                self.error(_("{where}: не понятны день или номер урока").format(where=where))
                continue
            if row.keep_teacher not in ("", KEEP_TEACHER):
                self.error(
                    _("{where}: в колонке «{column}» — «да» или пусто").format(
                        where=where, column="Учитель не менять"  # i18n-skip: заголовок колонки книги школы
                    )
                )
                continue
            if row.joint:
                joint[row.joint].append(row)
                continue
            unit = self._unit(row, [row])
            if unit is not None:
                units.append(unit)
        for code, rows in joint.items():
            first = rows[0]
            shape = {(r.weekday, r.slot, _key(r.subject), r.teacher, _key(r.room)) for r in rows}
            if len(shape) > 1:
                self.error(
                    _(
                        "«{sheet}», совместный урок {code}: день, урок, предмет, учитель и кабинет должны совпадать"
                    ).format(sheet=LESSONS, code=code)
                )
                continue
            if any(r.subgroup for r in rows):
                self.error(
                    _("«{sheet}», совместный урок {code}: собирается из групп, не из подгрупп").format(
                        sheet=LESSONS, code=code
                    )
                )
                continue
            unit = self._unit(first, rows)
            if unit is not None:
                unit["joint"] = code
                units.append(unit)
        if self.report.errors:
            return
        self.units = units
        if self.replace_from is not None:
            self._pass_journals(units)
        seen: set[tuple[int, int, int]] = set()
        for unit in units:
            key = (unit["cohort"].pk, unit["weekday"], unit["slot"])
            if key in seen:
                self.report.warn(
                    "group", _("{lesson}: у состава два урока в {when}").format(lesson=unit["label"], when=unit["when"])
                )
            seen.add(key)
            state = self._series(unit)
            if state == "created":
                section.created += 1
            elif state == "updated":
                section.updated += 1
            else:
                section.unchanged += 1
            if unit["teacher"] is None:
                self.report.series_without_teacher += 1
        without: dict[tuple[str, str], int] = defaultdict(int)
        for unit in units:
            if unit["teacher"] is None:
                without[(unit["cohort"].name, unit["subject"].title)] += 1
        self.report.without_teacher = [
            {"cohort": cohort, "subject": subject, "count": count}
            for (cohort, subject), count in sorted(without.items())
        ]

    def _unit(self, row: LessonRow, rows: list[LessonRow]) -> dict | None:
        """Урок недели: состав, предмет, учитель, день, номер, кабинет, время."""
        where = _where(LESSONS, row.row)
        subject = self.subject(row.subject, where)
        teacher = self.person(row.teacher, where)
        if row.teacher and teacher is None:
            return None
        if row.subgroup:
            cohort = self.subgroups.get(_key(row.subgroup))
            if cohort is None:
                self.error(
                    _("{where}: подгруппы {subgroup} нет на листе «{sheet}»").format(
                        where=where, subgroup=row.subgroup, sheet=SUBGROUPS
                    )
                )
                return None
            groups = [g for g in self.groups.values() if g.pk in _stream_group_ids(cohort)]
            name = row.subgroup
        else:
            groups = [self.group(r.group, _where(LESSONS, r.row)) for r in rows]
            if any(g is None for g in groups):
                return None
            if len(groups) == 1:
                cohort = group_cohort(groups[0])
            else:
                cohort, _state = self._stream(" + ".join(sorted(g.code for g in groups)), groups)
            name = " + ".join(g.code for g in groups)
        if subject is None:
            return None
        if row.keep_teacher == KEEP_TEACHER:
            if row.teacher:
                self.report.warn(
                    "teacher",
                    _("{where}: «Учитель не менять» — ID учителя {code} не берётся").format(
                        where=where, code=row.teacher
                    ),
                )
            teacher = self._current_teacher(cohort, subject, row.weekday, row.slot)
        parallel = row.parallel or (groups[0].parallel if groups else None)
        grid = self.grids.get(parallel or 0, {})
        bell = grid.get(row.slot)
        when = f"{WEEKDAY_WORDS[row.weekday]} {row.starts:%H:%M}" if row.starts else f"{WEEKDAY_WORDS[row.weekday]}"
        if grid and bell is None:
            # сводится в одну строку на параллель и номер (`_missing_bells`)
            span = f"{row.starts:%H:%M}–{row.ends:%H:%M}" if row.starts and row.ends else ""
            self.missing_bells[(parallel, row.slot, span)].append(f"{name} {WEEKDAY_WORDS[row.weekday]}")
        elif bell and row.starts and (row.starts, row.ends) != bell:
            self.report.warn(
                "time",
                _("{name}, {day} {slot} урок: в файле {file_time}, по звонкам {bell_time}").format(
                    name=name,
                    day=WEEKDAY_WORDS[row.weekday],
                    slot=row.slot,
                    file_time=f"{row.starts:%H:%M}–{row.ends:%H:%M}",
                    bell_time=f"{bell[0]:%H:%M}–{bell[1]:%H:%M}",
                ),
            )
        return {
            "cohort": cohort,
            "subject": subject,
            "report_role": self.report_role_for(row.subject),
            "teacher": teacher,
            "weekday": row.weekday,
            "slot": row.slot,
            "room": row.room[:40],
            "starts": row.starts or (bell[0] if bell else None),
            "ends": row.ends or (bell[1] if bell else None),
            "label": f"{name} {subject.title}",
            "when": when,
            "joint": "",
        }

    def _series(self, unit: dict) -> str:
        """Журнал и еженедельный урок: найти по ключу или завести."""
        course = Course.objects.filter(subject=unit["subject"], teacher=unit["teacher"], cohort=unit["cohort"]).first()
        if course is None:
            course = Course.objects.create(
                subject=unit["subject"],
                teacher=unit["teacher"],
                cohort=unit["cohort"],
                report_role=unit.get("report_role") or ReportRole.NONE,
            )
        series = (
            LessonSeries.objects.filter(course=course, weekday=unit["weekday"], slot=unit["slot"], ends__gte=self.cut)
            .order_by("-starts")
            .first()
        )
        if series is None:
            series = LessonSeries.objects.create(
                course=course,
                weekday=unit["weekday"],
                slot=unit["slot"],
                room=unit["room"],
                # замена: новое расписание — с даты замены, прежнее до неё остаётся
                starts=self.replace_from or SERIES_START,
                ends=self.year.ends,
                created_by=self.actor,
            )
            self.new_series.append(series)
            self.touched_series.add(series.pk)
            return "created"
        self.touched_series.add(series.pk)
        if series.room == unit["room"]:
            return "same"
        series.room = unit["room"]
        series.save(update_fields=["room"])
        Lesson.objects.filter(series=series, date__gte=self.cut, status=LessonStatus.PLANNED).update(room=unit["room"])
        return "updated"

    def _current_teacher(self, cohort: Cohort, subject: Subject, weekday: int, slot: int) -> User | None:
        """Учитель, что сейчас ведёт журнал состава по предмету. Журнала нет — без учителя.

        Журналов у состава по предмету несколько — берётся учитель урока
        в той же клетке (день и номер); не нашёлся один — без учителя
        и предупреждение.
        """
        key = (cohort.pk, subject.pk, weekday, slot)
        if key not in self.current_teachers:
            courses = list(Course.objects.filter(cohort=cohort, subject=subject).select_related("teacher"))
            live = [c for c in courses if c.series.filter(ends__gte=self.cut).exists()] or courses
            teachers = {c.teacher_id: c.teacher for c in live}
            if len(teachers) > 1:
                same_cell = {
                    c.teacher_id: c.teacher
                    for c in live
                    if c.series.filter(ends__gte=self.cut, weekday=weekday, slot=slot).exists()
                }
                teachers = same_cell if len(same_cell) == 1 else teachers
            found = None
            if len(teachers) == 1:
                found = next(iter(teachers.values()))
            elif teachers and (cohort.pk, subject.pk) not in self.keep_warned:
                # одно предупреждение на пару, а не на каждую клетку
                self.keep_warned.add((cohort.pk, subject.pk))
                self.report.warn(
                    "teacher",
                    _("{cohort} {subject}: «Учитель не менять», а журналов несколько — урок без учителя").format(
                        cohort=cohort.name, subject=subject.title
                    ),
                )
            self.current_teachers[key] = found
        return self.current_teachers[key]

    def _pass_journals(self, units: list[dict]) -> None:
        """Сменился учитель пары «состав + предмет» — журнал передаётся с даты замены.

        Передаётся, только когда ясно, кому: у пары ровно один журнал теряет
        учителя и ровно один учитель (или «без учителя») приходит. Иначе —
        как без замены: новый журнал, у прежнего серии заканчиваются.
        """
        wanted: dict[tuple[int, int], dict] = {}
        for unit in units:
            pair = wanted.setdefault(
                (unit["cohort"].pk, unit["subject"].pk), {"cohort": unit["cohort"], "subject": unit["subject"]}
            )
            teacher = unit["teacher"]
            pair.setdefault("teachers", {})[teacher.pk if teacher else None] = teacher
        for pair in wanted.values():
            live = (
                Course.objects.filter(cohort=pair["cohort"], subject=pair["subject"], series__ends__gte=self.cut)
                .distinct()
                .select_related("teacher")
            )
            current = {course.teacher_id: course for course in live}
            gone = set(current) - set(pair["teachers"])
            came = set(pair["teachers"]) - set(current)
            if len(gone) != 1 or len(came) != 1:
                continue
            course = current[gone.pop()]
            teacher = pair["teachers"][came.pop()]
            if Course.objects.filter(cohort=pair["cohort"], subject=pair["subject"], teacher=teacher).exists():
                # у нового учителя журнал этого состава уже есть — уроки пойдут в него
                continue
            was = user_name(course.teacher)
            reassign(course, teacher=teacher, since=self.cut, actor=self.actor, source=Source.IMPORT)
            self.report.journals_passed.append(
                {
                    "cohort": pair["cohort"].name,
                    "subject": pair["subject"].title,
                    "was": was,
                    "will": user_name(teacher),
                }
            )
        self.report.journals_passed.sort(key=lambda row: (row["subject"], row["cohort"]))

    def _groups_of(self, cohort: Cohort) -> set[int]:
        if cohort.pk not in self.cohort_groups:
            self.cohort_groups[cohort.pk] = _stream_group_ids(cohort)
        return self.cohort_groups[cohort.pk]

    def _replace(self) -> None:
        """Замена с даты: прежние еженедельные уроки пар из файла заканчиваются накануне.

        Пара — группа и предмет: подгруппа английского считается своими группами,
        поэтому прежняя подгруппа, которой в файле нет, тоже заменяется, если
        английский этих групп в файле есть. Уроки до даты не трогаются никогда;
        с даты — без данных удаляются, с данными остаются и называются.
        """
        cut = self.replace_from
        in_file = {(gid, unit["subject"].pk) for unit in self.units for gid in self._groups_of(unit["cohort"])}

        cells = {(gid, unit["weekday"], unit["slot"]) for unit in self.units for gid in self._groups_of(unit["cohort"])}

        def covered(course: Course) -> bool:
            return any((gid, course.subject_id) in in_file for gid in self._groups_of(course.cohort))

        later = Lesson.objects.filter(date__gte=cut, series__isnull=False).values("series_id")
        rows = (
            LessonSeries.objects.filter(course__archived_at__isnull=True)
            # закончившиеся с уроками после даты — те, что прошлая замена оставила из-за данных
            .filter(Q(ends__gte=cut) | Q(pk__in=later))
            .exclude(pk__in=self.touched_series)
            .select_related("course__subject", "course__cohort", "course__teacher")
            .order_by("course__subject__title", "course__cohort__name", "weekday", "slot")
        )
        replaced: list[LessonSeries] = []
        untouched: dict[tuple[str, str, str], list[int]] = defaultdict(lambda: [0, 0])
        for series in rows:
            course = series.course
            if covered(course):
                replaced.append(series)
            elif series.ends >= cut:
                row = untouched[(course.subject.title, course.cohort.name, user_name(course.teacher))]
                row[0] += 1
                # в той же клетке у этих групп урок из файла — после замены их будет два
                groups = self._groups_of(course.cohort)
                row[1] += any((gid, series.weekday, series.slot) in cells for gid in groups)
        self.report.untouched = [
            {"subject": subject, "cohort": cohort, "teacher": teacher, "count": count, "same_cell": same_cell}
            for (subject, cohort, teacher), (count, same_cell) in sorted(untouched.items())
        ]

        lessons = (
            Lesson.objects.filter(series__in=replaced, date__gte=cut)
            .annotate(
                has_attendance=Exists(Attendance.objects.filter(lesson=OuterRef("pk"))),
                has_grades=Exists(Grade.objects.filter(lesson=OuterRef("pk"))),
                has_assignment=Exists(_assignments().filter(lesson=OuterRef("pk"))),
            )
            .select_related("course__subject", "course__cohort", "teacher")
            .order_by("date", "slot", "course__cohort__name")
        )
        gone: list[int] = []
        for lesson in lessons:
            what = _lesson_data(lesson)
            if not what:
                gone.append(lesson.pk)
                continue
            self.report.lessons_kept.append(
                {
                    "id": lesson.pk,
                    "date": lesson.date,
                    "slot": lesson.slot,
                    "cohort": lesson.course.cohort.name,
                    "subject": lesson.course.subject.title,
                    "teacher": user_name(lesson.teacher),
                    "data": ", ".join(what),
                }
            )
        if gone:
            Lesson.objects.filter(pk__in=gone).delete()
        self.report.lessons_deleted = len(gone)

        for series in replaced:
            if series.ends < cut:
                continue
            self.report.series_ended += 1
            series.ends = cut - dt.timedelta(days=1)
            if series.ends < series.starts and not Lesson.all_objects.filter(series=series).exists():
                # серия начиналась с даты замены или позже — от неё ничего не осталось
                series.delete()
            else:
                series.save(update_fields=["ends"])

        on_sheet = set(self.subgroups)
        self.report.subgroups_unused = sorted(
            cohort.name
            for cohort in Cohort.objects.filter(kind=CohortKind.SUBGROUP).prefetch_related("courses")
            if _key(cohort.name) not in on_sheet and any(covered(course) for course in cohort.courses.all())
        )

    def _missing_bells(self) -> None:
        """Урок с номером, которого нет в звонках его параллели: время у него не покажется."""
        for (parallel, slot, span), where in sorted(self.missing_bells.items()):
            shown = ", ".join(where[:6])
            if len(where) > 6:
                shown = _("{names} и ещё {rest}").format(names=shown, rest=len(where) - 6)
            lesson = (
                _("{slot} урока ({time})").format(slot=slot, time=span) if span else _("{slot} урока").format(slot=slot)
            )
            self.report.warn(
                "time",
                _(
                    "{lesson} нет в звонках {parallel} параллели — уроков в неделю: {count} ({names}). "
                    "Добавьте звонок на листе «{sheet}»"
                ).format(lesson=lesson, parallel=parallel, count=len(where), names=shown, sheet=BELLS),
            )

    def _stale(self) -> None:
        """Еженедельные уроки в LMS, которых нет в файле: не удаляются, только называются."""
        rows = (
            LessonSeries.objects.filter(ends__gte=self.day, course__archived_at__isnull=True)
            .exclude(pk__in=self.touched_series)
            .select_related("course__subject", "course__cohort", "course__teacher")
            .order_by("course__cohort__name", "weekday", "slot")
        )
        for series in rows:
            course = series.course
            self.report.stale_series.append(
                _("{subject} · {cohort} · {day} {slot} урок · {teacher}").format(
                    subject=course.subject.title,
                    cohort=course.cohort.name,
                    day=WEEKDAY_WORDS.get(series.weekday, ""),
                    slot=series.slot,
                    teacher=user_name(course.teacher) or _("учитель не назначен"),
                )
            )

    def _conflicts(self) -> None:
        """Накладки кабинетов и учителей по времени из файла — предупреждения."""
        by_day: dict[int, list[dict]] = defaultdict(list)
        for unit in getattr(self, "units", []):
            if unit["starts"] and unit["ends"]:
                by_day[unit["weekday"]].append(unit)
        for weekday, units in sorted(by_day.items()):
            units.sort(key=lambda u: (u["starts"], u["label"]))
            for i, a in enumerate(units):
                for b in units[i + 1 :]:
                    if b["starts"] >= a["ends"]:
                        break
                    when = f"{WEEKDAY_WORDS[weekday]} {max(a['starts'], b['starts']):%H:%M}"
                    if a["room"] and _key(a["room"]) == _key(b["room"]):
                        self.report.warn(
                            "room",
                            _("Накладка кабинета · каб. {room} · {when}: {first} / {second}").format(
                                room=a["room"], when=when, first=a["label"], second=b["label"]
                            ),
                        )
                    if a["teacher"] is not None and a["teacher"] == b["teacher"]:
                        self.report.warn(
                            "teacher",
                            _("Накладка учителя · {teacher} · {when}: {first} / {second}").format(
                                teacher=user_name(a["teacher"]), when=when, first=a["label"], second=b["label"]
                            ),
                        )

    def _materialize(self) -> None:
        """Строки уроков новых еженедельных — со дня импорта, не с 1 сентября."""
        calendar = school_calendar.load(self.year)
        if not calendar.quarters:
            self.report.warn("year", _("У учебного года нет четвертей — строки уроков не заведутся"))
        since = max(self.day, SERIES_START, self.replace_from or SERIES_START)
        self.report.lessons_from = since
        for series in self.new_series:
            if self.commit:
                self.report.lessons_to_create += materialize(series, calendar, since=since, actor=self.actor)
            else:
                self.report.lessons_to_create += len(series_dates(series, calendar, since=since))

    def _log(self) -> None:
        """След импорта в журнале расписания. Паролей в нём нет."""
        parts = [
            _("{section}: +{created}, изменено {updated}").format(
                section=str(s.title).lower(), created=s.created, updated=s.updated
            )
            for s in self.report.sections.values()
        ]
        if self.replace_from is not None:
            parts.append(
                _(
                    "замена с {date}: еженедельных уроков закончилось {ended}, уроков удалено {deleted}, "
                    "оставлено из-за данных {kept}"
                ).format(
                    date=f"{self.replace_from:%d.%m.%Y}",
                    ended=self.report.series_ended,
                    deleted=self.report.lessons_deleted,
                    kept=len(self.report.lessons_kept),
                )
            )
        AuditLog.objects.create(
            actor=self.actor,
            actor_role=getattr(self.actor, "role", "") or "",
            model_label="academics.Lesson",
            object_id="",
            field_name=SCHEDULE_EVENT,
            domain_code="academics",
            old_value="",
            new_value=_("Импорт расписания из файла — {sections}").format(sections="; ".join(parts))[:2000],
            source=Source.IMPORT,
        )
        for change in self.report.curator_changes:
            AuditLog.objects.create(
                actor=self.actor,
                actor_role=getattr(self.actor, "role", "") or "",
                model_label="students.StudyGroup",
                object_id=change["group"],
                domain_code="",
                field_name="curator",
                old_value=change["was"],
                new_value=change["will"],
                source=Source.IMPORT,
            )


def _stream_name(subgroup: str) -> str:
    """«EEP-8-1» → «EEP-8»: поток — название подгруппы без номера."""
    match = re.fullmatch(r"(.+?)[-\s]*\d+", subgroup)
    # имя потока — ключ повторного импорта, не подпись: не переводится
    return match.group(1) if match else f"{subgroup} · поток"  # i18n-skip: ключ записи потока


def _assignments():
    from homework.models import Assignment

    return Assignment.objects.all()


def _lesson_data(lesson: Lesson) -> list[str]:
    """Что на уроке есть, кроме расписания. Пусто — урок можно удалить.

    Данные — работа учителя и учеников: посещаемость, оценки, ДЗ, тема,
    назначенный СОР или СОЧ (решение владельца, 15.10.2026). Замена, перенос
    и отмена — нет: они про прежнюю сетку. Замечания к уроку не привязаны —
    они остаются у ученика при любом раскладе.
    """
    what = []
    if lesson.marked_at is not None or lesson.has_attendance:
        what.append(_("посещаемость"))
    if lesson.has_grades:
        what.append(_("оценки"))
    if lesson.has_assignment or lesson.homework.strip():
        what.append(_("ДЗ"))
    if lesson.topic.strip():
        what.append(_("тема"))
    if lesson.kind in (LessonKind.SOR, LessonKind.SOCH):
        what.append(lesson.get_kind_display())
    return what


def _stream_group_ids(cohort: Cohort) -> set[int]:
    from academics.cohorts import group_ids_of

    return set(group_ids_of(cohort))
