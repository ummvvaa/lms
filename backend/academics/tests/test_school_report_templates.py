"""Отчёты родителям по шаблонам школы: дни, оценки, пробник, язык, тексты ИИ, файлы, права."""

from __future__ import annotations

import io
import zipfile
from unittest import mock

import pytest

from academics import marks as marking
from academics import reports as reporting
from academics import school_reports
from academics.calendar import scale_of
from academics.models import (
    Course,
    DraftState,
    EnglishLevel,
    Lesson,
    LessonKind,
    ParentReport,
    ReportPeriod,
    ReportReview,
    ReportRole,
    ReportSection,
    ReportStatus,
    ReportTemplate,
    ReviewKind,
    Scheme,
)
from academics.schedule import create_once
from academics.tests.conftest import days, login, school_day
from directories.models import SportType
from students.models import AttemptFormat, ExamAttempt, SportProfile

pytestmark = pytest.mark.django_db


def lesson_on(subjects, teacher, cohort, day, slot, subject="alg", kind=LessonKind.FO, max_score=None):
    lesson = create_once(subject=subjects[subject], teacher=teacher, cohort=cohort, date=day, slot=slot, room="1")
    if kind != LessonKind.FO:
        Lesson.objects.filter(pk=lesson.pk).update(kind=kind, max_score=max_score)
        lesson.refresh_from_db()
    return lesson


@pytest.fixture
def two_days(year, subjects, teacher, other_teacher, cohorts, calendar, pupils):
    """Два учебных дня BOSTON: в первый Дамир не был ни на одном уроке, во второй опоздал."""
    first, second = school_day(-8, calendar), school_day(-3, calendar)
    lessons = {
        "a1": lesson_on(subjects, teacher, cohorts["boston"], first, 2),
        "a2": lesson_on(subjects, teacher, cohorts["boston"], first, 3),
        "b1": lesson_on(subjects, teacher, cohorts["boston"], second, 2),
        "b2": lesson_on(subjects, teacher, cohorts["boston"], second, 3),
    }
    for key in ("a1", "a2"):
        marking.save_attendance(
            lessons[key], [{"student": pupils["damir"].pk, "mark": "absent"}], actor=teacher, calendar=calendar
        )
    marking.save_attendance(
        lessons["b1"],
        [{"student": pupils["damir"].pk, "mark": "late", "arrived": "09:35"}],
        actor=teacher,
        calendar=calendar,
    )
    marking.save_attendance(lessons["b2"], [], actor=teacher, calendar=calendar)
    return lessons


def build(student, template, calendar, *, language="ru", start=None, end=None):
    return reporting.build_report(
        student,
        kind=ReportPeriod.CUSTOM,
        start=start or days(-30),
        end=end or days(0),
        calendar=calendar,
        config=reporting.report_settings(calendar),
        template=template,
        language=language,
    )


def codes(report) -> dict[str, str]:
    return {line.code: line.value for line in report.lines.all() if line.section == ReportSection.ATTENDANCE}


# --- Посещаемость днями ----------------------------------------------------------


def test_days_count_marked_days_and_whole_missed_days(two_days, pupils, calendar):
    report = build(pupils["damir"], ReportTemplate.PROGRESS, calendar)
    got = codes(report)
    assert got["days_total"] == "2", "учебный день — день с хотя бы одним отмеченным уроком"
    assert got["days_missed"] == "1", "пропущен — все уроки дня «не был»"
    assert got["days_present"] == "1"
    assert got["late"] == "1"
    totals = __import__("academics.results", fromlist=["student_attendance"]).student_attendance(
        pupils["damir"].pk, days(-30), days(0)
    )
    assert got["pct"] == f"{totals.pct}%", "процент — по минутам, из той же Presence"
    present = build(pupils["aliya"], ReportTemplate.PROGRESS, calendar)
    assert codes(present)["days_missed"] == "0" and codes(present)["days_total"] == "2"


def test_one_absent_lesson_of_two_is_not_a_missed_day(two_days, pupils, teacher, calendar):
    marking.save_attendance(
        two_days["b2"], [{"student": pupils["nurai"].pk, "mark": "absent"}], actor=teacher, calendar=calendar
    )
    got = codes(build(pupils["nurai"], ReportTemplate.REVIEW, calendar))
    assert got["days_missed"] == "0" and got["days_present"] == "2"


# --- Оценки ----------------------------------------------------------------------


@pytest.fixture
def grades(year, subjects, teacher, cohorts, calendar, pupils):
    scale = scale_of(calendar.year)
    one = lesson_on(subjects, teacher, cohorts["boston"], school_day(-6, calendar), 2)
    two = lesson_on(subjects, teacher, cohorts["boston"], school_day(-4, calendar), 2)
    sor = lesson_on(
        subjects, teacher, cohorts["boston"], school_day(-5, calendar), 4, kind=LessonKind.SOR, max_score=20
    )
    pe = lesson_on(subjects, teacher, cohorts["boston"], school_day(-4, calendar), 5, subject="pe")
    for lesson, value in ((one, 8), (two, 9), (sor, 17), (pe, 10)):
        marking.set_grade(lesson, pupils["aliya"], value, actor=teacher, calendar=calendar, scale=scale)
    return {"one": one, "two": two}


def grade_lines(report) -> dict[str, str]:
    return {line.title: line.value for line in report.lines.all() if line.section == ReportSection.GRADES}


def test_variant_one_averages_fo_and_variant_two_lists_them(grades, pupils, calendar):
    review = grade_lines(build(pupils["aliya"], ReportTemplate.REVIEW, calendar))
    assert review["Алгебра"] == "8,5", "средняя только по ФО, СОР 17 не входит"
    assert "Физкультура" not in review, "предмет «Только ФО» в табель не идёт"
    progress = grade_lines(build(pupils["aliya"], ReportTemplate.PROGRESS, calendar))
    assert progress["Алгебра"] == "8, 9", "список ФО по порядку уроков"
    assert "Физкультура" not in progress


def test_subject_without_fo_is_an_empty_cell(cohorts, year, subjects, pupils, calendar, teacher):
    lesson_on(subjects, teacher, cohorts["boston"], school_day(-4, calendar), 2)
    lines = grade_lines(build(pupils["nurai"], ReportTemplate.REVIEW, calendar))
    assert lines["Алгебра"] == ""
    report = ParentReport.objects.get(student=pupils["nurai"], template=ReportTemplate.REVIEW)
    assert any("Нет оценок ФО" in gap for gap in school_reports.gaps(report)), "куратор видит пометку"


def test_kazakh_report_takes_kazakh_subject_title(grades, pupils, calendar, subjects):
    subjects["alg"].title_kk = "Алгебра (қаз)"
    subjects["alg"].save()
    lines = grade_lines(build(pupils["aliya"], ReportTemplate.REVIEW, calendar, language="kk"))
    assert "Алгебра (қаз)" in lines


def test_default_language_is_the_group_language(pupils, boston):
    assert school_reports.default_language(pupils["aliya"]) == "ru"
    boston.language = "kk"
    boston.save()
    pupils["aliya"].refresh_from_db()
    assert school_reports.default_language(pupils["aliya"]) == "kk"


def test_build_endpoint_uses_group_language_by_default(grades, pupils, boston, as_curator):
    boston.language = "kk"
    boston.save()
    answer = as_curator.post(
        "/api/acad/reports/build/",
        {"template": "progress", "date_from": str(days(-30)), "date_to": str(days(0)), "group": "BOSTON"},
        format="json",
    )
    assert answer.status_code == 200, answer.content
    languages = set(ParentReport.objects.filter(template="progress").values_list("language", flat=True))
    assert languages == {"kk"}
    assert answer.json()["period"].startswith("progress:kk:custom:")


# --- Пробник ----------------------------------------------------------------------


def mock_attempt(student, exam, day, **scores):
    return ExamAttempt.objects.create(
        student=student, exam_type=exam, attempt_format=AttemptFormat.MOCK, date=day, **scores
    )


def test_latest_mock_of_the_period_and_both_exams(pupils, calendar, year):
    student = pupils["aliya"]
    mock_attempt(student, "IELTS", days(-40), total_score=5)
    mock_attempt(student, "IELTS", days(-20), total_score="5.5", listening=5, reading=6, writing="5.5", speaking=6)
    mock_attempt(student, "IELTS", days(-10), total_score=6, listening="5.5", reading=7, writing="5.5", speaking="6.5")
    ExamAttempt.objects.create(
        student=student, exam_type="IELTS", attempt_format=AttemptFormat.OFFICIAL, date=days(-5), total_score=8
    )
    report = build(student, ReportTemplate.PROGRESS, calendar)
    ielts = {line.code: line.value for line in report.lines.all() if line.section == ReportSection.IELTS}
    assert ielts == {"listening": "5.5", "reading": "7", "writing": "5.5", "speaking": "6.5", "overall": "6"}
    assert not any(line.section == ReportSection.SAT for line in report.lines.all()), "SAT нет — блока нет"
    mock_attempt(student, "SAT", days(-3), total_score=1170, verbal=560, math=610)
    report = build(student, ReportTemplate.PROGRESS, calendar)
    sat = {line.code: line.value for line in report.lines.all() if line.section == ReportSection.SAT}
    assert sat == {"verbal": "560", "math": "610", "overall": "1170"}, "оба пробника — оба блока"


def test_no_mock_in_period_is_a_note_for_the_curator(pupils, calendar, year):
    mock_attempt(pupils["aliya"], "IELTS", days(-90), total_score=6)
    report = build(pupils["aliya"], ReportTemplate.PROGRESS, calendar)
    assert not any(line.section == ReportSection.IELTS for line in report.lines.all())
    assert any("Mock Test" in gap for gap in school_reports.gaps(report))


def test_variant_one_takes_level_on_period_end_and_sport(pupils, calendar, year):
    student = pupils["aliya"]
    EnglishLevel.objects.create(student=student, level="A2", since=days(-50))
    EnglishLevel.objects.create(student=student, level="B1", since=days(-5))
    EnglishLevel.objects.create(student=student, level="B2", since=days(1))
    SportProfile.objects.filter(student=student).update(sport_type=SportType.objects.create(name="Теннис"))
    report = build(student, ReportTemplate.REVIEW, calendar)
    profile = {line.code: line.value for line in report.lines.all() if line.section == ReportSection.PROFILE}
    assert profile == {"english_level": "B1", "sport": "Теннис"}


# --- Черновик ИИ ---------------------------------------------------------------------


def test_ai_is_not_called_when_there_is_no_data(pupils, calendar, year):
    from academics.report_drafts import draft

    report = build(pupils["nurai"], ReportTemplate.PROGRESS, calendar)
    with mock.patch("suggestions.llm.complete") as complete:
        draft(report)
    complete.assert_not_called()
    report.refresh_from_db()
    assert report.draft_state == DraftState.SKIPPED
    assert report.summary == "" and report.character == ""


def test_ai_unavailable_leaves_fields_empty(grades, pupils, calendar):
    from academics.report_drafts import draft
    from suggestions.llm import LLMUnavailable

    report = build(pupils["aliya"], ReportTemplate.PROGRESS, calendar)
    with mock.patch("suggestions.llm.complete", side_effect=LLMUnavailable("Лимит расходов на модель исчерпан")):
        draft(report)
    report.refresh_from_db()
    assert report.draft_state == DraftState.FAILED and "куратор" in report.draft_note
    assert report.summary == "" and not any(row.text for row in report.reviews.all())


@pytest.fixture
def sat_course(year, subjects, cohorts, make_user, calendar, pupils):
    """SAT Verbal у BOSTON с комментарием учителя к оценке."""
    from academics.models import Subject

    sat = Subject.objects.create(code="sat", title="SAT", short_title="SAT", scheme=Scheme.FO, order=9)
    trainer = make_user("teacher", "trainer@example.kz", full_name="Тренерова Мадина Ержановна")
    lesson = create_once(
        subject=sat, teacher=trainer, cohort=cohorts["boston"], date=school_day(-4, calendar), slot=6, room="1"
    )
    Course.objects.filter(pk=lesson.course_id).update(report_role=ReportRole.SAT_VERBAL)
    marking.set_grade(
        lesson,
        pupils["aliya"],
        9,
        comment="Хорошо читает тексты",
        actor=trainer,
        calendar=calendar,
        scale=scale_of(calendar.year),
    )
    return lesson.course


def test_ai_draft_fills_only_fields_with_data_and_hides_the_name(grades, sat_course, pupils, calendar, subjects):
    from academics.report_drafts import draft
    from suggestions.llm import LLMResponse

    Lesson.objects.filter(pk=grades["one"].pk).first().grades.update(comment="Старается на уроках")
    report = build(pupils["aliya"], ReportTemplate.PROGRESS, calendar)
    algebra = grades["one"].course
    answer = LLMResponse(
        content="",
        parsed={
            "eep": "Выдуманный отзыв без данных",
            "sat_verbal": "{name} хорошо читает тексты.",
            "sat_math": "Выдумка",
            "mock_comment": "Пробника нет, а текст есть",
            "character": "не для этого варианта",
            "summary": "{name} старается.",
            "subjects": [{"course": algebra.pk, "text": "{name} старается на уроках."}, {"course": 99999, "text": "x"}],
        },
        model="test",
    )
    with mock.patch("suggestions.llm.complete", return_value=answer) as complete:
        draft(report)
    sent = complete.call_args.kwargs["user"]
    assert "Ахметова" not in sent and "Алия" not in sent, "ни имени, ни фамилии в запросе"
    assert "Хорошо читает тексты" in sent and "Старается на уроках" in sent
    report.refresh_from_db()
    reviews = {row.kind: row for row in report.reviews.all()}
    assert reviews[ReviewKind.SAT_VERBAL].text == "Алия хорошо читает тексты."
    assert reviews[ReviewKind.SAT_VERBAL].teacher_name == "Тренерова Мадина Ержановна"
    assert reviews[ReviewKind.EEP].text == "", "журнала GE/EEP нет — отзыв не пишется"
    assert reviews[ReviewKind.SAT_MATH].text == ""
    assert report.mock_comment == "", "пробника за период нет — комментария нет"
    assert report.character == "", "у варианта 2 характеристики нет"
    assert report.summary == "Алия старается."
    extra = report.reviews.filter(kind=ReviewKind.SUBJECT)
    assert [(row.course_id, row.subject_title) for row in extra] == [(algebra.pk, "Алгебра")], "чужой номер отброшен"
    assert report.draft_state == DraftState.DONE


def test_redraft_keeps_curator_edit_unless_asked(grades, sat_course, pupils, calendar):
    from academics.report_drafts import draft
    from suggestions.llm import LLMResponse

    report = build(pupils["aliya"], ReportTemplate.PROGRESS, calendar)
    report.summary = "Слово куратора"
    report.save()
    answer = LLMResponse(content="", parsed={"summary": "От ИИ", "subjects": []}, model="test")
    with mock.patch("suggestions.llm.complete", return_value=answer):
        draft(report)
        report.refresh_from_db()
        assert report.summary == "Слово куратора"
        draft(report, overwrite=True)
    report.refresh_from_db()
    assert report.summary == "От ИИ"


def test_building_orders_a_draft_once(grades, pupils, as_curator, django_capture_on_commit_callbacks):
    with mock.patch("academics.report_drafts.draft") as draft, django_capture_on_commit_callbacks(execute=True):
        for _ in range(2):
            as_curator.post(
                "/api/acad/reports/build/",
                {
                    "template": "review",
                    "language": "ru",
                    "date_from": str(days(-30)),
                    "date_to": str(days(0)),
                    "student": pupils["aliya"].pk,
                },
                format="json",
            )
    assert draft.call_count == 1, "черновик заказывается при первой сборке, дальше — «Написать заново»"


# --- Правка текстов --------------------------------------------------------------------


def test_curator_edits_texts_and_removes_only_subject_blocks(grades, pupils, calendar, as_curator):
    report = build(pupils["aliya"], ReportTemplate.PROGRESS, calendar)
    extra = ReportReview.objects.create(report=report, kind=ReviewKind.SUBJECT, order=20, text="x", by_ai=True)
    verbal = report.reviews.get(kind=ReviewKind.SAT_VERBAL)
    answer = as_curator.patch(
        f"/api/acad/reports/{report.pk}/",
        {"summary": "Итог куратора", "reviews": [{"id": verbal.pk, "text": "Отзыв куратора"}]},
        format="json",
    )
    assert answer.status_code == 200
    body = answer.json()["school"]
    assert body["texts"]["summary"] == "Итог куратора"
    assert next(row for row in body["reviews"] if row["id"] == verbal.pk)["text"] == "Отзыв куратора"
    assert as_curator.delete(f"/api/acad/reports/{report.pk}/reviews/{verbal.pk}/").status_code == 400
    assert as_curator.delete(f"/api/acad/reports/{report.pk}/reviews/{extra.pk}/").status_code == 200
    assert not ReportReview.objects.filter(pk=extra.pk).exists()


# --- Файлы ---------------------------------------------------------------------------


def checked(report):
    report.status = ReportStatus.CHECKED
    report.save()
    return report


@pytest.mark.parametrize("template", [ReportTemplate.REVIEW, ReportTemplate.PROGRESS])
@pytest.mark.parametrize("language", ["ru", "kk"])
def test_word_file_is_filled_and_has_no_tags(grades, pupils, calendar, template, language, curator):
    from docx import Document

    from academics.report_files import render_docx

    report = build(pupils["aliya"], template, calendar, language=language)
    report.summary = report.character = "Текст & <проверка>"
    report.save()
    report.reviews.filter(kind=ReviewKind.SAT_VERBAL).update(text="Отзыв тренера")
    document = Document(io.BytesIO(render_docx(report)))
    text = "\n".join(p.text for p in document.paragraphs)
    cells = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert "Ахметова Алия" in text + cells
    assert "{{" not in text + cells and "{%" not in text + cells
    assert "Текст & <проверка>" in text, "спецсимволы экранируются, а не ломают файл"
    if template == ReportTemplate.PROGRESS:
        # эмодзи заголовка блока ставит шаблон школы по виду блока
        assert "\U0001f4d6 SAT Verbal" in text and "Отзыв тренера" in text
    assert "Асель Куратор" in text or "Куратор" in text
    section = document.sections[0]
    assert round(section.page_width.mm) == 210 and round(section.page_height.mm) == 297, "лист A4"


def test_pdf_and_word_come_from_one_docx(grades, pupils, calendar, as_curator):
    report = checked(build(pupils["aliya"], ReportTemplate.PROGRESS, calendar))
    pdf = as_curator.get(f"/api/acad/reports/{report.pk}/pdf/?type=pdf")
    assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF"
    word = as_curator.get(f"/api/acad/reports/{report.pk}/pdf/?type=docx")
    assert word.status_code == 200 and word.content[:2] == b"PK"
    report.refresh_from_db()
    assert report.status == ReportStatus.EXPORTED


def test_draft_is_not_downloaded_by_the_curator(grades, pupils, calendar, as_curator):
    report = build(pupils["aliya"], ReportTemplate.REVIEW, calendar)
    assert as_curator.get(f"/api/acad/reports/{report.pk}/pdf/?type=docx").status_code == 400


def test_group_zip_through_the_queue(grades, pupils, calendar, as_curator, curator, kymbat):
    rows = [checked(build(pupils[key], ReportTemplate.REVIEW, calendar)) for key in ("aliya", "damir", "nurai")]
    period = f"review:ru:custom:{rows[0].period_start}:{rows[0].period_end}"
    answer = as_curator.post(
        "/api/acad/reports/export/", {"group": "BOSTON", "period": period, "format": "docx"}, format="json"
    )
    assert answer.status_code == 200, answer.content
    job = answer.json()["job"]
    state = as_curator.get(f"/api/acad/reports/export/{job}/").json()
    assert state["state"] == "done" and state["total"] == 3
    archive = zipfile.ZipFile(io.BytesIO(as_curator.get(f"/api/acad/reports/export/{job}/file/").content))
    assert len(archive.namelist()) == 3 and all(name.endswith(".docx") for name in archive.namelist())
    assert ParentReport.objects.filter(status=ReportStatus.EXPORTED).count() == 3
    stranger = login(kymbat)
    assert stranger.get(f"/api/acad/reports/export/{job}/").status_code == 404, "чужое задание — 404"


# --- Права ---------------------------------------------------------------------------


def test_access_by_roles(grades, pupils, calendar, curator, kymbat, saltanat, admin, teacher, make_user, chicago):
    own = build(pupils["aliya"], ReportTemplate.REVIEW, calendar)
    foreign = build(pupils["stranger"], ReportTemplate.REVIEW, calendar)
    for user in (curator, kymbat, saltanat, admin):
        assert login(user).get(f"/api/acad/reports/{own.pk}/").status_code == 200, user.role
    assert login(curator).get(f"/api/acad/reports/{foreign.pk}/").status_code == 404, "чужая группа — 404"
    assert login(teacher).get(f"/api/acad/reports/{own.pk}/").status_code == 404, "учителю отчётов нет"
    assert login(pupils["aliya"].user).get(f"/api/acad/reports/{own.pk}/").status_code == 403
    answer = login(curator).post(
        "/api/acad/reports/build/",
        {"template": "review", "date_from": str(days(-30)), "date_to": str(days(0)), "group": chicago.code},
        format="json",
    )
    assert answer.status_code == 404


# --- Уровень английского и раздел журнала ---------------------------------------------------


def test_english_level_is_set_by_eep_teacher_curator_kymbat(
    eng_lesson, pupils, other_teacher, teacher, curator, kymbat, cohorts
):
    Course.objects.filter(pk=eng_lesson.course_id).update(report_role=ReportRole.EEP)
    url = f"/api/acad/students/{pupils['aliya'].pk}/english-level/"
    answer = login(other_teacher).post(url, {"level": "B1"}, format="json")
    assert answer.status_code == 200 and answer.json()["level"] == "B1"
    assert login(teacher).post(url, {"level": "B2"}, format="json").status_code in (403, 404)
    assert login(curator).post(url, {"level": "B2"}, format="json").json()["level"] == "B2", "тот же день — правка"
    assert EnglishLevel.objects.filter(student=pupils["aliya"]).count() == 1
    assert login(kymbat).post(url, {"level": "Z9"}, format="json").status_code == 400
    assert login(pupils["aliya"].user).post(url, {"level": "C1"}, format="json").status_code == 404
    student_view = login(pupils["aliya"].user).get("/api/acad/me/grades/")
    assert "english" not in student_view.json(), "ученику уровень не показывается"


def test_report_role_is_set_by_kymbat_only(lesson, as_kymbat, as_curator):
    url = f"/api/acad/journals/{lesson.course_id}/report-role/"
    assert as_curator.post(url, {"report_role": "sat_math"}, format="json").status_code == 403
    answer = as_kymbat.post(url, {"report_role": "sat_math"}, format="json")
    assert answer.status_code == 200 and answer.json()["report_role"] == "sat_math"
    assert as_kymbat.post(url, {"report_role": "nope"}, format="json").status_code == 400


def test_report_roles_command_sets_sat_by_teacher(sat_course, cohorts):
    from django.core.management import call_command

    Course.objects.filter(pk=sat_course.pk).update(report_role="")
    call_command("report_roles", "sat_math", "Тренерова", "--subject", "SAT")
    assert Course.objects.get(pk=sat_course.pk).report_role == ReportRole.SAT_MATH


def test_standard_report_still_works(grades, pupils, calendar, as_curator):
    start, end = reporting.month_bounds(days(0))
    rows = reporting.build_for_period(kind=ReportPeriod.MONTH, start=start, end=end, calendar=calendar)
    assert all(row.template == ReportTemplate.STANDARD for row in rows)
    listed = as_curator.get("/api/acad/reports/").json()
    assert listed["period"]["template"] == "standard"
