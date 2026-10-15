"""Импорт расписания и сотрудников из книги школы.

Два источника книги. Маленькая книга собирается здесь же: на ней быстро
проверяются правила — кто кем становится, что предупреждение, а что ошибка.
Большая — обезличенная копия настоящей книги школы (`fixtures/schedule_import.xlsx`):
та же структура, 50 сотрудников, 31 группа, 1108 уроков, те же накладки
и совместные уроки, а людей нет. На ней — разбор целиком и повторный импорт.
"""

from __future__ import annotations

import datetime as dt
import io
from pathlib import Path
from unittest.mock import Mock

import pytest
from openpyxl import Workbook, load_workbook

from academics import schedule_import
from academics.cohorts import group_ids_of
from academics.models import (
    AcademicYear,
    Attendance,
    Bell,
    BellSchedule,
    Cohort,
    CohortKind,
    Course,
    GradingScale,
    Lesson,
    LessonSeries,
    Mark,
    Quarter,
    Scheme,
    Subject,
    TeacherProfile,
)
from academics.rights import grades_lesson, marks_lesson
from academics.tests.conftest import days, login, make_student
from accounts.curators import assign, curator_of
from accounts.models import CuratorAssignment, Role, User
from core.models import AuditLog
from students.models import StudyGroup

REAL_BOOK = Path(__file__).parent / "fixtures" / "schedule_import.xlsx"


# --- Книги -------------------------------------------------------------------------


def _book(**rows: list[dict]) -> bytes:
    """Книга со всеми листами: `staff=`, `groups=`, … — строки словарями по заголовкам."""
    keys = {
        "Сотрудники": "staff",
        "Группы": "groups",
        "Предметы": "subjects",
        "Звонки": "bells",
        "Подгруппы": "subgroups",
        "Уроки": "lessons",
    }
    wb = Workbook()
    wb.remove(wb.active)
    wb.create_sheet("Инструкция").append(["Для людей: этот лист не импортируется"])
    for title, required in schedule_import.SHEETS.items():
        # необязательная колонка появляется, только если она есть в строках
        extra = [h for h in schedule_import.OPTIONAL.get(title, ()) if any(h in r for r in rows.get(keys[title], []))]
        headers = [*required, *extra]
        sheet = wb.create_sheet(title)
        sheet.append(headers)
        for row in rows.get(keys[title], []):
            sheet.append([row.get(h) for h in headers])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _staff(code, last, first, role, subjects="", phone="", login_="", email=""):
    return {
        "ID": code,
        "ФИО": f"{last} {first} Тестовна",
        "Фамилия": last,
        "Имя": first,
        "Отчество": "Тестовна",
        "Роль": role,
        "Предметы": subjects,
        "Телефон": phone,
        "Почта": email,
        "Логин": login_,
    }


def _lesson(parallel, day, slot, starts, ends, subject, teacher, room, group="", subgroup="", joint=""):
    return {
        "Параллель": parallel,
        "Группа": group,
        "Подгруппа": subgroup,
        "День": day,
        "Урок": slot,
        "Начало": starts,
        "Конец": ends,
        "Предмет": subject,
        "ID учителя": teacher,
        "Кабинет": room,
        "Совместный урок": joint,
    }


STAFF = [
    _staff("С01", "Алгебраева", "Айгерим", "учитель", "Алгебра, Английский язык (EEP)", "+7 700 000 00 01", "t01.test"),
    _staff("С02", "Физикова", "Дана", "учитель + куратор", "Физика", "87000000002", "t02.test"),
    _staff("С03", "Кураторова", "Сауле", "куратор", "—", "+77000000003", "c03.test"),
    _staff("С04", "Завучева", "Бану", "администрация", "Методист", "+77000000004", "a04.test"),
]
GROUPS = [
    {
        "Группа": "BOSTON",
        "Параллель": 11,
        "Литера": "В",
        "Язык обучения": "русский",
        "ID куратора": "С02",
        "Домашний кабинет": "206",
    },
    {
        "Группа": "NYU",
        "Параллель": 9,
        "Литера": "А",
        "Язык обучения": "казахский",
        "ID куратора": "С03",
        "Домашний кабинет": "401",
    },
    {
        "Группа": "YALE",
        "Параллель": 9,
        "Литера": "Б",
        "Язык обучения": "казахский",
        "ID куратора": "С03",
        "Домашний кабинет": None,
    },
]
SUBJECTS = [{"Предмет": s} for s in ("Алгебра", "Физика", "Классный час", "Creative Writing", "Английский язык (EEP)")]
BELLS = [
    {"Параллели": "11", "Урок": 1, "Начало": "10:15", "Конец": "10:55"},
    {"Параллели": "11", "Урок": 2, "Начало": "11:00", "Конец": "11:40"},
    {"Параллели": "11", "Урок": "обед", "Начало": "11:45", "Конец": "12:25"},
    {"Параллели": "8–9", "Урок": 1, "Начало": "8:00", "Конец": "8:40"},
    {"Параллели": "8–9", "Урок": 2, "Начало": "8:45", "Конец": "9:25"},
]
SUBGROUPS = [
    {
        "Подгруппа": "EEP-9-1",
        "Параллель": 9,
        "Предмет": "Английский язык (EEP)",
        "Из групп": "NYU, YALE",
        "ID учителя": "С01",
        "Основной кабинет": "305",
    }
]
LESSONS = [
    _lesson(11, "Пн", 1, "10:15", "10:55", "Алгебра", "С01", "206", group="BOSTON"),
    # учителя нет ни в расписании, ни в списке — это не ошибка
    _lesson(11, "Пн", 2, "11:00", "11:40", "Creative Writing", "", "206", group="BOSTON"),
    _lesson(11, "Вт", 1, "10:15", "10:55", "Классный час", "С02", "206", group="BOSTON"),
    # совместный урок: один урок для двух групп, не накладка
    _lesson(9, "Пн", 1, "8:00", "8:40", "Физика", "С02", "актовый", group="NYU", joint="СОВМ-01"),
    _lesson(9, "Пн", 1, "8:00", "8:40", "Физика", "С02", "актовый", group="YALE", joint="СОВМ-01"),
    _lesson(9, "Вт", 2, "8:45", "9:25", "Английский язык (EEP)", "С01", "305", subgroup="EEP-9-1"),
    # накладка кабинета 401 — предупреждение
    _lesson(9, "Ср", 1, "8:00", "8:40", "Алгебра", "С01", "401", group="NYU"),
    _lesson(9, "Ср", 1, "8:00", "8:40", "Физика", "С02", "401", group="YALE"),
]


def small_book(**override) -> bytes:
    sheets = {
        "staff": STAFF,
        "groups": GROUPS,
        "subjects": SUBJECTS,
        "bells": BELLS,
        "subgroups": SUBGROUPS,
        "lessons": LESSONS,
    }
    sheets.update(override)
    return _book(**sheets)


def _apply(data: bytes, actor) -> dict:
    return schedule_import.apply(data, actor=actor, expected=schedule_import.fingerprint(data))


def _counts() -> tuple:
    return (
        User.objects.count(),
        StudyGroup.objects.count(),
        Cohort.objects.count(),
        Course.objects.count(),
        LessonSeries.objects.count(),
        Lesson.objects.count(),
        BellSchedule.objects.count(),
        Bell.objects.count(),
        CuratorAssignment.objects.count(),
        TeacherProfile.objects.count(),
    )


def _section(report: dict, code: str) -> dict:
    return next(s for s in report["sections"] if s["code"] == code)


@pytest.fixture
def short_year(db) -> AcademicYear:
    """Текущий год с одной короткой четвертью: строк уроков немного, прогон быстрый."""
    year = AcademicYear.objects.create(title="2026–2027", starts=days(-40), ends=days(13), is_current=True)
    Quarter.objects.create(year=year, number=1, title="1 четверть", starts=days(-40), ends=days(13))
    GradingScale.objects.create(year=year)
    return year


# --- Разбор ------------------------------------------------------------------------


def test_real_book_parses_every_sheet_and_skips_the_ones_for_people():
    book = schedule_import.parse(REAL_BOOK.read_bytes())
    assert book.errors == []
    assert len(book.staff) == 50
    assert len(book.groups) == 31
    assert len(book.subjects) == 24
    assert len(book.subgroups) == 51
    assert len(book.lessons) == 1108
    # у каждой строки урока — либо группа, либо подгруппа
    assert all(bool(row.group) != bool(row.subgroup) for row in book.lessons)
    assert sum(1 for row in book.lessons if not row.teacher) == 11
    assert {row.role_text for row in book.staff} == {"учитель", "куратор", "администрация"}
    assert {tuple(row.parallels) for row in book.bells} == {(8, 9), (10,), (11,)}
    grading = {row.title: row.grading for row in book.subjects}
    assert grading["SAT"] == grading["Классный час"] == grading["Профориентация"] == "без оценок"
    assert grading["Алгебра"] == "фо, сор, соч"


def test_cells_are_read_by_meaning_not_by_format():
    assert schedule_import._time("8:00") == dt.time(8, 0)
    assert schedule_import._time(dt.time(13, 15)) == dt.time(13, 15)
    assert schedule_import._time("обед") is None
    assert schedule_import._parallels("8–9") == [8, 9]
    assert schedule_import._parallels("10") == [10]
    assert schedule_import._text(202.0) == "202"
    assert schedule_import._key("  Ахметова   Алия ") == "ахметова алия"
    assert schedule_import._stream_name("EEP-8-1") == "EEP-8"
    assert schedule_import._stream_name("GE-10-12") == "GE-10"


def test_missing_sheet_and_missing_column_are_errors_before_the_database():
    wb = load_workbook(io.BytesIO(small_book()))
    del wb["Звонки"]
    buffer = io.BytesIO()
    wb.save(buffer)
    report = schedule_import.preview(buffer.getvalue())
    assert report["errors"] == ["В файле нет листов: «Звонки»"]

    wb = load_workbook(io.BytesIO(small_book()))
    wb["Уроки"].delete_cols(10)  # «Кабинет»
    buffer = io.BytesIO()
    wb.save(buffer)
    assert "нет колонок: «Кабинет»" in schedule_import.parse(buffer.getvalue()).errors[0]
    assert schedule_import.parse(b"not a workbook").errors == ["Файл не читается как книга Excel (.xlsx)"]


# --- Маленькая книга: правила ---------------------------------------------------------


def test_preview_writes_nothing_and_counts_what_apply_will_do(short_year, admin):
    before = _counts()
    data = small_book()
    report = schedule_import.preview(data)
    assert _counts() == before
    assert report["errors"] == []
    assert report["applied"] is False
    assert report["credentials"] == []
    assert _section(report, "staff")["created"] == 4
    assert _section(report, "groups")["created"] == 3
    # 8 строк: совместный урок — один на две группы
    assert _section(report, "lessons")["created"] == 7

    applied = _apply(data, admin)
    assert applied["applied"] is True
    assert applied["sections"] == report["sections"]
    assert applied["lessons_to_create"] == report["lessons_to_create"] > 0


def test_apply_builds_people_groups_bells_subgroups_and_lessons(short_year, admin):
    report = _apply(small_book(), admin)
    assert report["errors"] == []

    teacher = User.objects.get(login="t01.test")
    assert teacher.role == Role.TEACHER
    assert teacher.email is None
    assert teacher.phone == "+77000000001"
    assert teacher.must_change_password and teacher.temp_password_expires_at is not None
    assert set(teacher.teacher_profile.subjects.values_list("title", flat=True)) == {"Алгебра", "Английский язык (EEP)"}

    # «учитель + куратор» — куратор с профилем учителя
    both = User.objects.get(login="t02.test")
    assert both.role == Role.CURATOR
    assert list(both.teacher_profile.subjects.values_list("title", flat=True)) == ["Физика"]
    assert not TeacherProfile.objects.filter(user__login="c03.test").exists()

    # администрация — выключена, без пароля, без прав администратора, в отчёте
    staff = User.objects.get(login="a04.test")
    assert not staff.is_active and not staff.has_usable_password()
    assert staff.role != Role.ADMIN and not staff.is_staff and not staff.is_superuser
    assert [row["login"] for row in report["needs_role"]] == ["a04.test"]
    assert {row["login"] for row in report["credentials"]} == {"t01.test", "t02.test", "c03.test"}

    boston = StudyGroup.objects.get(code="BOSTON")
    assert (boston.parallel, boston.letter, boston.language, boston.home_room) == (11, "В", "ru", "206")
    assert StudyGroup.objects.get(code="YALE").home_room == ""
    assert curator_of(boston).curator == both
    assert curator_of(StudyGroup.objects.get(code="NYU")).curator.login == "c03.test"

    # звонки: две сетки, обед не урок, группы привязаны по параллели
    eleven = BellSchedule.objects.get(title="Звонки 11")
    assert list(eleven.bells.values_list("number", flat=True).order_by("number")) == [1, 2]
    assert set(eleven.groups.values_list("code", flat=True)) == {"BOSTON"}
    assert set(BellSchedule.objects.get(title="Звонки 8–9").groups.values_list("code", flat=True)) == {"NYU", "YALE"}
    assert report["skipped_bells"] == 1

    # подгруппа внутри потока: группы — группы потока, учеников пока нет
    eep = Cohort.objects.get(kind=CohortKind.SUBGROUP, name="EEP-9-1")
    assert eep.group_id is None and eep.stream.name == "EEP-9" and eep.room == "305"
    assert {StudyGroup.objects.get(pk=pk).code for pk in group_ids_of(eep)} == {"NYU", "YALE"}
    assert not eep.memberships.exists()

    # совместный урок — один журнал на поток из двух групп
    physics = Course.objects.get(subject__title="Физика", cohort__kind=CohortKind.STREAM)
    assert physics.series.count() == 1
    assert {StudyGroup.objects.get(pk=pk).code for pk in group_ids_of(physics.cohort)} == {"NYU", "YALE"}

    # урок без учителя: журнал и строки без учителя
    creative = Course.objects.get(subject__title="Creative Writing")
    assert creative.teacher is None
    assert creative.lessons.exists() and not creative.lessons.exclude(teacher=None).exists()
    assert report["series_without_teacher"] == 1

    # еженедельный урок с 1 сентября, строки — не раньше дня импорта
    series = LessonSeries.objects.filter(course__subject__title="Алгебра").first()
    assert series.starts == schedule_import.SERIES_START
    assert not Lesson.objects.filter(date__lt=days(0)).exists()

    warnings = [w["text"] for w in report["warnings"]]
    assert any(text.startswith("Накладка кабинета · каб. 401 · Ср 08:00") for text in warnings)
    assert not any("актовый" in text for text in warnings)


def test_second_import_of_the_same_file_duplicates_nothing(short_year, admin):
    data = small_book()
    _apply(data, admin)
    before = _counts()
    report = _apply(data, admin)
    assert _counts() == before
    assert all(s["created"] == 0 and s["updated"] == 0 for s in report["sections"])
    assert report["credentials"] == []
    assert report["curator_changes"] == []


def test_passwords_are_returned_once_and_never_written_to_the_log(short_year, admin):
    report = _apply(small_book(), admin)
    passwords = [row["password"] for row in report["credentials"]]
    assert passwords
    for password in passwords:
        assert not AuditLog.objects.filter(new_value__contains=password).exists()
        assert not AuditLog.objects.filter(old_value__contains=password).exists()
    assert AuditLog.objects.filter(new_value__startswith="Импорт расписания из файла", source="import").count() == 1


def test_existing_people_are_found_by_login_name_or_phone_and_keep_how_they_sign_in(short_year, admin, make_user):
    # по телефону: почта и способ входа остаются, ФИО обновляется
    by_phone = make_user(Role.TEACHER, "old.teacher@example.kz", full_name="Старое Имя", phone="+77000000001")
    # по логину: был учителем — стал куратором, как в файле
    by_login = User.objects.create_user(
        email=None, login="t02.test", password=None, full_name="Кто-то", role=Role.TEACHER
    )
    # по ФИО: директор из администрации — его не трогают и в «назначьте роль» не зовут
    director = make_user(Role.DIRECTOR_EXAM, "kymbat.test@example.kz", full_name="завучева  бану тестовна")
    # по телефону: директор, которому файл даёт «куратора», — роль остаётся, это в отчёте
    sport = make_user(Role.DIRECTOR_SPORT, "sport.test@example.kz", full_name="Спортов Нурлан", phone="+77000000003")

    report = _apply(small_book(), admin)
    by_phone.refresh_from_db()
    assert by_phone.email == "old.teacher@example.kz" and by_phone.login is None
    assert by_phone.full_name == "Алгебраева Айгерим Тестовна"
    assert by_phone.has_usable_password()
    by_login.refresh_from_db()
    assert by_login.role == Role.CURATOR and by_login.login == "t02.test"
    director.refresh_from_db()
    assert director.role == Role.DIRECTOR_EXAM and director.is_active and director.login is None
    sport.refresh_from_db()
    assert sport.role == Role.DIRECTOR_SPORT and sport.login is None
    assert report["needs_role"] == []
    assert report["roles_kept"] == [{"full_name": "Спортов Нурлан", "role": "Директор спорта", "file": "куратор"}]
    assert any(w["kind"] == "curator" and "не куратор в LMS" in w["text"] for w in report["warnings"])
    # пароли — только новым, а новых нет
    assert report["credentials"] == []
    assert _section(report, "staff") == {
        "code": "staff",
        "title": "Сотрудники",
        "created": 0,
        "updated": 4,
        "unchanged": 0,
    }


def test_account_with_email_and_no_login_is_found_by_email(short_year, admin, make_user):
    """Как на проде: вход по почте, логина нет, ФИО короткое — находится по «Почте», дубля нет."""
    short = make_user(Role.CURATOR, "Sau.Kuratorova@School.kz", full_name="Кураторова Сауле")
    staff = [dict(row) for row in STAFF]
    staff[2]["Почта"] = "sau.kuratorova@school.kz"
    before = User.objects.count()
    report = _apply(small_book(staff=staff), admin)
    assert report["errors"] == []
    short.refresh_from_db()
    assert short.login is None
    # почта как была: импорт её не трогает
    assert short.email == "sau.kuratorova@school.kz"
    assert short.full_name == "Кураторова Сауле Тестовна"
    assert short.has_usable_password()
    assert not User.objects.filter(login="c03.test").exists()
    assert User.objects.count() == before + 3
    assert "c03.test" not in {row["login"] for row in report["credentials"]}
    assert curator_of(StudyGroup.objects.get(code="NYU")).curator == short


def test_email_of_a_student_is_an_error(short_year, admin, make_user):
    make_user(Role.STUDENT, "pupil@school.kz", full_name="Ученица")
    staff = [dict(row) for row in STAFF]
    staff[3]["Почта"] = "Pupil@school.kz"
    report = schedule_import.preview(small_book(staff=staff))
    assert report["errors"] == ["«Сотрудники», строка 5: почта pupil@school.kz — у ученика, не у сотрудника"]


def test_director_who_teaches_marks_own_lessons_and_nothing_more(short_year, admin, make_user):
    director = make_user(Role.DIRECTOR_TALENT, "arman.test@school.kz", full_name="Алгебраев А.")
    idle = make_user(Role.DIRECTOR_SPORT, "sport.idle@school.kz", full_name="Без уроков")
    staff = [dict(row) for row in STAFF]
    staff[0]["Почта"] = "arman.test@school.kz"
    report = _apply(small_book(staff=staff), admin)
    director.refresh_from_db()
    # роль не меняется — только в отчёт
    assert director.role == Role.DIRECTOR_TALENT and director.login is None
    assert report["roles_kept"] == [{"full_name": "Алгебраев А.", "role": "Директор талантов", "file": "учитель"}]
    # правило «отмечает тот, кто ведёт» — для предметов, которые ведутся в LMS
    Subject.objects.update(in_lms=True)

    client = login(director)
    assert client.get("/api/auth/me/").json()["teaches"] is True
    own = Lesson.objects.filter(teacher=director).order_by("date").first()
    # неделя первого своего урока, а не сегодняшняя: учебный год файла может ещё не начаться
    week = client.get("/api/acad/lessons/", {"from": own.date.isoformat()}).json()
    assert week["lessons"] and {row["teacher"]["id"] for row in week["lessons"]} == {director.pk}
    other = Lesson.objects.exclude(teacher=director).exclude(teacher=None).first()
    assert client.get(f"/api/acad/lessons/{own.pk}/").json()["may_mark"] is True
    assert client.get(f"/api/acad/lessons/{other.pk}/").status_code == 404
    assert marks_lesson(director, own) and not marks_lesson(director, other)
    # оценки и правка расписания — не его: права роли не расширяются
    assert not grades_lesson(director, own)
    assert client.post(f"/api/acad/lessons/{own.pk}/edit/", {"room": "999"}, format="json").status_code == 403

    # директор без уроков: как было — расписания нет
    idle_client = login(idle)
    assert idle_client.get("/api/auth/me/").json()["teaches"] is False
    assert idle_client.get("/api/acad/lessons/").status_code == 403


def test_name_is_not_guessed(short_year, admin, make_user):
    """Без совпадения ФИО, логина или телефона — новая выключенная учётка, а не похожая чужая."""
    kymbat = make_user(Role.DIRECTOR_EXAM, "kymbat.test@example.kz", full_name="Завучева Бану")
    report = _apply(small_book(), admin)
    kymbat.refresh_from_db()
    assert kymbat.login is None and kymbat.phone == ""
    assert [row["login"] for row in report["needs_role"]] == ["a04.test"]


def test_two_people_with_one_name_stop_the_import(short_year, admin, make_user):
    make_user(Role.TEACHER, "one@example.kz", full_name="Алгебраева Айгерим Тестовна")
    make_user(Role.TEACHER, "two@example.kz", full_name="Алгебраева Айгерим Тестовна")
    before = _counts()
    report = _apply(small_book(), admin)
    assert report["applied"] is False
    assert any("несколько сотрудников" in text for text in report["errors"])
    assert _counts() == before


def test_curator_change_is_listed_and_closes_the_old_assignment(short_year, admin, make_user):
    boston = StudyGroup.objects.create(code="BOSTON", parallel=11)
    old = make_user(Role.CURATOR, "old.curator@example.kz", full_name="Прежняя Кураторша")
    assign(group=boston, curator=old, since=days(-30), actor=admin)
    report = _apply(small_book(), admin)
    [change] = report["curator_changes"]
    assert (change["group"], change["parallel"], change["was"]) == ("BOSTON", 11, "Прежняя Кураторша")
    assert change["will"] == "Физикова Дана Тестовна"
    assert curator_of(boston).curator.login == "t02.test"
    assert CuratorAssignment.objects.get(curator=old).until == days(0)


def test_broken_references_are_errors_and_nothing_is_written(short_year, admin):
    lessons = [
        *LESSONS,
        _lesson(11, "Чт", 1, "10:15", "10:55", "Алгебра", "С99", "206", group="BOSTON"),
        _lesson(11, "Чт", 2, "11:00", "11:40", "Алгебра", "С01", "206", group="NOWHERE"),
        _lesson(11, "Пт", 1, "10:15", "10:55", "Алгебра", "С01", "206", group="BOSTON", subgroup="EEP-9-1"),
    ]
    before = _counts()
    report = _apply(small_book(lessons=lessons), admin)
    assert report["applied"] is False
    assert _counts() == before
    joined = " | ".join(report["errors"])
    assert "сотрудника с ID С99 нет" in joined
    assert "группы NOWHERE нет" in joined
    assert "либо «Группа», либо «Подгруппа»" in joined


def test_grading_column_sets_the_scheme_of_new_subjects_only(short_year, admin):
    # SAT уже заведён с полной схемой: импорт её не меняет, только называет расхождение
    sat = Subject.objects.create(code="sat", title="SAT", short_title="SAT", scheme=Scheme.KZ)
    subjects = [
        {"Предмет": "Алгебра", "Оценивание": "ФО,СОР,  СОЧ"},
        {"Предмет": "Физика", "Оценивание": None},
        {"Предмет": "Классный час", "Оценивание": "Без оценок"},
        {"Предмет": "Creative Writing", "Оценивание": "ФО, СОР, СОЧ"},
        {"Предмет": "Английский язык (EEP)", "Оценивание": "ФО, СОР, СОЧ"},
        {"Предмет": "SAT", "Оценивание": "без оценок"},
    ]
    report = _apply(small_book(subjects=subjects), admin)
    assert report["errors"] == []
    schemes = dict(Subject.objects.values_list("title", "scheme"))
    assert schemes["Алгебра"] == Scheme.KZ
    assert schemes["Классный час"] == Scheme.FO
    sat.refresh_from_db()
    assert sat.scheme == Scheme.KZ
    assert [w["text"] for w in report["warnings"] if w["kind"] == "scheme"] == [
        "Предмет «SAT»: в LMS — «ФО, СОР и СОЧ, итог за четверть», " "в файле — «Только ФО из 10»; схема не меняется"
    ]
    # пустая ячейка — как без колонки: ФО, СОР и СОЧ, и просьба проверить
    assert schemes["Физика"] == Scheme.KZ
    assert [w["text"] for w in report["warnings"] if w["kind"] == "subject"] == [
        "Новый предмет «Физика»: оценивание в файле не указано — ФО, СОР и СОЧ, проверьте"
    ]


def test_book_without_grading_column_works_as_before(short_year, admin):
    report = _apply(small_book(), admin)
    assert report["errors"] == []
    # без колонки — ФО, СОР и СОЧ; предметы школы, которые ведутся в LMS, — только ФО из 10
    assert set(Subject.objects.filter(in_lms=False).values_list("scheme", flat=True)) == {Scheme.KZ}
    assert set(Subject.objects.filter(in_lms=True).values_list("title", flat=True)) == {
        "Creative Writing",
        "Английский язык (EEP)",
    }
    assert set(Subject.objects.filter(in_lms=True).values_list("scheme", flat=True)) == {Scheme.FO}
    assert len([w for w in report["warnings"] if w["kind"] == "subject"]) == 3


def test_unknown_grading_is_an_error(short_year, admin):
    subjects = [*SUBJECTS, {"Предмет": "Черчение", "Оценивание": "зачёт"}]
    report = schedule_import.preview(small_book(subjects=subjects))
    assert report["errors"] == ["«Предметы», строка 7: оценивание «зачёт» — нужно «ФО, СОР, СОЧ» или «без оценок»"]


def test_apply_takes_only_the_file_seen_in_preview(short_year, admin):
    data = small_book()
    with pytest.raises(schedule_import.ImportRefused):
        schedule_import.apply(data, actor=admin, expected="0" * 64)


def test_without_current_year_import_refuses(db, admin):
    report = schedule_import.preview(small_book())
    assert report["errors"] == ["Нет текущего учебного года: заведите его на странице «Учебный год»"]


# --- Урок без учителя и классный час куратора ---------------------------------------


def test_curator_marks_the_lesson_they_teach_and_only_it(short_year, admin):
    _apply(small_book(), admin)
    curator = User.objects.get(login="t02.test")
    # временный пароль сменён — иначе дальше экрана смены пароля не пустят
    User.objects.filter(pk=curator.pk).update(must_change_password=False)
    curator.refresh_from_db()
    own = Lesson.objects.filter(course__subject__title="Классный час").first()
    # урок из книги может стоять на следующей неделе (классный час во вторник,
    # а прогон в среду) — переносим на ту же пару неделями раньше, в прошлое
    past = own.date
    while past >= days(0):
        past -= dt.timedelta(days=7)
    Lesson.objects.filter(pk=own.pk).update(date=past)
    own.refresh_from_db()
    other = Lesson.objects.filter(course__subject__title="Алгебра", course__cohort__group__code="BOSTON").first()
    # классный час ведётся в LMS только если школа так решит: правило проверяется на таком
    Subject.objects.update(in_lms=True)
    own.refresh_from_db()
    other.refresh_from_db()
    assert marks_lesson(curator, own)
    assert not marks_lesson(curator, other)
    client = login(curator)
    response = client.get(f"/api/acad/lessons/{own.pk}/")
    assert response.status_code == 200, response.json()
    assert response.json()["may_mark"] is True
    marked = client.post(f"/api/acad/lessons/{own.pk}/attendance/", {"all_present": True}, format="json")
    assert marked.status_code == 200, marked.json()


def test_lesson_without_teacher_opens_and_is_not_nagged(short_year, admin):
    _apply(small_book(), admin)
    lesson = Lesson.objects.filter(course__subject__title="Creative Writing").first()
    payload = login(admin).get(f"/api/acad/lessons/{lesson.pk}/").json()
    assert payload["lesson"]["teacher"] is None and payload["may_remind"] is False
    assert login(admin).post(f"/api/acad/lessons/{lesson.pk}/remind/").status_code == 400
    # неделя расписания с уроками без учителя в одном слоте собирается
    assert login(admin).get("/api/acad/schedule/").status_code == 200
    assert login(admin).get("/api/acad/cohorts/").status_code == 200


# --- API ------------------------------------------------------------------------------


def _upload(client, path: str, data: bytes, **extra):
    upload = io.BytesIO(data)
    upload.name = "raspisanie.xlsx"
    return client.post(path, {"file": upload, **extra}, format="multipart")


def test_only_admin_imports(short_year, admin, kymbat, curator, teacher):
    data = small_book()
    for user in (kymbat, curator):
        client = login(user)
        assert _upload(client, "/api/acad/schedule/import/preview/", data).status_code == 403
        assert _upload(client, "/api/acad/schedule/import/apply/", data, fingerprint="x").status_code == 403
    # учителю чужие маршруты закрывает шлюз — как несуществующие
    assert _upload(login(teacher), "/api/acad/schedule/import/preview/", data).status_code == 404

    client = login(admin)
    preview = _upload(client, "/api/acad/schedule/import/preview/", data)
    assert preview.status_code == 200 and preview.json()["errors"] == []
    wrong = _upload(client, "/api/acad/schedule/import/apply/", data, fingerprint="0" * 64)
    assert wrong.status_code == 400
    applied = _upload(client, "/api/acad/schedule/import/apply/", data, fingerprint=preview.json()["fingerprint"])
    assert applied.status_code == 200
    assert applied["Cache-Control"] == "private, no-store"
    assert len(applied.json()["credentials"]) == 3


def test_curator_only_import_counts_once_but_preview_and_repeat_do_not(short_year, admin, monkeypatch):
    _apply(small_book(), admin)
    changed_groups = [{**row, "ID куратора": "С03"} if row["Группа"] == "BOSTON" else row for row in GROUPS]
    data = small_book(groups=changed_groups)
    tracker = Mock()
    monkeypatch.setattr("core.usage.track", tracker)
    client = login(admin)

    preview = _upload(client, "/api/acad/schedule/import/preview/", data)
    assert preview.status_code == 200
    tracker.assert_not_called()

    applied = _upload(client, "/api/acad/schedule/import/apply/", data, fingerprint=schedule_import.fingerprint(data))
    report = applied.json()
    assert applied.status_code == 200 and report["applied"], report
    assert not any(part["created"] or part["updated"] for part in report["sections"])
    assert report["lessons_to_create"] == 0 and len(report["curator_changes"]) == 1
    assert curator_of(StudyGroup.objects.get(code="BOSTON")).curator.login == "c03.test"
    tracker.assert_called_once()
    assert tracker.call_args.args[1] == "import.schedule.apply"

    repeated = _upload(client, "/api/acad/schedule/import/apply/", data, fingerprint=schedule_import.fingerprint(data))
    assert repeated.status_code == 200 and repeated.json()["curator_changes"] == []
    tracker.assert_called_once()


def test_staff_passwords_download_as_a_book(short_year, admin):
    report = _apply(small_book(), admin)
    response = login(admin).post(
        "/api/users/handout/export/", {"rows": report["credentials"], "kind": "staff"}, format="json"
    )
    assert response.status_code == 200
    assert "paroli-sotrudnikov.xlsx" in response["Content-Disposition"]
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert [c.value for c in sheet[1]] == ["ФИО", "Почта или логин", "Временный пароль", "Срок действия ссылки"]
    assert sheet.max_row == 4


# --- Копия настоящей книги: целиком и повторно ------------------------------------------


def test_real_book_imports_once_and_the_second_time_changes_nothing(short_year, admin):
    data = REAL_BOOK.read_bytes()
    report = _apply(data, admin)
    assert report["errors"] == []
    assert _section(report, "staff")["created"] == 50
    assert _section(report, "groups")["created"] == 31
    assert _section(report, "subgroups")["created"] == 51 + 10  # подгруппы и их потоки
    # 1108 строк, 25 совместных уроков по две группы — 1083 урока в неделю
    assert _section(report, "lessons")["created"] == 1083
    assert report["series_without_teacher"] == 11
    assert len(report["needs_role"]) == 3
    assert len(report["credentials"]) == 47
    rooms = [w["text"] for w in report["warnings"] if w["kind"] == "room"]
    assert len(rooms) == 2 and all("каб. 506" in text for text in rooms)
    assert not [w for w in report["warnings"] if w["kind"] == "teacher"]
    # звонки у 8–9 и 10 — по 8 уроков: время каждого урока совпадает со звонками
    assert [w for w in report["warnings"] if w["kind"] == "time"] == []
    assert list(
        BellSchedule.objects.get(title="Звонки 8–9").bells.order_by("number").values_list("number", flat=True)
    ) == list(range(1, 9))
    # «без оценок» — «Только ФО»; предметы, которые ведутся в LMS, — тоже только ФО из 10
    assert set(Subject.objects.filter(scheme=Scheme.FO).values_list("title", flat=True)) == {
        "SAT",
        "Классный час",
        "Профориентация",
        "Английский язык (EEP)",
        "Английский язык (GE)",
        "Creative Writing",
    }
    assert Subject.objects.filter(scheme=Scheme.KZ).count() == 18
    # в LMS ведутся пять предметов школы (SAT — один), остальные — только расписание
    assert set(Subject.objects.filter(in_lms=True).values_list("title", flat=True)) == {
        "SAT",
        "Профориентация",
        "Английский язык (EEP)",
        "Английский язык (GE)",
        "Creative Writing",
    }
    assert not [w for w in report["warnings"] if w["kind"] in ("subject", "scheme")]
    assert set(BellSchedule.objects.values_list("title", flat=True)) == {"Звонки 8–9", "Звонки 10", "Звонки 11"}

    before = _counts()
    again = _apply(data, admin)
    assert _counts() == before
    assert all(s["created"] == 0 and s["updated"] == 0 for s in again["sections"])
    assert again["credentials"] == [] and again["curator_changes"] == [] and again["stale_series"] == []


# --- Замена расписания с даты -------------------------------------------------------

#: дата замены: неделя до неё и неделя после неё — внутри короткой четверти
CUT = days(7)
BOSTON_PHYSICS = _lesson(11, "Ср", 2, "11:00", "11:40", "Физика", "С02", "206", group="BOSTON")
OLD_LESSONS = [*LESSONS, BOSTON_PHYSICS]
NEW_LESSONS = [
    # алгебра BOSTON переехала с понедельника на четверг
    _lesson(11, "Чт", 1, "10:15", "10:55", "Алгебра", "С01", "206", group="BOSTON"),
    {**LESSONS[1], "Учитель не менять": "да"},
    # у классного часа учителя в файле нет: журнал остаётся без учителя с даты замены
    {**LESSONS[2], "ID учителя": ""},
    *LESSONS[3:5],
    # английский NYU и YALE — в новой подгруппе, прежней в файле нет
    {**LESSONS[5], "Подгруппа": "EEP-9-2"},
    *LESSONS[6:],
    # физики BOSTON в файле нет совсем — пара «группа + предмет» не трогается
]
NEW_SUBGROUPS = [{**SUBGROUPS[0], "Подгруппа": "EEP-9-2"}]


def _new_book() -> bytes:
    return small_book(lessons=NEW_LESSONS, subgroups=NEW_SUBGROUPS)


def _replace(data: bytes, actor, cut: dt.date = CUT) -> dict:
    return schedule_import.apply(data, actor=actor, expected=schedule_import.fingerprint(data, cut), replace_from=cut)


def _series(subject: str, cohort: str, weekday: int, slot: int) -> LessonSeries:
    return LessonSeries.objects.get(
        course__subject__title=subject, course__cohort__name=cohort, weekday=weekday, slot=slot
    )


@pytest.fixture
def old_schedule(short_year, admin) -> User:
    """Прежнее расписание; у Creative Writing BOSTON учителя назначили в LMS руками."""
    assert _apply(small_book(lessons=OLD_LESSONS), admin)["errors"] == []
    from academics.schedule import reassign

    writing = Course.objects.get(subject__title="Creative Writing", cohort__name="BOSTON")
    reassign(writing, teacher=User.objects.get(login="t02.test"), since=days(0))
    return admin


def test_replace_keeps_a_lesson_with_attendance_and_deletes_planned_ones(old_schedule):
    algebra = _series("Алгебра", "BOSTON", 1, 1)
    marked = Lesson.objects.filter(series=algebra, date__gte=CUT).first()
    planned = list(Lesson.objects.filter(series=algebra, date__gte=CUT).exclude(pk=marked.pk))
    planned += list(Lesson.objects.filter(series=_series("Английский язык (EEP)", "EEP-9-1", 2, 2), date__gte=CUT))
    assert planned
    pupil = make_student(StudyGroup.objects.get(code="BOSTON"), "Отмеченова", "Аня", "anya@example.kz")
    Attendance.objects.create(lesson=marked, student=pupil, mark=Mark.ABSENT)

    report = _replace(_new_book(), old_schedule)

    assert report["errors"] == [] and report["applied"]
    assert Lesson.objects.filter(pk=marked.pk).exists()
    assert not Lesson.objects.filter(pk__in=[row.pk for row in planned]).exists()
    kept = [row for row in report["lessons_kept"] if row["id"] == marked.pk]
    assert kept and kept[0]["data"] == "посещаемость" and kept[0]["cohort"] == "BOSTON"
    algebra.refresh_from_db()
    assert algebra.ends == CUT - dt.timedelta(days=1)
    assert report["series_ended"] >= 1 and report["lessons_deleted"] >= len(planned)
    # новый урок недели — с даты замены, а не с 1 сентября
    thursday = _series("Алгебра", "BOSTON", 4, 1)
    assert thursday.starts == CUT
    assert not Lesson.objects.filter(series=thursday, date__lt=CUT).exists()
    assert Lesson.objects.filter(series=thursday, date__gte=CUT).exists()


def test_lessons_before_the_date_are_never_touched(old_schedule):
    algebra = _series("Алгебра", "BOSTON", 1, 1)
    before = list(Lesson.objects.filter(series=algebra, date__lt=CUT).values_list("pk", "date", "teacher_id"))
    assert before

    _replace(_new_book(), old_schedule)

    assert list(Lesson.objects.filter(series=algebra, date__lt=CUT).values_list("pk", "date", "teacher_id")) == before


def test_pair_missing_from_the_file_is_not_touched(old_schedule):
    physics = _series("Физика", "BOSTON", 3, 2)
    lessons = set(Lesson.objects.filter(series=physics).values_list("pk", flat=True))
    ends = physics.ends

    report = _replace(_new_book(), old_schedule)

    physics.refresh_from_db()
    assert physics.ends == ends
    assert set(Lesson.objects.filter(series=physics).values_list("pk", flat=True)) == lessons
    assert report["untouched"] == [
        {"subject": "Физика", "cohort": "BOSTON", "teacher": "Физикова Дана Тестовна", "count": 1}
    ]
    assert report["subject_changed"] == []


def test_missing_pair_in_the_cell_of_a_file_lesson_is_replaced_as_subject_changed(old_schedule):
    """Физику BOSTON по средам вторым уроком в новом файле занял другой предмет."""
    lessons = [*NEW_LESSONS, {**LESSONS[1], "День": "Ср", "Учитель не менять": "да"}]
    physics = _series("Физика", "BOSTON", 3, 2)
    before_cut = set(Lesson.objects.filter(series=physics, date__lt=CUT).values_list("pk", flat=True))
    assert Lesson.objects.filter(series=physics, date__gte=CUT).exists()

    report = _replace(small_book(lessons=lessons, subgroups=NEW_SUBGROUPS), old_schedule)

    physics.refresh_from_db()
    assert physics.ends == CUT - dt.timedelta(days=1)
    assert not Lesson.objects.filter(series=physics, date__gte=CUT).exists()
    assert set(Lesson.objects.filter(series=physics, date__lt=CUT).values_list("pk", flat=True)) == before_cut
    assert report["subject_changed"] == [
        {"subject": "Физика", "cohort": "BOSTON", "teacher": "Физикова Дана Тестовна", "count": 1}
    ]
    assert not [row for row in report["untouched"] if row["subject"] == "Физика"]
    # повтор с той же датой ничего не меняет и пару больше не называет
    again = _replace(small_book(lessons=lessons, subgroups=NEW_SUBGROUPS), old_schedule)
    assert again["series_ended"] == 0 and again["lessons_deleted"] == 0 and again["subject_changed"] == []


def test_keep_teacher_column_keeps_the_teacher_of_the_journal(old_schedule):
    writing = Course.objects.get(subject__title="Creative Writing", cohort__name="BOSTON")
    teacher = User.objects.get(login="t02.test")
    assert writing.teacher == teacher

    report = _replace(_new_book(), old_schedule)

    assert report["errors"] == []
    assert Course.objects.filter(subject__title="Creative Writing", cohort__name="BOSTON").get() == writing
    writing.refresh_from_db()
    assert writing.teacher == teacher
    assert not Lesson.objects.filter(course=writing, date__gte=CUT).exclude(teacher=teacher).exists()
    assert not [row for row in report["without_teacher"] if row["subject"] == "Creative Writing"]


def test_keep_teacher_without_a_journal_is_a_lesson_without_teacher(short_year, admin):
    lessons = [{**LESSONS[1], "Учитель не менять": "да", "ID учителя": ""}]
    report = _apply(small_book(lessons=lessons), admin)
    assert report["errors"] == []
    assert Course.objects.get(subject__title="Creative Writing").teacher is None
    assert report["without_teacher"] == [{"cohort": "BOSTON", "subject": "Creative Writing", "count": 1}]


def test_unknown_keep_teacher_value_is_an_error(short_year, admin):
    report = schedule_import.preview(small_book(lessons=[{**LESSONS[1], "Учитель не менять": "может быть"}]))
    assert report["errors"] and "Учитель не менять" in report["errors"][0]


def test_teacher_change_passes_the_journal_from_the_date(old_schedule):
    homeroom = Course.objects.get(subject__title="Классный час", cohort__name="BOSTON")
    was = homeroom.teacher
    series = _series("Классный час", "BOSTON", 2, 1)

    report = _replace(_new_book(), old_schedule)

    # журнал тот же, учителя у него нет с даты замены; прежние уроки — за прежним
    assert Course.objects.filter(subject__title="Классный час", cohort__name="BOSTON").get() == homeroom
    homeroom.refresh_from_db()
    assert homeroom.teacher is None
    assert not Lesson.objects.filter(course=homeroom, date__gte=CUT, teacher__isnull=False).exists()
    assert not Lesson.objects.filter(course=homeroom, date__lt=CUT).exclude(teacher=was).exists()
    series.refresh_from_db()
    assert series.ends == AcademicYear.objects.get().ends
    assert report["journals_passed"] == [
        {"cohort": "BOSTON", "subject": "Классный час", "was": "Физикова Дана Тестовна", "will": ""}
    ]
    assert AuditLog.objects.filter(new_value__contains="без учителя вместо Физикова", source="import").exists()


def test_subgroup_missing_from_the_file_is_replaced_and_listed_unused(old_schedule):
    old = _series("Английский язык (EEP)", "EEP-9-1", 2, 2)

    report = _replace(_new_book(), old_schedule)

    old.refresh_from_db()
    assert old.ends == CUT - dt.timedelta(days=1)
    assert not Lesson.objects.filter(series=old, date__gte=CUT).exists()
    assert Cohort.objects.filter(kind=CohortKind.SUBGROUP, name="EEP-9-1").exists()
    assert report["subgroups_unused"] == ["EEP-9-1"]
    assert _series("Английский язык (EEP)", "EEP-9-2", 2, 2).starts == CUT


def test_repeat_of_the_same_file_and_date_changes_nothing(old_schedule):
    algebra = _series("Алгебра", "BOSTON", 1, 1)
    marked = Lesson.objects.filter(series=algebra, date__gte=CUT).first()
    marked.topic = "Квадратные уравнения"
    marked.save(update_fields=["topic"])
    data = _new_book()
    first = _replace(data, old_schedule)
    assert first["series_ended"] and first["lessons_deleted"] and first["journals_passed"]
    before = _counts()
    ends = dict(LessonSeries.objects.values_list("pk", "ends"))
    teachers = dict(Lesson.objects.values_list("pk", "teacher_id"))

    again = _replace(data, old_schedule)

    assert _counts() == before
    assert dict(LessonSeries.objects.values_list("pk", "ends")) == ends
    assert dict(Lesson.objects.values_list("pk", "teacher_id")) == teachers
    assert again["series_ended"] == 0 and again["lessons_deleted"] == 0 and again["journals_passed"] == []
    assert again["lessons_to_create"] == 0
    assert all(s["created"] == 0 and s["updated"] == 0 for s in again["sections"])
    # урок с данными по-прежнему назван: решение о нём за человеком
    assert [row["id"] for row in again["lessons_kept"]] == [marked.pk]
    assert again["lessons_kept"][0]["data"] == "тема"


def test_preview_of_replace_writes_nothing_and_matches_apply(old_schedule):
    data = _new_book()
    before = _counts()
    seen = schedule_import.preview(data, replace_from=CUT)
    assert _counts() == before
    done = _replace(data, old_schedule)
    for key in ("series_ended", "lessons_deleted", "lessons_to_create", "journals_passed", "subgroups_unused"):
        assert seen[key] == done[key], key
    assert len(seen["lessons_kept"]) == len(done["lessons_kept"])


def test_replace_date_is_part_of_the_fingerprint_and_not_in_the_past(old_schedule):
    data = _new_book()
    with pytest.raises(schedule_import.ImportRefused):
        schedule_import.apply(data, actor=old_schedule, expected=schedule_import.fingerprint(data), replace_from=CUT)
    with pytest.raises(schedule_import.ImportRefused):
        schedule_import.preview(data, replace_from=days(-1))
    late = schedule_import.preview(data, replace_from=days(400))
    assert late["errors"]


def test_without_replace_the_import_only_lists_what_is_missing(old_schedule):
    algebra = _series("Алгебра", "BOSTON", 1, 1)
    count = Lesson.objects.filter(series=algebra).count()
    report = _apply(_new_book(), old_schedule)
    algebra.refresh_from_db()
    assert algebra.ends == AcademicYear.objects.get().ends
    assert Lesson.objects.filter(series=algebra).count() == count
    assert any("Алгебра · BOSTON" in text for text in report["stale_series"])
    assert report["series_ended"] == 0 and report["journals_passed"] == []


def test_replace_through_the_api(old_schedule):
    client = login(old_schedule)
    data = _new_book()
    seen = _upload(client, "/api/acad/schedule/import/preview/", data, replace_from=CUT.isoformat()).json()
    assert seen["replace_from"] == CUT.isoformat() and seen["series_ended"] > 0
    # отпечаток предпросмотра без замены к замене не подходит
    wrong = _upload(
        client,
        "/api/acad/schedule/import/apply/",
        data,
        fingerprint=schedule_import.fingerprint(data),
        replace_from=CUT.isoformat(),
    )
    assert wrong.status_code == 400
    done = _upload(
        client,
        "/api/acad/schedule/import/apply/",
        data,
        fingerprint=seen["fingerprint"],
        replace_from=CUT.isoformat(),
    )
    assert done.status_code == 200 and done.json()["applied"]
    bad = _upload(client, "/api/acad/schedule/import/preview/", data, replace_from="15.10.2026")
    assert bad.status_code == 400
