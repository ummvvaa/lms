"""Состав подгрупп английского файлом школы: проверка, запись, повтор.

Ученики вымышленные. Поток BOSTON + CHICAGO: две подгруппы EEP и одна GE
(у ученика одна из них). Файл — «Подгруппа | ФИО | Группа».
"""

from __future__ import annotations

import datetime as dt
import io

import pytest
from openpyxl import Workbook

from academics.cohorts import group_cohort, make_stream, member_ids
from academics.models import Cohort, CohortKind, CohortMembership, Subject
from academics.subgroup_members import plan, subgroup_key
from academics.tests.conftest import login, make_student
from core.models import AuditLog
from students.models import StudyGroup

pytestmark = pytest.mark.django_db

URL = "/api/acad/cohorts/members/"


@pytest.fixture
def english(subjects, cohorts, pupils):
    # как в школе: EEP и GE — разные потоки и предметы, но из одних групп
    parts = [cohorts["boston"], cohorts["chicago"]]
    eep, ge = make_stream(name="EEP", parts=parts), make_stream(name="GE", parts=parts)
    general = Subject.objects.create(code="ge", title="Английский язык (GE)", short_title="Англ. GE", order=4)
    made = {}
    for name, stream, subject, room in (
        ("EEP-1", eep, subjects["eng"], "203"),
        ("EEP-2", eep, subjects["eng"], "204"),
        ("GE-1", ge, general, "205"),
    ):
        made[name] = Cohort.objects.create(
            kind=CohortKind.SUBGROUP, stream=stream, subject=subject, name=name, room=room
        )
    return made


def book(*rows, header=("Подгруппа", "ФИО", "Группа"), title="Лист1") -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.append(["Состав подгрупп английского"])  # заголовок над таблицей не мешает
    sheet.append(list(header))
    for row in rows:
        sheet.append(list(row))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def upload(client, content: bytes, **extra):
    file = io.BytesIO(content)
    file.name = "состав.xlsx"
    return client.post(URL, {"file": file, **extra}, format="multipart")


FULL = (
    ("EEP-1", "Ахметова Алия", "BOSTON"),
    ("ЕЕР-1", "Сериков Дамир", "BOSTON"),  # кириллица в названии подгруппы
    ("GE-1", "Абдрахман Нурай", "BOSTON"),
    ("EEP-2", "Чужестранцев Ерлан", "CHICAGO"),
)


def test_subgroup_names_match_whatever_the_alphabet():
    assert subgroup_key("ЕЕР 8 – 1") == subgroup_key("EEP8-1") == "EEP8-1"
    assert subgroup_key("ГЕ-10.1-2") == "GE-10.1-2"


def test_preview_writes_nothing_and_apply_sets_members_from_the_year_start(english, pupils, kymbat, year):
    client = login(kymbat)
    preview = upload(client, book(*FULL))
    assert preview.status_code == 200, preview.content
    data = preview.json()
    assert data["ok"] is True and data["errors"] == []
    assert {row["name"]: row["students"] for row in data["subgroups"]} == {"EEP-1": 2, "EEP-2": 1, "GE-1": 1}
    assert not CohortMembership.objects.filter(cohort__in=english.values()).exists(), "проверка ничего не пишет"

    applied = upload(client, book(*FULL), apply="true")
    assert applied.status_code == 200, applied.content
    assert applied.json()["applied"] is True
    assert set(member_ids(english["EEP-1"])) == {pupils["aliya"].pk, pupils["damir"].pk}
    assert member_ids(english["GE-1"]) == [pupils["nurai"].pk]
    assert {row.since for row in CohortMembership.objects.filter(cohort__in=english.values())} == {year.starts}
    assert AuditLog.objects.filter(new_value__startswith="Изменён состав EEP-1", source="import").exists()

    # повтор того же файла ничего не удваивает
    upload(client, book(*FULL), apply="true")
    assert CohortMembership.objects.filter(cohort__in=english.values(), until__isnull=True).count() == 4


def test_student_in_two_subgroups_is_an_error_and_nothing_is_written(english, kymbat):
    client = login(kymbat)
    rows = (*FULL, ("GE-1", "Ахметова Алия", "BOSTON"))
    response = upload(client, book(*rows), apply="true")
    assert response.status_code == 400
    assert any("У ученика одна подгруппа" in text for text in response.json()["errors"])
    assert not CohortMembership.objects.filter(cohort__in=english.values()).exists()


def test_unknown_subgroup_student_or_group_is_named(english, kymbat):
    data = upload(
        login(kymbat),
        book(
            ("EEP-9", "Ахметова Алия", "BOSTON"),
            ("EEP-1", "Несуществующий Ученик", "BOSTON"),
            ("EEP-1", "Сериков Дамир", "LONDON"),
        ),
    ).json()
    texts = " ".join(data["errors"])
    assert "подгруппы «EEP-9» нет" in texts
    assert "ученик «Несуществующий Ученик» не найден" in texts
    assert "группы «LONDON» нет" in texts
    assert data["ok"] is False


def test_moving_to_another_english_closes_the_old_subgroup(english, pupils, kymbat):
    client = login(kymbat)
    upload(client, book(*FULL), apply="true", since="2026-09-01")
    # Алия перешла из EEP в GE; EEP-1 в новом файле нет вовсе
    moved = (("GE-1", "Ахметова Алия", "BOSTON"), ("GE-1", "Абдрахман Нурай", "BOSTON"))
    response = upload(client, book(*moved), apply="true", since="2026-10-05")
    assert response.status_code == 200, response.content
    assert pupils["aliya"].pk not in member_ids(english["EEP-1"], dt.date(2026, 10, 6))
    assert pupils["aliya"].pk in member_ids(english["EEP-1"], dt.date(2026, 9, 20)), "прошлое не переписывается"
    assert set(member_ids(english["GE-1"], dt.date(2026, 10, 6))) == {pupils["aliya"].pk, pupils["nurai"].pk}
    # Дамира в новом файле нет — его состав не тронут
    assert pupils["damir"].pk in member_ids(english["EEP-1"], dt.date(2026, 10, 6))


def test_students_left_out_and_subgroups_left_out_are_warned(english, kymbat):
    data = upload(login(kymbat), book(("EEP-1", "Ахметова Алия", "BOSTON"))).json()
    texts = " ".join(data["warnings"])
    assert "EEP-2" in texts and "состав не меняется" in texts
    assert "Сериков Дамир" in texts and "английского в расписании у них не будет" in texts


def test_sheet_named_as_a_subgroup_and_split_name_columns(english, pupils):
    content = book(("Ахметова", "Алия"), ("Сериков", "Дамир"), header=("Фамилия", "Имя"), title="EEP-1")
    found = plan(content)
    assert found.errors == []
    assert found.wanted == {english["EEP-1"].pk: [pupils["aliya"].pk, pupils["damir"].pk]}
    other = plan(book(("Ахметова", "Алия"), header=("Фамилия", "Имя"), title="Лист1"))
    assert any("нет колонки «Подгруппа»" in text for text in other.errors)


def test_without_group_column_the_student_is_found_in_the_stream_groups(english, pupils):
    found = plan(book(("EEP-2", "Чужестранцев Ерлан"), header=("Подгруппа", "Ученик")))
    assert found.errors == []
    assert found.wanted == {english["EEP-2"].pk: [pupils["stranger"].pk]}


def test_namesakes_are_not_guessed(english, boston, kymbat):
    make_student(boston, "Ахметова", "Алия", "aliya2@example.kz")
    data = upload(login(kymbat), book(("EEP-1", "Ахметова Алия", "BOSTON"))).json()
    # тёзки в одной группе: скобки не помогут, нужно полное ФИО
    assert any("подходят несколько" in text and "ФИО полностью" in text for text in data["errors"]), data["errors"]


def test_only_schedule_editors_load_members(english, curator, teacher, admin):
    content = book(*FULL)
    assert upload(login(curator), content).status_code == 403
    assert upload(login(teacher), content).status_code in (403, 404)
    assert upload(login(admin), content).status_code == 200


# --- Блоки школы: «учитель, уровень, кабинет» над нумерованным списком ---------


def school_book(blocks, *, title="Первая параллель", extra=True) -> bytes:
    """Книга, как ведёт её школа: блоки рядом и друг под другом, лишние листы рядом.

    `blocks` — (строка заголовка, колонка ФИО, заголовок, [ФИО]).
    """
    workbook = Workbook()
    if extra:
        test = workbook.active
        test.title = "Тест "
        test.append(["Аты-жөні", "Балл", "Деңгейі"])
        test.append(["Ахметова Алия", 30, "B2"])
        sheet = workbook.create_sheet(title)
    else:
        sheet = workbook.active
        sheet.title = title
    for row, column, header, names in blocks:
        sheet.cell(row=row, column=column, value=header)
        for index, name in enumerate(names, start=1):
            sheet.cell(row=row + index, column=column - 1, value=float(index))
            sheet.cell(row=row + index, column=column, value=name)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_school_blocks_go_to_subgroups_by_room(english, pupils):
    content = school_book(
        [
            (1, 2, "Тестова А, A2, 203 каб", ["Ахметова Алия", "Сериков Дамир"]),
            (1, 5, "B1, 204 каб", ["Чужестранцев Ерлан"]),
            (8, 2, "Учебная Б, C1, 205 каб", ["Абдрахман Нурай"]),
        ]
    )
    found = plan(content)
    assert found.errors == []
    assert found.wanted == {
        english["EEP-1"].pk: [pupils["aliya"].pk, pupils["damir"].pk],
        english["EEP-2"].pk: [pupils["stranger"].pk],
        english["GE-1"].pk: [pupils["nurai"].pk],
    }
    assert found.skipped == ["Тест "], "лист тестов пропущен, а не прочитан как подгруппа"


def test_report_shows_which_block_went_where(english, kymbat):
    content = school_book([(1, 2, "Тестова А, A2, 203 каб", ["Ахметова Алия"])])
    data = upload(login(kymbat), content).json()
    assert data["blocks"] == [
        {
            "sheet": "Первая параллель",
            "header": "Тестова А, A2, 203 каб",
            "subgroup": "EEP-1",
            "teacher": "",
            "students": 1,
        }
    ]
    assert data["skipped"] == ["Тест "]


def test_room_of_two_subgroups_is_decided_by_the_students(english, subjects, pupils, make_user):
    """Кабинет 203 есть и у подгруппы другого потока: блок идёт туда, где его ученики."""
    lisbon = StudyGroup.objects.create(code="LISBON", parallel=9)
    stray = make_student(lisbon, "Далёкая", "Мира", "mira@example.kz")
    other = make_stream(name="EEP-9", parts=[group_cohort(lisbon)])
    nine = Cohort.objects.create(
        kind=CohortKind.SUBGROUP, stream=other, subject=subjects["eng"], name="EEP-9-1", room="203"
    )
    found = plan(
        school_book(
            [(1, 2, "A2, 203 каб", ["Ахметова Алия", "Сериков Дамир"]), (1, 5, "A2, 203 каб", ["Далёкая Мира"])]
        )
    )
    assert found.errors == []
    assert found.wanted == {english["EEP-1"].pk: [pupils["aliya"].pk, pupils["damir"].pk], nine.pk: [stray.pk]}


def test_subgroup_named_in_the_header_wins_over_the_room(english, pupils):
    found = plan(school_book([(1, 2, "EEP-2, Тестова А, A2, 203 каб", ["Ахметова Алия"])]))
    assert found.wanted == {english["EEP-2"].pk: [pupils["aliya"].pk]}


def test_block_without_a_subgroup_or_two_blocks_in_one_are_errors(english):
    unknown = plan(school_book([(1, 2, "A2, 999 каб", ["Ахметова Алия"])]))
    assert any("нет подгруппы потока с кабинетом 999" in text for text in unknown.errors)
    twice = plan(school_book([(1, 2, "A2, 203 каб", ["Ахметова Алия"]), (1, 5, "B1, 203 каб", ["Сериков Дамир"])]))
    assert any("в подгруппу EEP-1 ведёт и блок" in text for text in twice.errors)
    assert not twice.ok


def test_namesake_is_settled_by_the_group_in_brackets(english, chicago, pupils):
    twin = make_student(chicago, "Ахметова", "Алия", "aliya.twin@example.kz")
    blocks = [(1, 2, "A2, 203 каб", ["Ахметова Алия"])]
    stuck = plan(school_book(blocks))
    # подсказка называет группы самих тёзок, а не образец
    assert any("одну из: BOSTON, CHICAGO" in text for text in stuck.errors), stuck.errors
    found = plan(school_book([(1, 2, "A2, 203 каб", ["Ахметова Алия (CHICAGO)", "Ахметова Алия (BOSTON)"])]))
    assert found.errors == [], found.errors
    assert found.wanted == {english["EEP-1"].pk: [twin.pk, pupils["aliya"].pk]}


def test_book_without_blocks_or_columns_explains_itself(english):
    found = plan(school_book([], extra=True))
    assert any("ни блоков подгрупп, ни колонок" in text for text in found.errors)


def test_student_of_another_group_is_named_with_its_group(english, subjects, year):
    """Ученик не из групп подгруппы: ошибка говорит его группу; группа в скобках — явное согласие."""
    from academics.cohorts import groups_on
    from academics.subgroup_members import apply

    lisbon = StudyGroup.objects.create(code="LISBON", parallel=11)
    stray = make_student(lisbon, "Далёкая", "Мира", "mira@example.kz")
    stuck = plan(school_book([(1, 2, "A2, 203 каб", ["Далёкая Мира"])]))
    assert any(
        "ученик LISBON, а EEP-1 собрана из групп BOSTON, CHICAGO" in text and "«Далёкая Мира (LISBON)»" in text
        for text in stuck.errors
    ), stuck.errors

    found = plan(school_book([(1, 2, "A2, 203 каб", ["Далёкая Мира (LISBON)"])]))
    assert found.errors == [], found.errors
    assert found.wanted == {english["EEP-1"].pk: [stray.pk]}
    assert any("встанет и в неделю LISBON" in text for text in found.warnings)
    apply(found, year.starts)
    assert lisbon.pk in groups_on(english["EEP-1"]), "урок подгруппы стоит в неделе группы ученика"


def test_missing_name_suggests_only_students_not_yet_in_the_file(english, pupils):
    found = plan(school_book([(1, 2, "A2, 203 каб", ["Ахметова Алия", "Сериков Данияр", "Ахметова Алина"])]))
    texts = " ".join(found.errors)
    # Дамира в файле ещё нет — он похожий; Алия уже стоит — её не предлагают
    assert "«Сериков Данияр» не найден. Похожие из тех, кого ещё нет в файле: Сериков Дамир (BOSTON)" in texts
    assert "«Ахметова Алина» не найден — напишите ФИО как в LMS" in texts
    assert "Ахметова Алия (BOSTON)" not in texts
    assert not found.ok
    # фамилия по отцу против фамилии в LMS: сходство ФИО низкое, общее имя — подсказка
    other = plan(school_book([(1, 2, "A2, 203 каб", ["Ахметова Алия", "Тестұлы Дамир"])]))
    assert any(
        "Похожие из тех, кого ещё нет в файле: Сериков Дамир (BOSTON)" in text for text in other.errors
    ), other.errors


def test_same_line_twice_in_a_subgroup_counts_once(english, pupils):
    found = plan(school_book([(1, 2, "A2, 203 каб", ["Ахметова Алия", "Сериков Дамир", "ахметова  Алия"])]))
    assert found.errors == [], found.errors
    assert found.wanted == {english["EEP-1"].pk: [pupils["aliya"].pk, pupils["damir"].pk]}
    assert any("стоит второй раз" in text and "учтён один раз" in text for text in found.warnings)
    # иначе написанная строка, похожая на того же ученика, — ошибка с обеими записями
    other = plan(school_book([(1, 2, "A2, 203 каб", ["Ахметова Алия", "Ахметова Алияя"])]))
    assert any("«Ахметова Алияя» — тот же ученик, что «Ахметова Алия»" in text for text in other.errors), other.errors
