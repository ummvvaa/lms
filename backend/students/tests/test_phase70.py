"""Фаза 70: карточка ученика строго по таблице.

Три обещания владельцу, и каждое здесь проверено:

**В карточке ровно то, что есть в таблице.** Карточки «Цели поступления»
нет ни у одной роли; сами поля остались — их спрашивает анкета первого
входа, а читают подбор вузов и стипендии.

**Блок «Поступление» одинаков у всех.** До 70-й он собирался в двух
местах: у куратора одиннадцать строк, у Асем — три. Владелец домена
видел меньше куратора, и это была не разница прав, а разница кода.

**Правят трое.** Асем как владелец домена, администратор и куратор
своей группы — телефон, почту Common App и папку. Внутренние признаки
поступления куратору по-прежнему закрыты, чужому куратору закрыто всё.
"""

# ruff: noqa: F811 — фикстуры фазы 65 импортированы по имени
from __future__ import annotations

import pytest
from rest_framework.test import APIClient

from core import domains
from students.tests.test_phase65 import (  # noqa: F401 — общие фикстуры школы
    admin,
    asem,
    boston,
    chicago,
    curator,
    klass,
    kymbat,
    stranger,
)


def login(user) -> APIClient:
    client = APIClient()
    client.force_login(user)
    return client


#: Строки блока в порядке колонок таблицы Асем — то, что человек видит
#: сверху вниз. ФИО в блоке нет: оно в шапке карточки
BLOCK_KEYS = (
    "student_phone",
    "email",
    "credentials",
    "common_app_email",
    "drive_folder_url",
    "documents",
    "gpa",
    "attempts",
)


# --- Карточка «Цели поступления» ------------------------------------------------


def test_no_field_asks_for_a_card_of_its_own():
    """Раскладка карточки знает два значения: «в блоке» и «не показывать».

    Третье — отдельная карточка под поля, которых нет в таблице, — ушло
    вместе с «Целями поступления»: правило одно на весь реестр.
    """
    placements = {spec.card for _, _, spec in domains.iter_field_specs()}
    assert placements <= {"main", "none"}


def test_goal_fields_stayed_in_the_registry():
    """Поля остались: их спрашивает анкета, а читают подбор и стипендии."""
    model = domains.DOMAINS["admission"].model("students.AdmissionProfile")
    names = {spec.name for spec in model.fields}
    assert {"target_country", "target_major", "target_level"} <= names
    # приоритет стоимости удалён: его не читал никто, кроме счётчика
    assert "cost_priority" not in names


@pytest.mark.django_db
def test_the_questionnaire_still_fills_the_goal(klass):
    """Анкета первого входа работает по-прежнему — она и заполняет цели."""
    from engagement import onboarding

    student = klass[0]
    onboarding.answer(student, code="target_country", value="Канада")

    student.admission.refresh_from_db()
    assert student.admission.target_country == "Канада"
    assert onboarding.state(student)["answered"] == 1


@pytest.mark.django_db
def test_the_picker_still_reads_the_goal(klass):
    """Подбор читает цель ученика — поле живо, хотя карточки больше нет."""
    from universities.picker import _profile_line

    student = klass[0]
    student.admission.target_country = "Канада"
    student.admission.save(update_fields=["target_country"])

    assert "цель — Канада" in _profile_line(student)


# --- Блок «Поступление» ---------------------------------------------------------


@pytest.mark.django_db
def test_the_block_is_the_same_for_curator_asem_and_admin(klass, curator, asem, admin):
    """Один состав строк у всех троих: блок собирается одним местом."""
    student = klass[0]
    by_curator = login(curator).get(f"/api/curator/students/{student.pk}/").json()["admission"]
    by_asem = login(asem).get(f"/api/students/{student.pk}/").json()["admission_block"]
    by_admin = login(admin).get(f"/api/students/{student.pk}/").json()["admission_block"]

    assert set(by_curator) == set(by_asem) == set(by_admin)
    for key in BLOCK_KEYS:
        assert key in by_curator, key


@pytest.mark.django_db
def test_the_block_holds_the_columns_of_the_table(klass, asem):
    """Ничего сверх колонок: два пароля, три документа, по три попытки."""
    block = login(asem).get(f"/api/students/{klass[0].pk}/").json()["admission_block"]

    assert [row["kind"] for row in block["credentials"]] == ["email", "common_app"]
    assert [doc["code"] for doc in block["documents"]] == ["passport", "transcript", "recommendation"]
    assert [slot["exam"] for slot in block["attempts"]] == ["IELTS", "SAT"]
    assert all(slot["slots"] == 3 for slot in block["attempts"])


@pytest.mark.django_db
def test_the_student_gets_no_block_at_all(klass):
    """У ученика свой экран: пароли и документы школы в этом составе не его."""
    student = klass[0]
    body = login(student.user).get(f"/api/students/{student.pk}/").json()
    assert body["admission_block"] is None


# --- Кто правит блок ------------------------------------------------------------


@pytest.mark.django_db
def test_curator_of_the_group_edits_the_three_fields(klass, curator):
    """Телефон, почту Common App и папку куратор правит сам (фаза 70)."""
    student = klass[0]
    answer = login(curator).patch(
        f"/api/profiles/admission/{student.pk}/",
        {"student_phone": "+77071234567"},
        format="json",
    )

    assert answer.status_code == 200, answer.content
    student.admission.refresh_from_db()
    assert student.admission.student_phone == "+77071234567"


@pytest.mark.django_db
def test_curator_does_not_touch_the_inner_marks(klass, curator):
    """Статус поступления и признаки кабинета остаются за Асем."""
    student = klass[0]
    answer = login(curator).patch(
        f"/api/profiles/admission/{student.pk}/",
        {"status": "A"},
        format="json",
    )

    assert answer.status_code == 403
    student.admission.refresh_from_db()
    assert student.admission.status == ""


@pytest.mark.django_db
def test_a_stranger_curator_edits_nothing(stranger, curator):
    """Чужая группа — чужой ученик: правки нет, и существования не видно."""
    answer = login(curator).patch(
        f"/api/profiles/admission/{stranger.pk}/",
        {"student_phone": "+77070000000"},
        format="json",
    )

    assert answer.status_code in (403, 404)
    stranger.admission.refresh_from_db()
    assert stranger.admission.student_phone == ""


@pytest.mark.django_db
def test_the_block_says_who_may_edit_it(klass, curator, asem, kymbat, stranger):
    """Право приходит с сервера: экран его не вычисляет."""
    student = klass[0]
    mine = login(curator).get(f"/api/curator/students/{student.pk}/").json()["admission"]
    owner = login(asem).get(f"/api/students/{student.pk}/").json()["admission_block"]
    other = login(kymbat).get(f"/api/students/{student.pk}/").json()["admission_block"]

    assert mine["may_edit"] is True
    assert owner["may_edit"] is True
    # директор чужого домена блок видит, но не правит (инвариант №1)
    assert other["may_edit"] is False


# --- Вузы ученика ---------------------------------------------------------------


@pytest.fixture
def programs(db):
    """Две программы справочника — на них и строится список."""
    from universities.models import Program, University

    rows = []
    for name, country in (("Toronto", "Канада"), ("Melbourne", "Австралия")):
        uni = University.objects.create(name=name, country=country)
        rows.append(Program.objects.create(university=uni, name=f"{name} CS", is_active=True))
    return rows


def add_to_list(student, program, *, by="student"):
    from universities.models import AddedBy, StudentUniversity

    return StudentUniversity.objects.create(
        student=student,
        program=program,
        added_by=AddedBy.STUDENT if by == "student" else AddedBy.DIRECTOR,
        is_confirmed=by != "student",
    )


@pytest.mark.django_db
def test_the_priority_is_exactly_one_and_goes_first(klass, programs):
    """Второй приоритетный снимает первого — «главных» вузов не бывает двух."""
    from universities.models import StudentUniversity

    student = klass[0]
    first = add_to_list(student, programs[0])
    second = add_to_list(student, programs[1])
    client = login(student.user)

    assert client.post(f"/api/catalog/priority/{first.pk}/").status_code == 200
    assert client.post(f"/api/catalog/priority/{second.pk}/").status_code == 200

    marked = list(StudentUniversity.objects.filter(student=student, is_priority=True))
    assert [row.pk for row in marked] == [second.pk]
    # и он идёт первым: порядок задан моделью, а не экраном
    assert next(iter(student.universities.all())).pk == second.pk


@pytest.mark.django_db
def test_the_student_marks_even_a_row_of_the_director(klass, programs):
    """Пометить можно и строку Асем: место в списке — выбор ученика."""
    student = klass[0]
    theirs = add_to_list(student, programs[0], by="director")

    assert login(student.user).post(f"/api/catalog/priority/{theirs.pk}/").status_code == 200
    theirs.refresh_from_db()
    assert theirs.is_priority is True
    # содержимое строки при этом не тронуто: её по-прежнему ведёт директор
    assert theirs.added_by == "director"


@pytest.mark.django_db
def test_nobody_marks_a_list_that_is_not_theirs(klass, programs):
    """Чужой список не помечают: приоритет — выбор самого ученика."""
    student, other = klass[0], klass[1]
    entry = add_to_list(student, programs[0])

    answer = login(other.user).post(f"/api/catalog/priority/{entry.pk}/")

    assert answer.status_code == 404
    entry.refresh_from_db()
    assert entry.is_priority is False


@pytest.mark.django_db
def test_the_student_removes_only_their_own_row(klass, programs):
    """Убрать можно своё; строку директора снимает тот, кто её завёл."""
    student = klass[0]
    mine = add_to_list(student, programs[0])
    theirs = add_to_list(student, programs[1], by="director")
    client = login(student.user)

    assert client.delete(f"/api/catalog/remove/{mine.pk}/").status_code in (200, 204)
    assert client.delete(f"/api/catalog/remove/{theirs.pk}/").status_code == 403


@pytest.mark.django_db
def test_the_card_shows_the_priority_first(klass, programs, curator):
    """У куратора приоритетный тоже первым и с пометкой."""
    student = klass[0]
    add_to_list(student, programs[0])
    second = add_to_list(student, programs[1])
    login(student.user).post(f"/api/catalog/priority/{second.pk}/")

    rows = login(curator).get(f"/api/curator/students/{student.pk}/").json()["universities"]

    assert rows[0]["id"] == second.pk
    assert rows[0]["is_priority"] is True


@pytest.mark.django_db
def test_the_student_changes_the_tier_of_their_own_row(klass, programs):
    """«Изменить» правит ровно то, что ученик вносил, — категорию."""
    student = klass[0]
    mine = add_to_list(student, programs[0])
    theirs = add_to_list(student, programs[1], by="director")
    client = login(student.user)

    assert client.post(f"/api/catalog/tier/{mine.pk}/", {"tier": "reach"}, format="json").status_code == 200
    mine.refresh_from_db()
    assert mine.tier == "reach"

    # строка директора не правится: её категория — решение школы
    assert client.post(f"/api/catalog/tier/{theirs.pk}/", {"tier": "reach"}, format="json").status_code == 403
    theirs.refresh_from_db()
    assert theirs.tier == "target"


@pytest.mark.django_db
def test_an_unknown_tier_is_refused_with_words(klass, programs):
    """Категория — из реестра: чужое слово не проходит."""
    student = klass[0]
    mine = add_to_list(student, programs[0])

    answer = login(student.user).post(f"/api/catalog/tier/{mine.pk}/", {"tier": "любимый"}, format="json")

    assert answer.status_code == 400
    assert "категория" in str(answer.json().get("detail", "")).lower()


# --- Контакты родителей ---------------------------------------------------------


CONTACT = {
    "full_name": "Ержанова Гульнара",
    "relation": "mother",
    "phone": "+77071112233",
    "email": "mother@example.kz",
    "preferred_channel": "",
    "note": "",
    "is_primary": True,
}


@pytest.mark.django_db
def test_curator_adds_a_contact_to_a_student_of_their_group(klass, curator):
    """Кнопки не было с 66-й: право есть, а завести контакт было нечем."""
    student = klass[0]
    answer = login(curator).post("/api/contacts/", {**CONTACT, "student": student.pk}, format="json")

    assert answer.status_code == 201, answer.content
    assert student.contacts.filter(full_name="Ержанова Гульнара").exists()


@pytest.mark.django_db
def test_curator_adds_nothing_to_a_stranger(stranger, curator):
    """Чужая группа: ученика не видно, и контакт к нему не завести."""
    answer = login(curator).post("/api/contacts/", {**CONTACT, "student": stranger.pk}, format="json")

    assert answer.status_code == 404
    assert stranger.contacts.count() == 0


@pytest.mark.django_db
def test_curator_removes_a_contact_of_their_own_group(klass, curator):
    """Убрать контакт может тот же, кто его завёл, — по своим группам."""
    student = klass[0]
    created = login(curator).post("/api/contacts/", {**CONTACT, "student": student.pk}, format="json").json()

    answer = login(curator).delete(f"/api/contacts/{created['id']}/")

    assert answer.status_code in (200, 204)
    assert student.contacts.count() == 0


@pytest.mark.django_db
def test_the_admin_keeps_contacts_everywhere(stranger, admin):
    """Администратор ведёт контакты по всей школе (фаза 68)."""
    answer = login(admin).post("/api/contacts/", {**CONTACT, "student": stranger.pk}, format="json")

    assert answer.status_code == 201, answer.content


@pytest.mark.django_db
def test_a_director_of_another_domain_adds_nothing(klass, kymbat):
    """Контакты — домен директора школы; чужому директору там нечего делать."""
    answer = login(kymbat).post("/api/contacts/", {**CONTACT, "student": klass[0].pk}, format="json")

    assert answer.status_code == 403
    assert klass[0].contacts.count() == 0


# --- Раздача паролей без писем --------------------------------------------------


@pytest.mark.django_db
def test_the_handout_sends_no_letters(klass, admin, mailoutbox):
    """Писем не уходит вовсе: пароль раздают файлом, из рук в руки (фаза 70)."""
    from accounts import handout
    from accounts.models import User

    mailoutbox.clear()
    # с галочкой «включая тех, кто уже менял пароль»: защита заданных
    # паролей — предмет фазы 69, здесь проверяется, что писем нет
    outcome = handout.issue(User.objects.filter(student__in=klass), actor=admin, include_ready=True)

    assert outcome["issued"] >= 1
    assert mailoutbox == []
    assert "Письма не рассылались" in outcome["detail"]


@pytest.mark.django_db
def test_the_file_has_a_sheet_for_every_group_and_one_for_the_staff(klass, stranger, admin):
    """Лист — группа: его отдают куратору целиком, чужого в нём нет."""
    from io import BytesIO

    from openpyxl import load_workbook

    from accounts import handout
    from accounts.models import User

    people = User.objects.filter(pk__in=[s.user_id for s in [*klass, stranger] if s.user_id] + [admin.pk])
    outcome = handout.issue(people, actor=admin, include_ready=True)
    book = load_workbook(BytesIO(handout.export(outcome["rows"]).content))

    assert "Сотрудники" in book.sheetnames
    assert {"CHICAGO", "BOSTON"} <= set(book.sheetnames)
    # в листе группы — только её ученики
    chicago_emails = {row[1] for row in list(book["CHICAGO"].iter_rows(values_only=True))[1:]}
    assert chicago_emails == {s.user.email for s in klass}
    assert admin.email in {row[1] for row in list(book["Сотрудники"].iter_rows(values_only=True))[1:]}


@pytest.mark.django_db
def test_an_empty_sheet_is_not_created(admin):
    """Лист без строк не заводится: пустую вкладку человек открывает зря."""
    from accounts import handout

    pages = handout.sheets_of([{"email": "admin@school.kz", "group": ""}])

    assert [title for title, _ in pages] == ["Сотрудники"]
