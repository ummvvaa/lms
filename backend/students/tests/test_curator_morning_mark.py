"""Отметка куратора утром и отчёт «Отсутствующие» (решение владельца, 19.10.2026).

Куратор отмечает своих учеников за сегодня, в тот же день правит; с уроками
отметка не связана; в журнале — только «не был» и исправления; отчёт по дням
и группам читают куратор (свои), директор школы и администратор (все).
"""

from __future__ import annotations

import datetime as dt
from io import BytesIO

import pytest
from django.utils import timezone
from rest_framework.test import APIClient

from accounts.curators import assign
from accounts.models import Role
from core.models import AuditLog
from students.models import AttendanceDay, BehaviorProfile, Student, StudyGroup

pytestmark = pytest.mark.django_db


@pytest.fixture
def groups(db):
    return {
        "boston": StudyGroup.objects.create(code="BOSTON", parallel=10),
        "chicago": StudyGroup.objects.create(code="CHICAGO", parallel=10),
    }


@pytest.fixture
def pupils(groups):
    rows = {
        "a": Student.objects.create(
            last_name="Абаева", first_name="Айгерим", group=groups["boston"], graduation_year=2028
        ),
        "b": Student.objects.create(
            last_name="Бекова", first_name="Дана", group=groups["boston"], graduation_year=2028
        ),
        "c": Student.objects.create(
            last_name="Серик", first_name="Ерлан", group=groups["chicago"], graduation_year=2028
        ),
    }
    for student in rows.values():
        BehaviorProfile.objects.create(student=student)
    return rows


@pytest.fixture
def admin(make_user):
    return make_user(Role.ADMIN, email="admin@example.kz")


@pytest.fixture
def curator(make_user, groups, admin):
    user = make_user(Role.CURATOR, email="curator@example.kz", full_name="Нурланова Асель")
    assign(group=groups["boston"], curator=user, since=dt.date.today() - dt.timedelta(days=10), actor=admin)
    return user


@pytest.fixture
def other_curator(make_user, groups, admin):
    user = make_user(Role.CURATOR, email="other@example.kz", full_name="Касенова Айдана")
    assign(group=groups["chicago"], curator=user, since=dt.date.today() - dt.timedelta(days=10), actor=admin)
    return user


@pytest.fixture
def saltanat(make_user):
    return make_user("director_behavior", email="saltanat@example.kz")


def client_of(user) -> APIClient:
    api = APIClient()
    api.force_login(user)
    return api


def today() -> str:
    return timezone.localdate().isoformat()


def sheet(pupils, **present) -> list[dict]:
    return [{"student": student.pk, "present": present.get(key, True)} for key, student in pupils.items() if key != "c"]


def test_curator_marks_today_and_the_journal_keeps_only_absences_and_fixes(curator, pupils):
    api = client_of(curator)
    saved = api.post(
        "/api/attendance/save/", {"group": "BOSTON", "date": today(), "rows": sheet(pupils, b=False)}, format="json"
    )
    assert saved.status_code == 200, saved.content
    body = saved.json()
    assert body["written"] == 2 and body["absent"] == 1 and body["may_mark"] is True
    assert body["noted"]["by"] == "Нурланова Асель" and body["noted"]["at"]
    assert AttendanceDay.objects.filter(date=timezone.localdate()).count() == 2

    entries = AuditLog.objects.filter(field_name="attendance_day").order_by("pk")
    assert [(e.student_id, e.new_value) for e in entries] == [
        (pupils["b"].pk, f"{timezone.localdate():%d.%m.%Y}: не был")
    ]
    assert entries[0].actor == curator and entries[0].actor_role == Role.CURATOR

    # поправил в тот же день: Дана была, Айгерим — нет
    fixed = api.post(
        "/api/attendance/save/", {"group": "BOSTON", "date": today(), "rows": sheet(pupils, a=False)}, format="json"
    )
    assert fixed.status_code == 200 and fixed.json()["absent"] == 1
    texts = sorted(AuditLog.objects.filter(field_name="attendance_day").values_list("student_id", "new_value"))
    assert texts == sorted(
        [
            (pupils["b"].pk, f"{timezone.localdate():%d.%m.%Y}: не был"),
            (pupils["b"].pk, f"{timezone.localdate():%d.%m.%Y}: был (исправлено)"),
            (pupils["a"].pk, f"{timezone.localdate():%d.%m.%Y}: не был (исправлено)"),
        ]
    )
    # один день — одна строка на ученика, повтор не дублирует
    assert AttendanceDay.objects.filter(date=timezone.localdate()).count() == 2


def test_only_today_only_own_group_only_the_curator(curator, other_curator, saltanat, admin, pupils):
    api = client_of(curator)
    yesterday = (timezone.localdate() - dt.timedelta(days=1)).isoformat()
    late = api.post(
        "/api/attendance/save/", {"group": "BOSTON", "date": yesterday, "rows": sheet(pupils)}, format="json"
    )
    assert late.status_code == 400 and "сегодняшний" in late.json()["detail"]
    assert (
        api.post("/api/attendance/save/", {"group": "CHICAGO", "date": today(), "rows": []}, format="json").status_code
        == 404
    )
    assert (
        client_of(other_curator)
        .post("/api/attendance/save/", {"group": "BOSTON", "date": today(), "rows": []}, format="json")
        .status_code
        == 404
    )

    refused = client_of(saltanat).post(
        "/api/attendance/save/", {"group": "BOSTON", "date": today(), "rows": sheet(pupils)}, format="json"
    )
    assert refused.status_code == 403 and "куратор" in refused.json()["detail"]
    assert (
        client_of(admin)
        .post("/api/attendance/save/", {"group": "BOSTON", "date": today(), "rows": []}, format="json")
        .status_code
        == 403
    )
    assert not AttendanceDay.objects.exists()

    # лист: куратору сегодня — отмечать можно, вчера — нет; Салтанат читает
    assert api.get("/api/attendance/?group=BOSTON").json()["may_mark"] is True
    assert api.get(f"/api/attendance/?group=BOSTON&date={yesterday}").json()["may_mark"] is False
    assert client_of(saltanat).get("/api/attendance/?group=BOSTON").json()["may_mark"] is False


def test_the_morning_mark_does_not_touch_lesson_attendance(curator, pupils):
    api = client_of(curator)
    api.post(
        "/api/attendance/save/", {"group": "BOSTON", "date": today(), "rows": sheet(pupils, a=False)}, format="json"
    )
    pupils["a"].behavior.refresh_from_db()
    assert pupils["a"].behavior.attendance_percent is None
    from academics.models import Attendance

    assert not Attendance.objects.exists()
    day = api.get("/api/acad/attendance/?group=BOSTON&view=day").json()
    assert day["totals"]["absent"] == 0


def test_absences_report_by_days_and_groups(curator, other_curator, saltanat, admin, pupils):
    client_of(curator).post(
        "/api/attendance/save/",
        {"group": "BOSTON", "date": today(), "rows": sheet(pupils, a=False, b=False)},
        format="json",
    )
    client_of(other_curator).post(
        "/api/attendance/save/",
        {"group": "CHICAGO", "date": today(), "rows": [{"student": pupils["c"].pk, "present": False}]},
        format="json",
    )
    # вчерашняя строка — из прежней истории, в отчёт за период тоже попадает
    AttendanceDay.objects.create(student=pupils["c"], date=timezone.localdate() - dt.timedelta(days=1), mark="absent")

    mine = client_of(curator).get("/api/attendance/absences/").json()
    assert [(row["group_code"], row["count"]) for row in mine["rows"]] == [("BOSTON", 2)]
    assert [s["full_name"] for s in mine["rows"][0]["students"]] == ["Абаева Айгерим", "Бекова Дана"]
    assert client_of(curator).get("/api/attendance/absences/?group=CHICAGO").status_code == 404

    everything = client_of(saltanat).get("/api/attendance/absences/").json()
    assert [(row["group_code"], row["count"]) for row in everything["rows"]] == [
        ("BOSTON", 2),
        ("CHICAGO", 1),
        ("CHICAGO", 1),
    ]
    since = (timezone.localdate() - dt.timedelta(days=1)).isoformat()
    narrowed = client_of(admin).get(f"/api/attendance/absences/?group=CHICAGO&from={since}&to={since}").json()
    assert [row["count"] for row in narrowed["rows"]] == [1]

    export = client_of(saltanat).get("/api/attendance/absences/export/")
    assert export.status_code == 200
    from openpyxl import load_workbook

    rows = list(load_workbook(BytesIO(export.content)).active.iter_rows(values_only=True))
    assert rows[0][:4] == ("Дата", "День", "Группа", "Кого не было")
    assert rows[1][2] == "BOSTON" and "Абаева Айгерим" in rows[1][3]

    student = client_of(pupils["a"].user) if getattr(pupils["a"], "user", None) else None
    assert student is None or student.get("/api/attendance/absences/").status_code == 403


def test_curator_has_no_teacher_reminders(curator, pupils):
    api = client_of(curator)
    refused = api.post("/api/acad/lessons/1/remind/", {}, format="json")
    assert refused.status_code == 403 and "куратору не открыт" in refused.content.decode()
    day = api.get("/api/acad/attendance/?group=BOSTON&view=day").json()
    assert day["may_remind"] is False and day["shows_unmarked"] is False and day["unmarked"] == []
