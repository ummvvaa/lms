"""Фаза 72: мастер импорта — видно, что куда ложится.

Экран — фронт, здесь то, на чём он стоит: предпросмотр отдаёт шаг
«Что заполняем» из реестра (колонка → поле → домен → владелец → строк
с данными), нераспознанное поимённо, CSV — одним листом выбранной группы,
шаблон — из реестра, отчёт — по доменам и пропуски по видам, и счёт
«будет записано» сходится с тем, что записалось.
"""

# ruff: noqa: F811 — фикстуры фазы 65 импортированы по имени
from __future__ import annotations

from io import BytesIO

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook

from students import admission_import, import_registry
from students.tests.test_phase65 import (  # noqa: F401
    HEADER_19,
    admin,
    asem,
    boston,
    chicago,
    curator,
    klass,
    kymbat,
    stranger,
)
from students.tests.test_phase71 import book_of, login


@pytest.mark.django_db
def test_step_two_shows_exactly_the_registry(klass, asem):
    """Колонка → поле → домен → владелец → строк с данными — из реестра, не из своих списков."""
    sheets = admission_import.parse(book_of(klass))
    payload = admission_import.preview_payload(sheets)

    by_key = {row["key"]: row for row in payload["columns"]}
    # все колонки данных таблицы на месте, ФИО как служебная — нет
    assert set(by_key) == {spec.key for spec in import_registry.COLUMNS if spec.target != import_registry.MATCH}
    phone = by_key["phone"]
    assert phone == {
        "key": "phone",
        "title": "Номер телефона",
        "field_title": "Телефон ученика",
        "domain": "admission",
        "domain_title": "Поступление",
        "owner": "Асем",
        "kind": "телефон",
        "rows_with_data": 1,
    }
    assert by_key["gpa"]["owner"] == "Кымбат"
    assert by_key["ielts_2"]["field_title"] == "IELTS, результат 2"
    assert payload["unknown_columns"] == []
    assert payload["groups"] == ["CHICAGO"]


@pytest.mark.django_db
def test_unknown_columns_are_listed_by_name(klass):
    from students.tests.test_phase65 import book
    from students.tests.test_phase68_import import FULL, row19

    header = [*HEADER_19, "Любимый цвет", "Рост"]
    uploaded = book({"Chicago ": (header, [[*row19(klass[0].full_name, **FULL), "синий", "180"]])})
    payload = admission_import.preview_payload(admission_import.parse(uploaded))
    assert payload["unknown_columns"] == ["Любимый цвет", "Рост"]


@pytest.mark.django_db
def test_csv_is_one_sheet_of_the_chosen_group(klass, chicago, asem):
    """У CSV листов нет: файл — один лист группы с первого шага."""
    text = "ФИО,Номер телефона,Средний GPA\n" + f"{klass[0].full_name},8 707 389 63 73,4.4\n"
    uploaded = SimpleUploadedFile("chicago.csv", text.encode("utf-8"), content_type="text/csv")

    sheets = admission_import.parse(uploaded, group="CHICAGO")
    assert [sheet.group_code for sheet in sheets] == ["CHICAGO"]
    assert sheets[0].rows[0].values["phone"] == "+77073896373"

    with pytest.raises(admission_import.FileRejected, match="укажите группу"):
        admission_import.parse(SimpleUploadedFile("x.csv", text.encode("utf-8")))


@pytest.mark.django_db
def test_template_comes_from_the_registry(admin):
    answer = login(admin).get("/api/admission-imports/template/")
    assert answer.status_code == 200
    header = next(iter(load_workbook(BytesIO(answer.content)).active.iter_rows(values_only=True)))
    assert list(header) == ["№", *[spec.title for spec in import_registry.COLUMNS]]


@pytest.mark.django_db
def test_the_count_on_step_two_matches_what_was_written(klass, admin):
    """Счёт «будет записано» — поля выбранных доменов; отчёт считает те же поля."""
    sheets = admission_import.parse(book_of(klass))
    payload = admission_import.preview_payload(sheets)
    chosen = ["admission"]
    planned = [c for c in payload["columns"] if c["domain"] in chosen and c["rows_with_data"]]

    record = admission_import.apply(book_of(klass), actor=admin, domains=chosen)
    written = {
        line["text"].split(":")[0]: int(line["text"].rsplit("— ", 1)[1])
        for line in admission_import.report_rows(record)
        if line["kind"] == "домен"
    }
    # признак «Common App заведён» ставится сверх колонок — вычитаем его
    assert written == {"Поступление": len(planned) + 1}


@pytest.mark.django_db
def test_the_report_groups_skips_by_kind(klass, admin):
    from students.tests.test_phase65 import book
    from students.tests.test_phase68_import import FULL, row19

    header = [*HEADER_19, "Любимый цвет"]
    rows = [[*row19(klass[0].full_name, **FULL), "синий"], [*row19("Нет Такого", **FULL), "белый"]]
    record = admission_import.apply(book({"Chicago ": (header, rows)}), actor=admin, domains=["admission"])
    kinds = {row["kind"]: row["count"] for row in admission_import.record_payload(record)["skipped_by_kind"]}
    # вне «Поступления» — GPA, шесть баллов и три документа: десять колонок
    assert kinds == {"unknown": 1, "domain": 10, "row": 1}


@pytest.mark.django_db
def test_apply_names_the_first_student_for_the_card_link(klass, admin):
    payload = admission_import.record_payload(admission_import.apply(book_of(klass), actor=admin))
    assert payload["first_student"] == klass[0].pk


@pytest.mark.django_db
def test_owner_cannot_smuggle_a_foreign_domain_even_by_hand(klass, kymbat):
    """Кымбат подменяет запрос: домен «Поступление» — 403, в базе пусто."""
    answer = login(kymbat).post(
        "/api/admission-imports/apply/", {"file": book_of(klass), "domains": '["admission"]'}, format="multipart"
    )
    assert answer.status_code == 403
    klass[0].admission.refresh_from_db()
    assert klass[0].admission.student_phone == ""
