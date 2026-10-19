"""Отметка куратора сразу при нажатии и аналитика по отметкам (решение владельца, 19.10.2026)."""

from __future__ import annotations

import datetime as dt
from io import BytesIO

import pytest
from django.utils import timezone
from rest_framework.test import APIClient

from accounts.curators import assign
from accounts.models import Role
from core.models import AuditLog
from students.models import AttendanceDay, BehaviorProfile, Student, StudyGroup

pytestmark = pytest.mark.django_db


@pytest.fixture
def groups(db):
    return {
        "boston": StudyGroup.objects.create(code="BOSTON", parallel=10),
        "chicago": StudyGroup.objects.create(code="CHICAGO", parallel=10),
    }


@pytest.fixture
def pupils(groups):
    rows = {
        "a": Student.objects.create(
            last_name="Абаева", first_name="Айгерим", group=groups["boston"], graduation_year=2028
        ),
        "b": Student.objects.create(
            last_name="Бекова", first_name="Дана", group=groups["boston"], graduation_year=2028
        ),
        "c": Student.objects.create(
            last_name="Серик", first_name="Ерлан", group=groups["chicago"], graduation_year=2028
        ),
    }
    for student in rows.values():
        BehaviorProfile.objects.create(student=student)
    return rows


@pytest.fixture
def admin(make_user):
    return make_user(Role.ADMIN, email="admin@example.kz")


@pytest.fixture
def curator(make_user, groups, admin):
    user = make_user(Role.CURATOR, email="curator@example.kz", full_name="Нурланова Асель")
    for code in ("boston", "chicago"):
        assign(group=groups[code], curator=user, since=dt.date.today() - dt.timedelta(days=10), actor=admin)
    return user


@pytest.fixture
def saltanat(make_user):
    return make_user("director_behavior", email="saltanat@example.kz")


def client_of(user) -> APIClient:
    api = APIClient()
    api.force_login(user)
    return api


def today() -> dt.date:
    return timezone.localdate()


def test_opening_the_sheet_writes_present_for_everyone_and_a_tap_saves_at_once(curator, pupils):
    api = client_of(curator)
    sheet = api.get("/api/attendance/?group=BOSTON").json()
    assert sheet["saved"] is True and sheet["noted"]["by"] == "Нурланова Асель"
    assert all(row["mark"] == "present" for row in sheet["rows"]) and sheet["absent"] == 0
    assert AttendanceDay.objects.filter(date=today(), mark="present").count() == 2
    # открытие листа в журнал не пишет: все в школе — не событие
    assert not AuditLog.objects.filter(field_name="attendance_day").exists()

    tapped = api.post(
        "/api/attendance/mark/", {"group": "BOSTON", "student": pupils["a"].pk, "mark": "absent"}, format="json"
    )
    assert tapped.status_code == 200, tapped.content
    assert (
        tapped.json()["row"]["mark"] == "absent" and tapped.json()["absent"] == 1 and tapped.json()["late_count"] == 0
    )
    late = api.post(
        "/api/attendance/mark/",
        {"group": "BOSTON", "student": pupils["b"].pk, "mark": "late", "reason": "пробки"},
        format="json",
    ).json()
    assert late["row"] == {
        "student": pupils["b"].pk,
        "full_name": "Бекова Дана",
        "mark": "late",
        "reason": "пробки",
        "marked": True,
    }
    assert late["late_count"] == 1
    # по кругу: «опоздал» → «в школе» — комментарий у «в школе» не хранится
    back = api.post(
        "/api/attendance/mark/",
        {"group": "BOSTON", "student": pupils["b"].pk, "mark": "present", "reason": "пробки"},
        format="json",
    )
    assert back.json()["row"]["mark"] == "present" and back.json()["row"]["reason"] == ""
    assert AttendanceDay.objects.filter(date=today()).count() == 2

    texts = [e.new_value for e in AuditLog.objects.filter(field_name="attendance_day").order_by("pk")]
    stamp = f"{today():%d.%m.%Y}"
    assert texts == [f"{stamp}: не был (исправлено)", f"{stamp}: опоздал (исправлено)", f"{stamp}: был (исправлено)"]


def test_mark_refusals(curator, saltanat, admin, pupils):
    api = client_of(curator)
    assert (
        api.post(
            "/api/attendance/mark/", {"group": "BOSTON", "student": pupils["c"].pk, "mark": "absent"}, format="json"
        ).status_code
        == 404
    )
    bad = api.post(
        "/api/attendance/mark/", {"group": "BOSTON", "student": pupils["a"].pk, "mark": "sick"}, format="json"
    )
    assert bad.status_code == 400
    yesterday = (today() - dt.timedelta(days=1)).isoformat()
    assert (
        api.post(
            "/api/attendance/mark/",
            {"group": "BOSTON", "student": pupils["a"].pk, "mark": "absent", "date": yesterday},
            format="json",
        ).status_code
        == 400
    )
    assert (
        client_of(saltanat)
        .post("/api/attendance/mark/", {"group": "BOSTON", "student": pupils["a"].pk, "mark": "absent"}, format="json")
        .status_code
        == 403
    )
    assert (
        client_of(admin)
        .post("/api/attendance/mark/", {"group": "BOSTON", "student": pupils["a"].pk, "mark": "absent"}, format="json")
        .status_code
        == 403
    )
    assert (
        client_of(pupils["a"].user).get("/api/attendance/stats/").status_code == 403
        if getattr(pupils["a"], "user", None)
        else True
    )
    assert not AttendanceDay.objects.exists()


def seed_week(pupils):
    """Пять дней BOSTON: Айгерим два дня не была, Дана один раз опоздала; CHICAGO — один день."""
    base = today()
    for offset in range(5):
        day = base - dt.timedelta(days=offset)
        for student in (pupils["a"], pupils["b"]):
            AttendanceDay.objects.create(student=student, date=day, mark="present")
    AttendanceDay.objects.filter(student=pupils["a"], date__in=[base, base - dt.timedelta(days=1)]).update(
        mark="absent", reason="болела"
    )
    AttendanceDay.objects.filter(student=pupils["b"], date=base - dt.timedelta(days=2)).update(mark="late")
    AttendanceDay.objects.create(student=pupils["c"], date=base, mark="late", reason="проспал")


def test_stats_by_students_and_totals_with_dates_and_comments(curator, saltanat, pupils):
    seed_week(pupils)
    since, until = (today() - dt.timedelta(days=6)).isoformat(), today().isoformat()
    api = client_of(curator)
    report = api.get(f"/api/attendance/stats/?group=BOSTON&period=custom&from={since}&to={until}").json()
    assert report["period"]["code"] == "custom" and report["group"] == "BOSTON"
    assert {row["code"] for row in report["periods"]} >= {"week", "custom"}
    rows = {row["full_name"]: row for row in report["rows"]}
    aigerim, dana = rows["Абаева Айгерим"], rows["Бекова Дана"]
    assert (aigerim["days"], aigerim["present"], aigerim["absent"], aigerim["late"], aigerim["percent"]) == (
        5,
        3,
        2,
        0,
        60,
    )
    assert (dana["days"], dana["present"], dana["absent"], dana["late"], dana["percent"]) == (5, 4, 0, 1, 100)
    assert [(d["mark"], d["reason"]) for d in aigerim["dates"]] == [("absent", "болела"), ("absent", "болела")]
    assert report["totals"] == {"days": 5, "absent": 2, "late": 1, "percent": 80}

    everything = client_of(saltanat).get(f"/api/attendance/stats/?period=custom&from={since}&to={until}").json()
    assert everything["totals"]["days"] == 5 and everything["totals"]["late"] == 2
    assert {row["group"] for row in everything["rows"]} == {"BOSTON", "CHICAGO"}
    assert api.get("/api/attendance/stats/?group=NOPE").status_code == 404
    assert api.get("/api/attendance/stats/?period=week").status_code == 200

    export = api.get(f"/api/attendance/stats/export/?group=BOSTON&period=custom&from={since}&to={until}")
    assert export.status_code == 200
    from openpyxl import load_workbook

    sheet = list(load_workbook(BytesIO(export.content)).active.iter_rows(values_only=True))
    assert sheet[0][:7] == ("Ученик", "Группа", "Дней отмечено", "Был", "Не был", "Опоздал", "Процент")
    assert sheet[1][0] == "Абаева Айгерим" and "болела" in sheet[1][7]


def test_absences_export_lists_late_with_comments(curator, pupils):
    seed_week(pupils)
    export = client_of(curator).get("/api/attendance/absences/export/?group=CHICAGO")
    from openpyxl import load_workbook

    sheet = list(load_workbook(BytesIO(export.content)).active.iter_rows(values_only=True))
    assert sheet[0][5:7] == ("Кто опоздал", "Опозданий")
    assert sheet[1][5] == "Серик Ерлан (проспал)" and sheet[1][6] == 1
