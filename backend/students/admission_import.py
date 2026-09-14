"""Импорт таблицы поступления по реестру соответствий (фаза 65, движок — 71).

Асем ведёт учеников в Excel: по листу на группу, в строке — контакты,
пароли, ссылки на папку, паспорт, табель и рекомендацию, срок паспорта,
GPA и до трёх результатов IELTS и SAT. Данные ей приносят кураторы.
Решение владельца: таблица заливается один раз, дальше всё ведётся
в карточке. Повторный запуск не дублирует, а обновляет.

С фазы 71 разбор ничего не знает о колонках сам: какая колонка какое
поле какого домена заполняет и как её разбирать — записано в реестре
(`students.import_registry`), и здесь только движок:

* **лист — это группа** по своему имени (`Chicago ` с пробелом → CHICAGO);
  листа без группы мы не трогаем вовсе: угадывать, чей это класс, нельзя;
* **колонки ищутся по заголовкам через реестр**, не по позиции; колонка,
  которой реестр не знает, — не ошибка, а строка отчёта;
* **домены выбирает человек**: что не выбрано — не пишется, даже если
  колонка в файле есть, и это видно в отчёте;
* **пустая ячейка ничего не стирает**: пустота в таблице значит
  «не знаю», а не «нет»;
* **из текста попытка не создаётся** — в ячейках баллов встречаются
  заметки Асем себе, а не результаты;
* **почта из таблицы — личная почта ученика**, текст в карточке: с логином
  она не сверяется и предупреждений о несовпадении не даёт.

Пароли идут прямо в шифрованное хранилище (`students.credentials`)
и на шаге проверки не показываются — только «есть / нет».
"""

from __future__ import annotations

import datetime as dt
from dataclasses import dataclass, field

from django.db import transaction
from django.utils import timezone

from students import import_registry as registry
from students.import_registry import parse_date as parse_expiry  # noqa: F401
from students.import_registry import (  # noqa: F401 — разбор ячеек живёт в реестре, имена оставлены
    parse_email,
    parse_gpa,
    parse_link,
    parse_phone,
    parse_score,
    read_columns,
)


class FileRejected(ValueError):
    """Файл нельзя разбирать вовсе — дело не в отдельной строке."""


_text = registry._text


# --- Строки и листы -------------------------------------------------------


@dataclass
class Row:
    """Одна строка листа после разбора: что нашли и что не сошлось."""

    index: int
    raw_name: str
    student: int | None = None
    student_name: str = ""
    candidates: list[dict] = field(default_factory=list)
    #: значения по ключу колонки реестра — только разобранные, без пустых
    values: dict = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)
    error: str = ""
    skip: bool = False

    @property
    def scores(self) -> list[dict]:
        return [
            {"exam": spec.exam, "slot": spec.slot, "value": self.values[spec.key]}
            for spec in registry.COLUMNS
            if spec.target == registry.ATTEMPT and spec.key in self.values
        ]

    @property
    def links(self) -> list[dict]:
        return [
            {"doc_type": spec.doc_type, "url": self.values[spec.key]}
            for spec in registry.COLUMNS
            if spec.target == registry.DOCUMENT and spec.key in self.values
        ]

    @property
    def passwords(self) -> dict[str, bool]:
        return {
            spec.credential: spec.key in self.values for spec in registry.COLUMNS if spec.target == registry.CREDENTIAL
        }

    def as_dict(self) -> dict:
        gpa = self.values.get("gpa")
        return {
            "index": self.index,
            "raw_name": self.raw_name,
            "student": self.student,
            "student_name": self.student_name,
            "candidates": self.candidates,
            "phone": self.values.get("phone", ""),
            "email": self.values.get("email", ""),
            "common_app_email": self.values.get("common_app_email", ""),
            "drive_folder_url": self.values.get("drive_folder", ""),
            "gpa": float(gpa) if gpa is not None else None,
            "passport_expires": self.values.get("passport_expiry"),
            "scores": [{**s, "value": float(s["value"])} for s in self.scores],
            "links": self.links,
            "has_email_password": self.passwords.get("email", False),
            "has_common_app_password": self.passwords.get("common_app", False),
            "warnings": self.warnings,
            "error": self.error,
            "skip": self.skip,
        }


@dataclass
class Sheet:
    """Один лист файла: чья это группа и что в строках."""

    name: str
    group_code: str = ""
    group_id: int | None = None
    error: str = ""
    rows: list[Row] = field(default_factory=list)
    #: колонки реестра, найденные в заголовке листа
    columns: list[str] = field(default_factory=list)
    #: заголовки, которых реестр не знает
    unknown: list[str] = field(default_factory=list)

    def as_dict(self) -> dict:
        return {
            "name": self.name,
            "group_code": self.group_code,
            "group": self.group_id,
            "error": self.error,
            "columns": self.columns,
            "unknown_columns": self.unknown,
            "rows": [row.as_dict() for row in self.rows],
            "ready": sum(1 for row in self.rows if not row.error and not row.skip),
            "skipped": sum(1 for row in self.rows if row.error or row.skip),
        }


@dataclass
class Fix:
    """Правка человека по строке: кому отнести или пропустить."""

    student: int | None = None
    skip: bool = False


def read_sheets(uploaded, *, group: str = "") -> list[tuple[str, list[str], list[list]]]:
    """Прочитать книгу целиком: имя листа, заголовок и строки.

    Читаем сами, а не общим `read_table`: тому нужен один лист, а здесь
    лист — это группа, и все они нужны разом.
    """
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    name = (getattr(uploaded, "name", "") or "").lower()
    if name.endswith(".csv"):
        # у CSV листов нет: файл считается одним листом группы, которую
        # человек выбрал на первом шаге (фаза 72)
        if not group:
            raise FileRejected("Для CSV укажите группу: у файла нет листов, а лист — это группа")
        return [(group, *_read_csv(uploaded))]
    if not name.endswith((".xlsx", ".xlsm")):
        raise FileRejected("Файл читается из книги Excel (.xlsx) или CSV: лист — это группа")
    uploaded.seek(0)
    try:
        book = load_workbook(uploaded, read_only=True, data_only=True)
    except (InvalidFileException, KeyError, ValueError) as error:
        raise FileRejected(f"Файл не открылся как книга Excel: {error}") from error

    out = []
    for sheet_name in book.sheetnames:
        rows = [list(row) for row in book[sheet_name].iter_rows(values_only=True)]
        if not rows:
            out.append((sheet_name, [], []))
            continue
        header = [_text(cell) for cell in rows[0]]
        # пустые строки — хвост таблицы и разделитель под заголовком:
        # в одних листах он есть, в других данные идут сразу
        body = [row for row in rows[1:] if any(_text(cell) for cell in row)]
        out.append((sheet_name, header, body))
    book.close()
    return out


def _read_csv(uploaded) -> tuple[list[str], list[list]]:
    """CSV: заголовок и строки — тем же чтением, что и остальные файлы проекта."""
    import csv
    import io

    uploaded.seek(0)
    raw = uploaded.read()
    text = raw.decode("utf-8-sig") if isinstance(raw, bytes) else str(raw)
    sample = text[:2048]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows = [list(row) for row in csv.reader(io.StringIO(text), dialect)]
    if not rows:
        return [], []
    header = [_text(cell) for cell in rows[0]]
    body = [row for row in rows[1:] if any(_text(cell) for cell in row)]
    return header, body


def _group_of(sheet_name: str):
    """Группа по имени листа: `Chicago ` → CHICAGO."""
    from students.models import StudyGroup

    code = _text(sheet_name).upper()
    return StudyGroup.objects.filter(code__iexact=code).first()


def parse(uploaded, *, fixes: dict[str, Fix] | None = None, group: str = "") -> list[Sheet]:
    """Разобрать книгу целиком, ничего не записывая.

    Разбираются все колонки, какие нашлись: выбор доменов — дело
    применения и отчёта, разбор о нём не знает.
    """
    from students.models import Student
    from suggestions.name_matching import find

    fixes = fixes or {}
    sheets_raw = read_sheets(uploaded, group=group)
    if not sheets_raw:
        raise FileRejected("В книге нет ни одного листа")

    out: list[Sheet] = []
    for sheet_name, header, body in sheets_raw:
        sheet = Sheet(name=sheet_name, group_code=_text(sheet_name).upper())
        group = _group_of(sheet_name)
        if group is None:
            sheet.error = f"Группы «{sheet.group_code}» нет в системе — лист пропущен целиком"
            out.append(sheet)
            continue
        sheet.group_id = group.pk
        columns = registry.read_columns(header)
        sheet.columns = [spec.key for spec in registry.COLUMNS if spec.key in columns]
        sheet.unknown = registry.unknown_columns(header, columns)
        if "name" not in columns:
            sheet.error = "На листе не нашлась колонка «ФИО» — лист пропущен целиком"
            out.append(sheet)
            continue

        students = list(Student.objects.filter(group=group, is_active=True).select_related("group"))
        seen: dict[int, int] = {}
        for number, raw in enumerate(body, start=1):

            def cell(key: str, line: list = raw, columns: dict = columns):
                index = columns.get(key, -1)
                return line[index] if 0 <= index < len(line) else None

            raw_name = _text(cell("name"))
            # строка без ФИО — не ошибка: это хвост листа или его разметка
            if not raw_name:
                continue
            row = Row(index=number, raw_name=raw_name)
            fix = fixes.get(f"{sheet_name}:{number}")
            if fix and fix.skip:
                row.skip = True
            _resolve_student(row, students=students, fix=fix, finder=find)
            _resolve_values(row, cell=cell, keys=sheet.columns)
            if row.student is not None:
                if row.student in seen:
                    row.error = f"этот ученик уже был в строке №{seen[row.student]}"
                else:
                    seen[row.student] = number
            sheet.rows.append(row)
        out.append(sheet)
    return out


def _resolve_student(row: Row, *, students: list, fix: Fix | None, finder) -> None:
    """Найти ученика строки: правкой человека или нечётким сравнением ФИО."""
    if fix and fix.student:
        chosen = next((s for s in students if s.pk == fix.student), None)
        if chosen is not None:
            row.student, row.student_name = chosen.pk, chosen.full_name
            return
    outcome = finder(row.raw_name, students=students)
    row.candidates = [c.as_dict() for c in outcome.candidates]
    if outcome.is_confident and outcome.best:
        row.student, row.student_name = outcome.best.student_id, outcome.best.full_name
        return
    row.error = "похожих учеников несколько — выберите" if outcome.is_ambiguous else "ученик не найден в этой группе"


def _resolve_values(row: Row, *, cell, keys: list[str]) -> None:
    """Разобрать ячейки строки по реестру: каждую — своим типом."""
    for key in keys:
        spec = registry.spec_of(key)
        if spec.target == registry.MATCH:
            continue
        value, warning, fatal = registry.parse_cell(spec, cell(key))
        if fatal:
            row.error = warning
            continue
        if warning:
            row.warnings.append(warning)
        if value is not None:
            row.values[key] = value


def found_domains(sheets: list[Sheet]) -> list[str]:
    """Домены, для которых в файле нашлись колонки, — в порядке реестра."""
    keys = {key for sheet in sheets for key in sheet.columns}
    present = registry.domains_of_columns(keys)
    order: list[str] = []
    for spec in registry.COLUMNS:
        if spec.domain in present and spec.domain not in order:
            order.append(spec.domain)
    return order


def columns_payload(sheets: list[Sheet]) -> list[dict]:
    """Шаг «Что заполняем» (фаза 72): колонка → поле → домен → владелец → строк с данными.

    Всё из реестра: экран ничего не знает о колонках сам. Число строк —
    сколько строк файла несут значение в этой колонке; строки с ошибкой
    и пропущенные считаются тоже — человек смотрит на файл, а не на исход.
    """
    from core.domains import DOMAINS

    keys: list[str] = []
    for sheet in sheets:
        for key in sheet.columns:
            if key not in keys:
                keys.append(key)
    out = []
    for spec in registry.COLUMNS:
        if spec.key not in keys or spec.target == registry.MATCH:
            continue
        domain = DOMAINS.get(spec.domain)
        out.append(
            {
                "key": spec.key,
                "title": spec.title,
                "field_title": _field_title(spec),
                "domain": spec.domain,
                "domain_title": domain.title if domain else "",
                "owner": domain.owner_name if domain else "",
                "kind": registry.KIND_TITLES[spec.kind],
                "rows_with_data": sum(1 for sheet in sheets for row in sheet.rows if spec.key in row.values),
            }
        )
    return out


def _field_title(spec) -> str:
    """Человеческое имя поля — из реестра доменов, куда колонка ложится."""
    from core.domains import spec_of_field
    from students.models import CredentialKind, DocumentType

    if spec.target == registry.PROFILE:
        found = spec_of_field("students.AdmissionProfile", spec.field)
        return found.title if found else spec.field
    if spec.target == registry.EXAM_PROFILE:
        found = spec_of_field("students.ExamProfile", spec.field)
        return found.title if found else spec.field
    if spec.target == registry.ATTEMPT:
        return f"{spec.exam}, результат {spec.slot}"
    if spec.target == registry.DOCUMENT:
        return f"Документ «{DocumentType(spec.doc_type).label}»"
    if spec.target == registry.CREDENTIAL:
        return CredentialKind(spec.credential).label
    return spec.title


def unknown_payload(sheets: list[Sheet]) -> list[str]:
    """Нераспознанные заголовки всего файла — поимённо, без повторов."""
    seen: list[str] = []
    for sheet in sheets:
        for title in sheet.unknown:
            if title not in seen:
                seen.append(title)
    return seen


def preview_payload(sheets: list[Sheet]) -> dict:
    """Шаг «Проверка»: листы со строками, общие числа и домены файла. Паролей здесь нет."""
    rows = [row for sheet in sheets for row in sheet.rows]
    return {
        "sheets": [sheet.as_dict() for sheet in sheets],
        # шаг «Что заполняем» (фаза 72): колонки, нераспознанное, группы
        "columns": columns_payload(sheets),
        "unknown_columns": unknown_payload(sheets),
        "groups": [sheet.group_code for sheet in sheets if not sheet.error],
        # по умолчанию выбраны все домены, для которых нашлись колонки
        "domains": found_domains(sheets),
        "counts": {
            "sheets": len(sheets),
            "sheets_skipped": sum(1 for sheet in sheets if sheet.error),
            "rows": len(rows),
            "ready": sum(1 for row in rows if not row.error and not row.skip),
            "errors": sum(1 for row in rows if row.error and not row.skip),
            "warnings": sum(1 for row in rows if row.warnings),
            "skipped": sum(1 for row in rows if row.skip),
            "attempts": sum(len(row.scores) for row in rows if not row.error and not row.skip),
            "links": sum(len(row.links) for row in rows if not row.error and not row.skip),
            "passwords": sum(
                sum(1 for has in row.passwords.values() if has) for row in rows if not row.error and not row.skip
            ),
        },
    }


# --- Применение -----------------------------------------------------------


@transaction.atomic
def apply(uploaded, *, actor, fixes: dict[str, Fix] | None = None, domains: list[str] | None = None, group: str = ""):
    """Записать разобранное одной транзакцией и вернуть отчёт-запись.

    Половина таблицы записанной хуже, чем отказ: строки с ошибками
    здесь уже исправлены или пропущены человеком, а сбой на середине
    откатывает всё.

    `domains` — что заполнять; `None` значит все домены, для которых
    в файле нашлись колонки. Колонка невыбранного домена не пишется
    и попадает в отчёт: «пропущена: домен не выбран».
    """
    from students.models import AdmissionImport

    sheets = parse(uploaded, fixes=fixes, group=group)
    chosen = list(domains) if domains is not None else found_domains(sheets)
    first_student: int | None = None
    today = timezone.localdate()
    report: list[str] = []
    students_updated = 0
    attempts = links = passwords = skipped = 0
    written: dict[str, int] = {code: 0 for code in chosen}

    for sheet in sheets:
        if sheet.error:
            report.append(f"{sheet.name}\t—\t—\tлист\t{sheet.error}")
            continue
        for title in sheet.unknown:
            report.append(f"{sheet.name}\t—\t—\tколонка\tколонка «{title}» не распознана, пропущена")
        for key in sheet.columns:
            spec = registry.spec_of(key)
            if spec.domain and spec.domain not in chosen:
                report.append(
                    f"{sheet.name}\t—\t—\tколонка\t"
                    f"колонка «{spec.title}» пропущена: домен «{spec.domain_title}» не выбран"
                )
        for row in sheet.rows:
            if row.skip or row.error:
                skipped += 1
                reason = "пропущена человеком" if row.skip else row.error
                report.append(f"{sheet.name}\t{row.index}\t{row.raw_name}\tпропуск\t{reason}")
                continue
            outcome = _apply_row(row, actor=actor, today=today, domains=chosen)
            if first_student is None:
                first_student = row.student
            students_updated += 1 if outcome["changed"] else 0
            attempts += outcome["attempts"]
            links += outcome["links"]
            passwords += outcome["passwords"]
            for code, count in outcome["by_domain"].items():
                written[code] = written.get(code, 0) + count
            for warning in row.warnings:
                report.append(f"{sheet.name}\t{row.index}\t{row.student_name}\tвнимание\t{warning}")

    from core.domains import DOMAINS

    for code in chosen:
        title = DOMAINS[code].title if code in DOMAINS else code
        report.append(f"—\t—\t—\tдомен\t{title}: записано значений — {written.get(code, 0)}")

    record = AdmissionImport.objects.create(
        uploaded_by=actor if getattr(actor, "pk", None) else None,
        file_name=getattr(uploaded, "name", "") or "",
        domains=",".join(chosen),
        sheets=len(sheets),
        students_updated=students_updated,
        attempts_created=attempts,
        documents_created=links,
        credentials_saved=passwords,
        rows_skipped=skipped,
        report="\n".join(report),
    )
    # первый ученик файла — для ссылки «открыть карточку» на последнем шаге;
    # в записи не хранится: это удобство момента, а не факт загрузки
    record.first_student = first_student
    return record


def _apply_row(row: Row, *, actor, today: dt.date, domains: list[str]) -> dict:
    """Записать одну строку по реестру: профили, пароли, документы, попытки."""
    from core.audit import apply_changes
    from core.domains import Source
    from students.models import AdmissionProfile, ExamProfile, Student

    student = Student.objects.get(pk=row.student)
    changed = False
    by_domain: dict[str, int] = {}

    def allowed(spec) -> bool:
        return bool(spec.domain) and spec.domain in domains

    # профили: значения ложатся полями, пустая ячейка ничего не стирает
    profile_changes: dict[str, dict] = {}
    for key, value in row.values.items():
        spec = registry.spec_of(key)
        if spec.target in (registry.PROFILE, registry.EXAM_PROFILE) and allowed(spec):
            profile_changes.setdefault(spec.target, {})[spec.field] = value

    admission_changes = profile_changes.get(registry.PROFILE, {})
    if registry.PROFILE in profile_changes or _common_app_seen(row, domains):
        profile, _ = AdmissionProfile.objects.get_or_create(student=student)
        # пароль или почта Common App в таблице значат, что аккаунт заведён:
        # признак читают готовность и дашборд Асем, руками его никто не ставит
        if _common_app_seen(row, domains) and not profile.has_common_app:
            admission_changes["has_common_app"] = True
        if admission_changes:
            done = apply_changes(profile, admission_changes, actor=actor, source=Source.IMPORT)
            changed |= bool(done)
            by_domain["admission"] = by_domain.get("admission", 0) + len(admission_changes)

    exam_changes = profile_changes.get(registry.EXAM_PROFILE, {})
    if exam_changes:
        exam_profile, _ = ExamProfile.objects.get_or_create(student=student)
        changed |= bool(apply_changes(exam_profile, exam_changes, actor=actor, source=Source.IMPORT))
        by_domain["exam"] = by_domain.get("exam", 0) + len(exam_changes)

    passwords = _apply_passwords(row, student=student, actor=actor, domains=domains)
    links = _apply_links(row, student=student, actor=actor, domains=domains)
    attempts = _apply_scores(row, student=student, actor=actor, today=today, domains=domains)
    if passwords:
        by_domain["admission"] = by_domain.get("admission", 0) + passwords
    if links:
        by_domain["documents"] = by_domain.get("documents", 0) + links
    if attempts:
        by_domain["exam"] = by_domain.get("exam", 0) + attempts
    changed = changed or bool(passwords or links or attempts)
    return {"changed": changed, "attempts": attempts, "links": links, "passwords": passwords, "by_domain": by_domain}


def _common_app_seen(row: Row, domains: list[str]) -> bool:
    """Есть ли в строке признак заведённого Common App — и выбран ли его домен."""
    if "admission" not in domains:
        return False
    return bool(row.values.get("common_app_email") or row.values.get("common_app_password"))


def _apply_passwords(row: Row, *, student, actor, domains: list[str]) -> int:
    """Пароли — сразу в шифрованное хранилище, мимо любых ответов API."""
    from students import credentials

    saved = 0
    for spec in registry.COLUMNS:
        if spec.target != registry.CREDENTIAL or spec.domain not in domains:
            continue
        plaintext = row.values.get(spec.key, "")
        if plaintext and credentials.set_credential(student, spec.credential, plaintext, actor=actor):
            saved += 1
    return saved


def _apply_links(row: Row, *, student, actor, domains: list[str]) -> int:
    """Ссылки на паспорт, табель и рекомендацию — документами-ссылками.

    Тот же документ той же ссылкой второй раз не заводится: повторный
    запуск таблицы обновляет, а не плодит строки в чек-листе. Срок
    паспорта документ берёт из поля профиля (фаза 71): поле пишется
    и без ссылки, а появится ссылка — документ подхватит срок оттуда.
    """
    from core.audit import apply_changes
    from core.domains import Source
    from students import documents as documents_service
    from students.models import DocumentStatus, DocumentType, StudentDocument

    if "documents" not in domains:
        return 0
    made = 0
    for link in row.links:
        doc_type = link["doc_type"]
        expires = _passport_expiry(row, student=student) if doc_type == DocumentType.PASSPORT else None
        existing = (
            StudentDocument.objects.filter(student=student, doc_type=doc_type)
            .exclude(external_url="")
            .order_by("created_at", "id")
            .last()
        )
        wanted = {"external_url": link["url"]}
        if expires is not None:
            wanted["expires_at"] = expires
        if existing is not None:
            apply_changes(existing, wanted, actor=actor, source=Source.IMPORT)
            continue
        document = StudentDocument.objects.create(
            student=student,
            doc_type=doc_type,
            title=f"{DocumentType(doc_type).label}: ссылка из таблицы поступления",
            external_url=link["url"],
            expires_at=expires,
            uploaded_by=actor if getattr(actor, "pk", None) else None,
            status=DocumentStatus.PENDING,
        )
        # проверку документ-ссылка проходит ту же, что и файл (фаза 62)
        documents_service.submit(document, author=actor)
        made += 1
    return made


def _passport_expiry(row: Row, *, student) -> dt.date | None:
    """Срок для документа: из строки, если он там есть, иначе из профиля."""
    if row.values.get("passport_expiry") is not None:
        return row.values["passport_expiry"]
    admission = getattr(student, "admission", None)
    return getattr(admission, "passport_expires_at", None)


def _apply_scores(row: Row, *, student, actor, today: dt.date, domains: list[str]) -> int:
    """Баллы — официальными попытками с источником «импорт Асем».

    Даты в таблице нет: попытка получает дату загрузки и флаг «дата
    не указана», а карточка пишет «дата уточняется». Настоящую дату
    вносит ученик предложением — она флаг и снимет.

    Повторная загрузка не плодит попытки: место результата в таблице
    (IELTS-1, IELTS-2, …) — это и есть его номер, и та же клетка правит
    ту же попытку.
    """
    from core.audit import apply_changes
    from core.domains import Source
    from students.models import AttemptFormat, AttemptSource, ExamAttempt

    if "exam" not in domains:
        return 0
    made = 0
    by_exam: dict[str, list] = {}
    for exam in {spec.exam for spec in registry.COLUMNS if spec.target == registry.ATTEMPT}:
        by_exam[exam] = list(
            ExamAttempt.objects.filter(student=student, exam_type=exam, source=AttemptSource.ADMISSION_IMPORT).order_by(
                "created_at", "id"
            )
        )
    for score in sorted(row.scores, key=lambda s: (s["exam"], s["slot"])):
        exam, slot, value = score["exam"], score["slot"], score["value"]
        existing = by_exam[exam]
        if slot <= len(existing):
            apply_changes(existing[slot - 1], {"total_score": value}, actor=actor, source=Source.IMPORT)
            continue
        attempt = ExamAttempt.objects.create(
            student=student,
            exam_type=exam,
            attempt_format=AttemptFormat.OFFICIAL,
            source=AttemptSource.ADMISSION_IMPORT,
            date=today,
            date_unknown=True,
            total_score=value,
        )
        existing.append(attempt)
        made += 1
    return made


def report_rows(record) -> list[dict]:
    """Отчёт строками для экрана и выгрузки."""
    out = []
    for line in (record.report or "").splitlines():
        parts = line.split("\t")
        while len(parts) < 5:
            parts.append("")
        out.append({"sheet": parts[0], "row": parts[1], "student": parts[2], "kind": parts[3], "text": parts[4]})
    return out


def record_payload(record) -> dict:
    """Отчёт о загрузке — то, что показывает последний шаг мастера."""
    return {
        "id": record.pk,
        "created_at": record.created_at,
        "file_name": record.file_name,
        "uploaded_by": ((record.uploaded_by.full_name or record.uploaded_by.email) if record.uploaded_by_id else ""),
        "domains": [code for code in (record.domains or "").split(",") if code],
        "sheets": record.sheets,
        "students_updated": record.students_updated,
        "attempts_created": record.attempts_created,
        "documents_created": record.documents_created,
        "credentials_saved": record.credentials_saved,
        "rows_skipped": record.rows_skipped,
        "first_student": getattr(record, "first_student", None),
        "rows": report_rows(record),
        # пропуски по видам, а не одной кучей (фаза 72)
        "skipped_by_kind": _skipped_by_kind(record),
    }


def _skipped_by_kind(record) -> list[dict]:
    """Что пропущено и почему — по видам: колонка не распознана, домен не выбран, строка."""
    kinds = {
        "unknown": ("Колонки не распознаны", lambda r: r["kind"] == "колонка" and "не распознана" in r["text"]),
        "domain": ("Колонки вне выбранных доменов", lambda r: r["kind"] == "колонка" and "домен" in r["text"]),
        "sheet": ("Листы без группы", lambda r: r["kind"] == "лист"),
        "row": ("Строки с ошибкой", lambda r: r["kind"] == "пропуск" and "человеком" not in r["text"]),
        "manual": ("Строки, пропущенные человеком", lambda r: r["kind"] == "пропуск" and "человеком" in r["text"]),
    }
    rows = report_rows(record)
    return [
        {"kind": code, "title": title, "count": sum(1 for r in rows if test(r))}
        for code, (title, test) in kinds.items()
        if any(test(r) for r in rows)
    ]


def template_workbook() -> bytes:
    """Шаблон файла — заголовки колонок из реестра, лист-пример группы (фаза 72)."""
    from io import BytesIO

    from openpyxl import Workbook

    book = Workbook()
    page = book.active
    page.title = "ГРУППА"
    page.append(["№", *[spec.title for spec in registry.COLUMNS]])
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()
