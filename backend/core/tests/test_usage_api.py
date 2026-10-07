"""Права девяти ролей, безопасный пакет, календарь Алматы и общая выборка XLSX."""

from __future__ import annotations

import uuid
from datetime import UTC, datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook
from rest_framework.test import APIClient

from accounts.models import Role
from core import usage
from core.models import UsageEvent

pytestmark = pytest.mark.django_db
PERIOD = "from=2026-10-07&to=2026-10-07"


def client_of(user, **kwargs):
    client = APIClient(**kwargs)
    client.force_login(user)
    return client


def event_for(actor, stamp="2026-10-07T07:00:00+00:00", **extra):
    return UsageEvent.objects.create(
        event_id=uuid.uuid4(),
        actor=actor,
        role=extra.pop("role", actor.role),
        action=extra.pop("action", "screen.open"),
        screen=extra.pop("screen", "/dashboard"),
        occurred_at=datetime.fromisoformat(stamp),
        source="client",
        **extra,
    )


@pytest.fixture
def deliver(monkeypatch):
    batches = []

    def publish(*, args, **kwargs):
        batches.append(args[0])
        usage.store_events(args[0])

    monkeypatch.setattr("core.tasks.record_usage_events.apply_async", publish)
    return batches


@pytest.mark.parametrize("role", Role.values)
def test_every_role_can_write_only_own_events_and_get_registry(
    role, make_user, deliver, django_capture_on_commit_callbacks
):
    actor = make_user(role)
    client = client_of(actor)
    registry = client.get("/api/usage/registry/")
    assert registry.status_code == 200
    assert "/dashboard" in registry.json()["client_screens"]
    event = {"id": str(uuid.uuid4()), "action": "screen.open", "screen": "/dashboard"}
    with django_capture_on_commit_callbacks(execute=True):
        response = client.post("/api/usage/", {"events": [event]}, format="json")
    assert response.status_code == 202 and response.json() == {"accepted": 1}
    row = UsageEvent.objects.get()
    assert row.actor_id == actor.pk and row.role == role
    assert row.source == "client"


@pytest.mark.parametrize("role", Role.values)
def test_only_analytics_readers_get_summary_raw_events_or_export(role, make_user):
    client = client_of(make_user(role))
    allowed = role in (Role.ADMIN, Role.DIRECTOR_BEHAVIOR)
    for path in ["/api/usage/", "/api/usage/events/", "/api/usage/export/", "/api/usage/export/?preview=1"]:
        response = client.get(path)
        assert response.status_code == 200 if allowed else response.status_code in (403, 404)


def test_anonymous_and_missing_csrf_are_refused(make_user):
    event = {"events": [{"id": str(uuid.uuid4()), "action": "screen.open", "screen": "/dashboard"}]}
    assert APIClient().post("/api/usage/", event, format="json").status_code == 403
    assert APIClient().get("/api/usage/registry/").status_code == 403
    client = client_of(make_user(), enforce_csrf_checks=True)
    assert client.post("/api/usage/", event, format="json").status_code == 403


@pytest.mark.parametrize(
    "extra",
    [
        {"actor": 999},
        {"actor_id": 999},
        {"role": "admin"},
        {"occurred_at": "2026-01-01"},
        {"payload": {"student_id": 12}},
        {"object_id": 12},
    ],
)
def test_event_identity_time_payload_and_objects_cannot_be_supplied(extra, make_user):
    client = client_of(make_user())
    event = {"id": str(uuid.uuid4()), "action": "screen.open", "screen": "/dashboard", **extra}
    assert client.post("/api/usage/", {"events": [event]}, format="json").status_code == 400
    assert not UsageEvent.objects.exists()


@pytest.mark.parametrize(
    "overrides",
    [
        {"action": "journal.grade.set"},
        {"action": "unknown"},
        {"id": "invalid"},
        {"screen": "/students/12"},
        {"screen": "/dashboard?student=12"},
        {"screen": "/users"},
    ],
)
def test_server_keys_raw_paths_unknown_values_and_other_role_screens_are_refused(overrides, make_user):
    client = client_of(make_user(Role.STUDENT))
    event = {"id": str(uuid.uuid4()), "action": "screen.open", "screen": "/dashboard", **overrides}
    assert client.post("/api/usage/", {"events": [event]}, format="json").status_code == 400


def test_batch_limit_and_unknown_envelope_field(make_user):
    client = client_of(make_user())
    event = {"id": str(uuid.uuid4()), "action": "screen.open", "screen": "/dashboard"}
    assert client.post("/api/usage/", {"events": [event] * 101}, format="json").status_code == 400
    assert client.post("/api/usage/", {"events": [], "actor": 1}, format="json").status_code == 400


def test_repeated_client_uuid_is_private_and_actor_scoped(make_user, deliver, django_capture_on_commit_callbacks):
    actors = [make_user(Role.STUDENT, "one@example.kz"), make_user(Role.STUDENT, "two@example.kz")]
    event = {"id": str(uuid.uuid4()), "action": "screen.open", "screen": "/dashboard"}
    responses = []
    for actor in actors:
        client = client_of(actor)
        for _ in range(2):
            with django_capture_on_commit_callbacks(execute=True):
                response = client.post("/api/usage/", {"events": [event]}, format="json")
            responses.append((response.status_code, response.json()))
    assert responses == [(202, {"accepted": 1})] * 4
    assert UsageEvent.objects.count() == 2


def test_junior_and_forced_password_gates_allow_ingestion_but_not_analytics(
    make_user, student, deliver, django_capture_on_commit_callbacks
):
    student.group.parallel = 8
    student.group.save(update_fields=["parallel"])
    actor = make_user(Role.STUDENT, must_change_password=True)
    student.user = actor
    student.save(update_fields=["user"])
    client = client_of(actor)
    registry = client.get("/api/usage/registry/")
    assert registry.status_code == 200
    assert "/homework" in registry.json()["client_screens"]
    assert "/selection" not in registry.json()["client_screens"]
    event = {"id": str(uuid.uuid4()), "action": "screen.open", "screen": "/set-password"}
    with django_capture_on_commit_callbacks(execute=True):
        assert client.post("/api/usage/", {"events": [event]}, format="json").status_code == 202
    assert UsageEvent.objects.get().actor_id == actor.pk
    assert client.get("/api/usage/").status_code == 403


def test_summary_uses_inclusive_almaty_dates_and_event_role_snapshot(make_user):
    actor = make_user(Role.CURATOR, full_name="Сотрудник")
    event_for(actor, "2026-10-06T18:59:59+00:00")
    first = event_for(actor, "2026-10-06T19:00:00+00:00")
    last = event_for(actor, "2026-10-07T18:59:59.999999+00:00")
    event_for(actor, "2026-10-07T19:00:00+00:00")
    actor.role = Role.DIRECTOR_EXAM
    actor.save(update_fields=["role"])
    client = client_of(make_user(Role.ADMIN))
    body = client.get(f"/api/usage/?{PERIOD}&by=role").json()
    assert body["cards"]["actions"] == 2 and body["cards"]["active_users"] == 1
    assert body["results"][0]["key"] == Role.CURATOR
    body = client.get(f"/api/usage/events/?{PERIOD}").json()
    assert {row["id"] for row in body["results"]} == {first.pk, last.pk}
    body = client.get(f"/api/usage/?{PERIOD}&by=day").json()
    assert [(row["key"], row["count"]) for row in body["results"]] == [("2026-10-07", 2)]


def test_never_logged_in_is_all_time_real_active_accounts_regardless_of_filters(make_user):
    admin = make_user(Role.ADMIN)
    client = client_of(admin)
    make_user(Role.STUDENT, "never@example.kz")
    make_user(Role.STUDENT, "past@example.kz", last_login=datetime(2020, 1, 1, tzinfo=UTC))
    make_user(Role.STUDENT, "inactive@example.kz", is_active=False)
    make_user(Role.TEACHER, "fictional@example.kz", is_fictional=True)
    make_user(Role.STUDENT, "ignore@probe.local")
    make_user(Role.STUDENT, "probe-login@example.kz", login="probe_123")
    for params in [PERIOD, "from=2020-01-01&to=2020-01-02&action=auth.login&role=admin"]:
        cards = client.get(f"/api/usage/?{params}").json()["cards"]
        assert cards["never_logged_in"] == 1
        assert cards["actions"] == 0 and cards["most_frequent"] is None


def test_summary_and_export_share_all_filters_and_do_not_expose_raw_uuid_or_email(make_user):
    actor = make_user(Role.TEACHER, "private@example.kz", full_name="=HYPERLINK(1)")
    wanted = event_for(actor, action="journal.grade.set", screen="/journals/:id")
    event_for(actor, action="homework.check", screen="/homework-review/:id")
    client = client_of(make_user(Role.ADMIN))
    query = f"{PERIOD}&by=user&user={actor.pk}&role=teacher&action=journal.grade.set&screen=/journals/:id"
    summary = client.get(f"/api/usage/?{query}").json()
    assert summary["cards"]["actions"] == 1 and summary["results"][0]["count"] == 1
    response = client.get(f"/api/usage/export/?{query}")
    assert response.status_code == 200
    book = load_workbook(BytesIO(response.content))
    assert book.sheetnames == ["Сводка", "События"]
    assert book["Сводка"].max_row == book["События"].max_row == 2
    assert book["Сводка"]["A2"].data_type == book["События"]["A2"].data_type == "s"
    assert book["События"]["A2"].value.startswith("'=")
    assert book["События"]["E2"].value == datetime(2026, 10, 7, 12)
    events = client.get(f"/api/usage/events/?{query}").json()
    assert events["results"][0]["id"] == wanted.pk
    assert str(wanted.event_id) not in str(events) and actor.email not in str(events)
    assert response["Cache-Control"] == "private, no-store"


def test_export_preview_and_raw_log_are_paginated(make_user):
    actor = make_user(Role.ADMIN)
    for _ in range(53):
        event_for(actor)
    client = client_of(actor)
    events = client.get(f"/api/usage/events/?{PERIOD}").json()
    assert events["count"] == 53 and len(events["results"]) == 50 and events["next"]
    second = client.get(f"/api/usage/events/?{PERIOD}&page=2").json()
    assert len(second["results"]) == 3
    preview = client.get(f"/api/usage/export/?{PERIOD}&preview=1").json()
    assert len(preview["sheets"]) == 2
    assert preview["sheets"][1]["total"] == 53


@pytest.mark.parametrize(
    "query", ["from=2026-10-08&to=2026-10-07", "from=wrong", "by=unknown", "role=unknown", "user=-1"]
)
def test_invalid_filters_refused_in_both_summary_and_export(query, make_user):
    client = client_of(make_user(Role.ADMIN))
    assert client.get(f"/api/usage/?{query}").status_code == 400
    assert client.get(f"/api/usage/export/?{query}").status_code == 400
