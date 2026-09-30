"""Выгрузка переводов в xlsx для вычитки и загрузка правок обратно.

Работают на копии словарей во временной папке: тест не трогает исходники.
"""

from __future__ import annotations

import shutil
from io import StringIO

import pytest
from django.core.management import call_command
from openpyxl import load_workbook

from core import translations


@pytest.fixture
def dictionaries(tmp_path):
    folder = tmp_path / "i18n"
    folder.mkdir()
    for lang in translations.LANGS:
        shutil.copy(translations.frontend_src() / "i18n" / f"{lang}.ts", folder / f"{lang}.ts")
    return folder


def export(folder, tmp_path):
    out = tmp_path / "translations.xlsx"
    call_command("i18n_export", out=str(out), frontend_dir=str(folder), stdout=StringIO())
    return out


def rows_of(path):
    sheet = load_workbook(path).active
    header, *rows = list(sheet.iter_rows(values_only=True))
    return header, {row[0]: row for row in rows}


def test_export_has_every_key_with_translations_and_places(dictionaries, tmp_path):
    header, rows = rows_of(export(dictionaries, tmp_path))
    assert header == ("Ключ", "Русский", "Қазақша", "English", "Где встречается", "Комментарий")
    kk = translations.Dictionary.read(dictionaries / "kk.ts").entries
    assert set(kk) <= set(rows), "в выгрузке не все ключи словаря"
    assert rows["Выйти"][1] == "Выйти" and rows["Выйти"][2] and rows["Выйти"][3]
    assert rows["Выйти"][4] != "не используется"
    # подстановки и формы числа объяснены прямо в строке
    plural = next(key for key in rows if "|" in key and key in translations.plural_keys())
    assert "формы числа" in rows[plural][5]


def test_import_applies_edits_and_refuses_broken_ones(dictionaries, tmp_path):
    path = export(dictionaries, tmp_path)
    book = load_workbook(path)
    sheet = book.active
    placeholder_key = next(
        key for key in translations.Dictionary.read(dictionaries / "kk.ts").entries if "{" in key and "|" not in key
    )
    for row in sheet.iter_rows(min_row=2):
        if row[0].value == "Выйти":
            row[2].value = "Жүйеден шығу"
        if row[0].value == placeholder_key:
            row[2].value = "подстановка потерялась"
    book.save(path)

    output = StringIO()
    call_command("i18n_import", str(path), frontend_dir=str(dictionaries), stdout=output)
    kk = translations.Dictionary.read(dictionaries / "kk.ts").entries
    assert kk["Выйти"] == "Жүйеден шығу"
    assert kk[placeholder_key] != "подстановка потерялась"
    assert "подстановки не совпадают" in output.getvalue()

    # записанный словарь читается тем же разбором, что и правило ESLint и тесты
    en_before = translations.Dictionary.read(translations.frontend_src() / "i18n" / "en.ts").entries
    assert translations.Dictionary.read(dictionaries / "en.ts").entries == en_before


def test_dry_run_changes_nothing(dictionaries, tmp_path):
    path = export(dictionaries, tmp_path)
    book = load_workbook(path)
    for row in book.active.iter_rows(min_row=2):
        if row[0].value == "Выйти":
            row[3].value = "Sign out now"
    book.save(path)
    before = (dictionaries / "en.ts").read_text(encoding="utf-8")
    output = StringIO()
    call_command("i18n_import", str(path), dry_run=True, frontend_dir=str(dictionaries), stdout=output)
    assert "Изменится: kk — 0, en — 1" in output.getvalue()
    assert (dictionaries / "en.ts").read_text(encoding="utf-8") == before
