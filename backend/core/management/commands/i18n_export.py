"""Выгрузить переводы в xlsx для вычитки человеком.

    python manage.py i18n_export --out translations.xlsx

Два листа: «Интерфейс» (словари фронта) и «Сервер» (каталог Django gettext).
Колонки: ключ | русский | казахский | английский | где встречается | комментарий.
Носитель языка правит колонки «Қазақша» и «English», остальное не трогает;
правки загружает `i18n_import`. Комментарий к строке объясняет, что нельзя
ломать: подстановки `{имя}` и формы числа через черту.
"""

from __future__ import annotations

from pathlib import Path

from django.core.management.base import BaseCommand
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core import translations

HEADER = ("Ключ", "Русский", "Қазақша", "English", "Где встречается", "Комментарий")
WIDTHS = (40, 60, 60, 60, 40, 50)
#: столько мест показываем в «где встречается», остальное — числом
PLACES = 5
SHEETS = ("Интерфейс", "Сервер")


def hint(key: str, *, is_plural: bool) -> str:
    """Что нельзя сломать при правке этой строки."""
    notes = []
    names = sorted(translations.placeholders(key))
    if names:
        notes.append("подстановки " + ", ".join("{" + name + "}" for name in names) + " оставить как есть")
    if is_plural:
        notes.append("формы числа через «|»: русских три (1|2|5), казахская одна, английских две (one|other)")
    return "; ".join(notes)


def fill(sheet, rows: list[tuple[str, str, str, list[str], bool]]) -> None:
    sheet.append(HEADER)
    for key, kk, en, places, is_plural in rows:
        where = ", ".join(places[:PLACES]) + (f" и ещё {len(places) - PLACES}" if len(places) > PLACES else "")
        # у строк ключ и есть русский текст; ключ не правят — по нему загрузка
        # находит строку, правят только «Қазақша» и «English»
        sheet.append((key, key, kk, en, where or "не используется", hint(key, is_plural=is_plural)))
    bold = Font(bold=True)
    editable = PatternFill("solid", fgColor="FFF6E5")
    for column, width in enumerate(WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
        sheet.cell(row=1, column=column).font = bold
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        row[2].fill = editable
        row[3].fill = editable
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions


class Command(BaseCommand):
    help = "Выгрузить переводы интерфейса и сервера в xlsx для вычитки"

    def add_arguments(self, parser):
        parser.add_argument("--out", default="translations.xlsx", help="куда записать таблицу")
        parser.add_argument("--frontend-dir", help="папка словарей фронта (по умолчанию I18N_FRONTEND_DIR)")
        parser.add_argument("--locale-dir", help="папка переводов сервера (по умолчанию backend/locale)")

    def handle(self, *args, out: str, frontend_dir: str | None = None, locale_dir: str | None = None, **options):
        dictionaries = translations.read_dictionaries(Path(frontend_dir) if frontend_dir else None)
        usages = translations.frontend_usages()
        plural = translations.plural_keys()
        keys = sorted(set(usages) | set(dictionaries["kk"].entries) | set(dictionaries["en"].entries))
        interface = [
            (
                key,
                dictionaries["kk"].entries.get(key, ""),
                dictionaries["en"].entries.get(key, ""),
                usages.get(key, []),
                key in plural,
            )
            for key in keys
        ]
        server_rows = translations.server_entries(Path(locale_dir) if locale_dir else None)
        server = [(key, row["kk"], row["en"], row["where"], "|" in key) for key, row in sorted(server_rows.items())]

        book = Workbook()
        first = book.active
        first.title = SHEETS[0]
        fill(first, interface)
        fill(book.create_sheet(SHEETS[1]), server)
        book.save(out)
        self.stdout.write(f"Выгружено строк: интерфейс — {len(interface)}, сервер — {len(server)} → {out}")
