"""Загрузить вычитанные переводы из xlsx обратно в словари.

    python manage.py i18n_import translations.xlsx [--dry-run]

Берутся колонки «Ключ», «Қазақша» и «English» таблицы, выгруженной
`i18n_export`. Перевод, у которого сломаны подстановки `{имя}` или формы
числа, не загружается — строка попадает в список отказов с причиной.
Ключей, которых нет в словарях, команда не заводит: строку интерфейса
добавляет код, а не таблица.
"""

from __future__ import annotations

from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from openpyxl import load_workbook

from core import translations

COLUMNS = {"Ключ": "key", "Қазақша": "kk", "English": "en"}


class Command(BaseCommand):
    help = "Загрузить вычитанные переводы из xlsx в словари интерфейса"

    def add_arguments(self, parser):
        parser.add_argument("file", help="таблица, выгруженная i18n_export")
        parser.add_argument("--dry-run", action="store_true", help="только показать, что изменится")
        parser.add_argument("--frontend-dir", help="папка словарей фронта (по умолчанию I18N_FRONTEND_DIR)")

    def handle(self, *args, file: str, dry_run: bool = False, frontend_dir: str | None = None, **options):
        dictionaries = translations.read_dictionaries(Path(frontend_dir) if frontend_dir else None)
        plural = translations.plural_keys()
        sheet = load_workbook(file, read_only=True).active
        rows = sheet.iter_rows(values_only=True)
        header = [str(cell or "").strip() for cell in next(rows, ())]
        where = {role: header.index(title) for title, role in COLUMNS.items() if title in header}
        if set(where) != set(COLUMNS.values()):
            raise CommandError("в таблице нет колонок «Ключ», «Қазақша» и «English» — это не выгрузка i18n_export")

        changed = {lang: 0 for lang in translations.LANGS}
        refused: list[str] = []
        unknown: list[str] = []
        for number, row in enumerate(rows, start=2):
            key = str(row[where["key"]] or "")
            if not key:
                continue
            if key not in dictionaries["kk"].entries and key not in dictionaries["en"].entries:
                unknown.append(f"строка {number}: «{key}»")
                continue
            for lang in translations.LANGS:
                value = row[where[lang]]
                value = "" if value is None else str(value)
                if value == dictionaries[lang].entries.get(key):
                    continue
                problem = translations.check_translation(key, value, is_plural=key in plural)
                if problem:
                    refused.append(f"строка {number}, {lang}: {problem} — «{value}»")
                    continue
                dictionaries[lang].entries[key] = value
                changed[lang] += 1

        if not dry_run:
            for lang in translations.LANGS:
                if changed[lang]:
                    dictionaries[lang].write()
        verb = "Изменится" if dry_run else "Изменено"
        self.stdout.write(f"{verb}: kk — {changed['kk']}, en — {changed['en']}")
        for line in refused:
            self.stdout.write(f"Не загружено: {line}")
        for line in unknown:
            self.stdout.write(f"Нет в словарях: {line}")
