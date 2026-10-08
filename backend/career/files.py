"""Файл теста профориентации: разбор книги xlsx и шаблон для учителя.

Формат для людей — `guides/CAREER_TESTS.md`. Книга из пяти листов:
«Тест» (название, инструкция, порог), «Ответы» (варианты ответа и их
баллы), «Шкалы», «Вопросы» (номер, текст, шкала, знак, при нужде свои
варианты) и «Интерпретация» (диапазоны баллов с подписями). Разбор ничего
не пишет: отчёт с ошибками по листу и строке показывается до «Применить».
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

from django.utils.translation import gettext as _

SHEET_TEST = "Тест"  # i18n-skip: названия листов книги — формат файла
SHEET_OPTIONS = "Ответы"  # i18n-skip: формат файла
SHEET_SCALES = "Шкалы"  # i18n-skip: формат файла
SHEET_ITEMS = "Вопросы"  # i18n-skip: формат файла
SHEET_RANGES = "Интерпретация"  # i18n-skip: формат файла

KEY_TITLE = "название"  # i18n-skip: ключи листа «Тест»
KEY_INSTRUCTION = "инструкция"  # i18n-skip: ключи листа «Тест»
KEY_THRESHOLD = "порог для разбора"  # i18n-skip: ключи листа «Тест»

#: заголовки колонок (без регистра); у вопросов со своими вариантами —
#: «Вариант 1», «Шкала 1», «Баллы 1» и дальше
COL_LABEL = "подпись"  # i18n-skip: формат файла
COL_VALUE = "баллы"  # i18n-skip: формат файла
COL_CODE = "код"  # i18n-skip: формат файла
COL_TITLE = "название"  # i18n-skip: формат файла
COL_DESCRIPTION = "описание"  # i18n-skip: формат файла
COL_NUMBER = "номер"  # i18n-skip: формат файла
COL_TEXT = "текст"  # i18n-skip: формат файла
COL_SCALE = "шкала"  # i18n-skip: формат файла
COL_SIGN = "знак"  # i18n-skip: формат файла
COL_LOW = "от"  # i18n-skip: формат файла
COL_HIGH = "до"  # i18n-skip: формат файла
CHOICE_RE = re.compile(r"^(вариант|шкала|баллы)\s*(\d+)$")  # i18n-skip: формат файла

MAX_ITEMS = 500
MAX_SCALES = 60
MAX_OPTIONS = 12
DEFAULT_THRESHOLD = 1


@dataclass
class ParsedChoice:
    label: str
    scale: str
    value: int = 1


@dataclass
class ParsedItem:
    number: int
    text: str
    scale: str
    sign: int = 1
    choices: list[ParsedChoice] = field(default_factory=list)


@dataclass
class ParsedRange:
    scale: str
    low: int
    high: int
    label: str


@dataclass
class ParsedTest:
    title: str = ""
    instruction: str = ""
    threshold: int = DEFAULT_THRESHOLD
    options: list[tuple[str, int]] = field(default_factory=list)
    scales: list[tuple[str, str, str]] = field(default_factory=list)
    items: list[ParsedItem] = field(default_factory=list)
    ranges: list[ParsedRange] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.errors

    def report(self) -> dict:
        """Отчёт для экрана: что нашлось и что не так."""
        return {
            "ok": self.ok,
            "title": self.title,
            "instruction": self.instruction,
            "threshold": self.threshold,
            "options": [{"label": label, "value": value} for label, value in self.options],
            "scales": [{"code": code, "title": title} for code, title, _desc in self.scales],
            "items": len(self.items),
            "ranges": len(self.ranges),
            "errors": self.errors,
            "warnings": self.warnings,
        }


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _key(value) -> str:
    return re.sub(r"\s+", " ", _text(value)).casefold()


def _int(value) -> int | None:
    text = _text(value).replace("−", "-").replace("–", "-").replace("—", "-").replace("+", "")
    if not text:
        return None
    try:
        return int(text)
    except ValueError:
        return None


def _sign(value) -> int | None:
    text = _text(value).replace("−", "-").replace("–", "-")
    if text in ("", "+", "1", "+1"):
        return 1
    if text in ("-", "-1"):
        return -1
    return None


def _sheet(book, name: str):
    wanted = _key(name)
    for sheet in book.worksheets:
        if _key(sheet.title) == wanted:
            return sheet
    return None


def _rows(sheet) -> list[list]:
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def _header(row: list) -> dict[str, int]:
    """Заголовок → номер колонки; пустые и повторные пропускаются."""
    out: dict[str, int] = {}
    for index, cell in enumerate(row):
        key = _key(cell)
        if key and key not in out:
            out[key] = index
    return out


def _where(sheet: str, row: int) -> str:
    return _("лист «{sheet}», строка {row}").format(sheet=sheet, row=row)


def parse(data: bytes) -> ParsedTest:
    """Разобрать книгу. Ошибки — списком с листом и строкой, база не трогается."""
    from openpyxl import load_workbook

    parsed = ParsedTest()
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        parsed.errors.append(_("Файл не читается как книга xlsx"))
        return parsed

    missing = [name for name in (SHEET_TEST, SHEET_SCALES, SHEET_ITEMS) if _sheet(book, name) is None]
    if missing:
        parsed.errors.append(_("Нет листов: {sheets}").format(sheets=", ".join(f"«{n}»" for n in missing)))
        return parsed

    _parse_test(parsed, _sheet(book, SHEET_TEST))
    options_sheet = _sheet(book, SHEET_OPTIONS)
    if options_sheet is not None:
        _parse_options(parsed, options_sheet)
    _parse_scales(parsed, _sheet(book, SHEET_SCALES))
    _parse_items(parsed, _sheet(book, SHEET_ITEMS))
    ranges_sheet = _sheet(book, SHEET_RANGES)
    if ranges_sheet is not None:
        _parse_ranges(parsed, ranges_sheet)
    _cross_check(parsed)
    return parsed


def _parse_test(parsed: ParsedTest, sheet) -> None:
    for row in _rows(sheet):
        if len(row) < 2:
            continue
        key, value = _key(row[0]), _text(row[1])
        if key == KEY_TITLE:
            parsed.title = value[:200]
        elif key == KEY_INSTRUCTION:
            parsed.instruction = value
        elif key == KEY_THRESHOLD:
            number = _int(value)
            if number is None:
                parsed.errors.append(_("Лист «{sheet}»: «Порог для разбора» — целое число").format(sheet=SHEET_TEST))
            else:
                parsed.threshold = number
    if not parsed.title:
        parsed.errors.append(_("Лист «{sheet}»: нет строки «Название»").format(sheet=SHEET_TEST))


def _parse_options(parsed: ParsedTest, sheet) -> None:
    rows = _rows(sheet)
    if not rows:
        return
    head = _header(rows[0])
    if COL_LABEL not in head or COL_VALUE not in head:
        parsed.errors.append(_("Лист «{sheet}»: нужны колонки «Подпись» и «Баллы»").format(sheet=SHEET_OPTIONS))
        return
    seen: set[str] = set()
    for index, row in enumerate(rows[1:], start=2):
        label = _text(row[head[COL_LABEL]] if head[COL_LABEL] < len(row) else None)
        if not label:
            continue
        value = _int(row[head[COL_VALUE]] if head[COL_VALUE] < len(row) else None)
        if value is None:
            parsed.errors.append(
                _("{where}: у ответа «{label}» нет баллов").format(where=_where(SHEET_OPTIONS, index), label=label)
            )
            continue
        if label.casefold() in seen:
            parsed.errors.append(
                _("{where}: ответ «{label}» повторяется").format(where=_where(SHEET_OPTIONS, index), label=label)
            )
            continue
        seen.add(label.casefold())
        parsed.options.append((label[:60], value))
    if len(parsed.options) > MAX_OPTIONS:
        parsed.errors.append(
            _("Лист «{sheet}»: вариантов ответа не больше {n}").format(sheet=SHEET_OPTIONS, n=MAX_OPTIONS)
        )


def _parse_scales(parsed: ParsedTest, sheet) -> None:
    rows = _rows(sheet)
    if not rows:
        parsed.errors.append(_("Лист «{sheet}» пуст").format(sheet=SHEET_SCALES))
        return
    head = _header(rows[0])
    if COL_CODE not in head or COL_TITLE not in head:
        parsed.errors.append(_("Лист «{sheet}»: нужны колонки «Код» и «Название»").format(sheet=SHEET_SCALES))
        return
    seen: set[str] = set()
    for index, row in enumerate(rows[1:], start=2):
        code = _text(row[head[COL_CODE]] if head[COL_CODE] < len(row) else None)
        if not code:
            continue
        title = _text(row[head[COL_TITLE]] if head[COL_TITLE] < len(row) else None)
        if not title:
            parsed.errors.append(
                _("{where}: у шкалы «{code}» нет названия").format(where=_where(SHEET_SCALES, index), code=code)
            )
            continue
        if code.casefold() in seen:
            parsed.errors.append(
                _("{where}: код шкалы «{code}» повторяется").format(where=_where(SHEET_SCALES, index), code=code)
            )
            continue
        seen.add(code.casefold())
        description = ""
        if COL_DESCRIPTION in head and head[COL_DESCRIPTION] < len(row):
            description = _text(row[head[COL_DESCRIPTION]])
        parsed.scales.append((code[:40], title[:120], description))
    if not parsed.scales:
        parsed.errors.append(_("Лист «{sheet}»: ни одной шкалы").format(sheet=SHEET_SCALES))
    if len(parsed.scales) > MAX_SCALES:
        parsed.errors.append(_("Лист «{sheet}»: шкал не больше {n}").format(sheet=SHEET_SCALES, n=MAX_SCALES))


def _parse_items(parsed: ParsedTest, sheet) -> None:
    rows = _rows(sheet)
    if not rows:
        parsed.errors.append(_("Лист «{sheet}» пуст").format(sheet=SHEET_ITEMS))
        return
    head = _header(rows[0])
    for needed, title in ((COL_NUMBER, "Номер"), (COL_TEXT, "Текст")):  # i18n-skip: названия колонок — формат файла
        if needed not in head:
            parsed.errors.append(_("Лист «{sheet}»: нет колонки «{column}»").format(sheet=SHEET_ITEMS, column=title))
    if COL_NUMBER not in head or COL_TEXT not in head:
        return
    # колонки своих вариантов: номер варианта → (вариант, шкала, баллы)
    choice_cols: dict[int, dict[str, int]] = {}
    for key, index in head.items():
        match = CHOICE_RE.match(key)
        if match:
            choice_cols.setdefault(int(match.group(2)), {})[match.group(1)] = index
    seen: set[int] = set()

    def cell(row, column: str):
        index = head.get(column)
        return row[index] if index is not None and index < len(row) else None

    for index, row in enumerate(rows[1:], start=2):
        if all(_text(c) == "" for c in row):
            continue
        where = _where(SHEET_ITEMS, index)
        number = _int(cell(row, COL_NUMBER))
        text = _text(cell(row, COL_TEXT))
        if number is None or number < 1:
            parsed.errors.append(_("{where}: номер вопроса — целое число от 1").format(where=where))
            continue
        if number in seen:
            parsed.errors.append(_("{where}: номер {number} повторяется").format(where=where, number=number))
            continue
        if not text:
            parsed.errors.append(_("{where}: у вопроса {number} нет текста").format(where=where, number=number))
            continue
        seen.add(number)
        sign = _sign(cell(row, COL_SIGN))
        if sign is None:
            parsed.errors.append(_("{where}: знак — «+» или «−»").format(where=where))
            sign = 1
        choices: list[ParsedChoice] = []
        for order in sorted(choice_cols):
            cols = choice_cols[order]
            label = (
                _text(row[cols["вариант"]]) if "вариант" in cols and cols["вариант"] < len(row) else ""
            )  # i18n-skip: формат файла
            scale = (
                _text(row[cols["шкала"]]) if "шкала" in cols and cols["шкала"] < len(row) else ""
            )  # i18n-skip: формат файла
            value_raw = (
                row[cols["баллы"]] if "баллы" in cols and cols["баллы"] < len(row) else None
            )  # i18n-skip: формат файла
            if not label and not scale:
                continue
            if not label or not scale:
                parsed.errors.append(_("{where}: у варианта {n} нужны и текст, и шкала").format(where=where, n=order))
                continue
            value = _int(value_raw)
            if _text(value_raw) and value is None:
                parsed.errors.append(_("{where}: баллы варианта {n} — целое число").format(where=where, n=order))
                continue
            choices.append(ParsedChoice(label=label[:400], scale=scale, value=1 if value is None else value))
        scale = _text(cell(row, COL_SCALE))
        if not scale and not choices:
            parsed.errors.append(_("{where}: у вопроса {number} нет шкалы").format(where=where, number=number))
            continue
        parsed.items.append(ParsedItem(number=number, text=text, scale=scale, sign=sign, choices=choices))
    if not parsed.items:
        parsed.errors.append(_("Лист «{sheet}»: ни одного вопроса").format(sheet=SHEET_ITEMS))
    if len(parsed.items) > MAX_ITEMS:
        parsed.errors.append(_("Лист «{sheet}»: вопросов не больше {n}").format(sheet=SHEET_ITEMS, n=MAX_ITEMS))
    parsed.items.sort(key=lambda item: item.number)


def _parse_ranges(parsed: ParsedTest, sheet) -> None:
    rows = _rows(sheet)
    if not rows:
        return
    head = _header(rows[0])
    for needed, title in ((COL_LOW, "От"), (COL_HIGH, "До"), (COL_LABEL, "Подпись")):  # i18n-skip: формат файла
        if needed not in head:
            parsed.errors.append(_("Лист «{sheet}»: нет колонки «{column}»").format(sheet=SHEET_RANGES, column=title))
            return
    for index, row in enumerate(rows[1:], start=2):
        if all(_text(c) == "" for c in row):
            continue
        where = _where(SHEET_RANGES, index)
        low, high = _int(row[head[COL_LOW]] if head[COL_LOW] < len(row) else None), _int(
            row[head[COL_HIGH]] if head[COL_HIGH] < len(row) else None
        )
        label = _text(row[head[COL_LABEL]] if head[COL_LABEL] < len(row) else None)
        if low is None or high is None:
            parsed.errors.append(_("{where}: «От» и «До» — целые числа").format(where=where))
            continue
        if low > high:
            parsed.errors.append(_("{where}: «От» больше «До»").format(where=where))
            continue
        if not label:
            parsed.errors.append(_("{where}: нет подписи диапазона").format(where=where))
            continue
        scale = ""
        if COL_SCALE in head and head[COL_SCALE] < len(row):
            scale = _text(row[head[COL_SCALE]])
            if scale == "*":
                scale = ""
        parsed.ranges.append(ParsedRange(scale=scale, low=low, high=high, label=label[:120]))


def _cross_check(parsed: ParsedTest) -> None:
    """Связи между листами: шкалы вопросов есть, у вопросов без вариантов есть ответы."""
    codes = {code.casefold(): code for code, _title, _desc in parsed.scales}
    used: set[str] = set()
    needs_options = False
    for item in parsed.items:
        if item.choices:
            for choice in item.choices:
                if choice.scale.casefold() not in codes:
                    parsed.errors.append(
                        _("Вопрос {number}: шкалы «{scale}» нет на листе «{sheet}»").format(
                            number=item.number, scale=choice.scale, sheet=SHEET_SCALES
                        )
                    )
                else:
                    choice.scale = codes[choice.scale.casefold()]
                    used.add(choice.scale)
            if item.scale and item.scale.casefold() not in codes:
                parsed.errors.append(
                    _("Вопрос {number}: шкалы «{scale}» нет на листе «{sheet}»").format(
                        number=item.number, scale=item.scale, sheet=SHEET_SCALES
                    )
                )
            elif item.scale:
                item.scale = codes[item.scale.casefold()]
            continue
        needs_options = True
        if item.scale.casefold() not in codes:
            parsed.errors.append(
                _("Вопрос {number}: шкалы «{scale}» нет на листе «{sheet}»").format(
                    number=item.number, scale=item.scale, sheet=SHEET_SCALES
                )
            )
        else:
            item.scale = codes[item.scale.casefold()]
            used.add(item.scale)
    if needs_options and len(parsed.options) < 2:
        parsed.errors.append(
            _("Лист «{sheet}»: нужны хотя бы два варианта ответа — у вопросов нет своих").format(sheet=SHEET_OPTIONS)
        )
    for rng in parsed.ranges:
        if rng.scale and rng.scale.casefold() not in codes:
            parsed.errors.append(
                _("Интерпретация: шкалы «{scale}» нет на листе «{sheet}»").format(scale=rng.scale, sheet=SHEET_SCALES)
            )
        elif rng.scale:
            rng.scale = codes[rng.scale.casefold()]
    for code, title, _desc in parsed.scales:
        if code not in used:
            parsed.warnings.append(_("Шкала «{title}» не встречается ни в одном вопросе").format(title=title))
    numbers = [item.number for item in parsed.items]
    if numbers and numbers != list(range(1, len(numbers) + 1)):
        parsed.warnings.append(_("Номера вопросов идут не подряд: ученик увидит их по порядку номеров"))
    if not parsed.ranges:
        parsed.warnings.append(_("Диапазонов интерпретации нет: у баллов не будет подписей"))


def template() -> bytes:
    """Пустая книга с заголовками и образцом строк — скачивается с экрана."""
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = SHEET_TEST
    for row in TEMPLATE_TEST:
        sheet.append(list(row))
    for title, rows in (
        (SHEET_OPTIONS, TEMPLATE_OPTIONS),
        (SHEET_SCALES, TEMPLATE_SCALES),
        (SHEET_ITEMS, TEMPLATE_ITEMS),
        (SHEET_RANGES, TEMPLATE_RANGES),
    ):
        page = book.create_sheet(title)
        for row in rows:
            page.append(list(row))
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


# Содержимое шаблона — формат файла, а не интерфейс: строки не переводятся
TEMPLATE_TEST = (  # i18n-skip: содержимое шаблона файла
    ("Название", "Название теста"),
    ("Инструкция", "Что сказать ученику перед началом"),
    ("Порог для разбора", DEFAULT_THRESHOLD),
)
TEMPLATE_OPTIONS = (  # i18n-skip: содержимое шаблона файла
    ("Подпись", "Баллы", "Описание"),
    ("++", 2, "очень хотелось бы"),
    ("+", 1, "нравится"),
    ("0", 0, "затрудняюсь ответить"),
    ("−", -1, "не нравится"),
    ("−−", -2, "совсем не хотелось бы"),
)
TEMPLATE_SCALES = (  # i18n-skip: содержимое шаблона файла
    ("Код", "Название", "Описание"),
    ("bio", "Биология", ""),
    ("phys", "Физика", ""),
)
TEMPLATE_ITEMS = (  # i18n-skip: содержимое шаблона файла
    (
        "Номер",
        "Текст",
        "Шкала",
        "Знак",
        "Вариант 1",
        "Шкала 1",
        "Баллы 1",
        "Вариант 2",
        "Шкала 2",
        "Баллы 2",
    ),
    (1, "Изучать разнообразие животного и растительного мира", "bio", "+"),
    (2, "Проводить физические эксперименты", "phys", "+"),
)
TEMPLATE_RANGES = (  # i18n-skip: содержимое шаблона файла
    ("Шкала", "От", "До", "Подпись"),
    ("", -12, -6, "активно отрицается"),
    ("", -5, -1, "интереса не вызывает"),
    ("", 0, 0, "не определён"),
    ("", 1, 4, "интерес выражен слабо"),
    ("", 5, 7, "выраженный интерес"),
    ("", 8, 12, "ярко выраженный интерес"),
)
