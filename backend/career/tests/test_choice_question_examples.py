"""Образцы с вопросами со своими вариантами: Йовайши и ДДО Климова из `guides/examples`.

Файлы грузятся как есть, ученики отвечают через API ответов, суммы по шкалам
сверяются с подсчётом по ключу, прочитанному из книги напрямую, а не через разбор.
"""

from __future__ import annotations

from collections import Counter

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook

from career import analysis
from career.models import CareerAttempt, CareerTest
from career.scoring import bounds

from .conftest import EXAMPLE, activate, login

pytestmark = pytest.mark.django_db

YOVAISHI = EXAMPLE.with_name("yovaishi.xlsx")
KLIMOV = EXAMPLE.with_name("klimov_ddo.xlsx")


def _upload(teacher, path) -> CareerTest:
    client = login(teacher)
    preview = client.post(
        "/api/career/tests/preview/", {"file": SimpleUploadedFile(path.name, path.read_bytes())}, format="multipart"
    )
    assert preview.status_code == 200, preview.content
    report = preview.json()
    assert report["ok"] and report["errors"] == [] and report["warnings"] == [], report
    response = client.post(
        "/api/career/tests/", {"file": SimpleUploadedFile(path.name, path.read_bytes())}, format="multipart"
    )
    assert response.status_code == 201, response.content
    return CareerTest.objects.get(pk=response.data["id"])


def _key(path) -> list[list[tuple[str, str, int]]]:
    """Ключ из книги без разбора платформы: по вопросу — (вариант, шкала, баллы)."""
    rows = list(load_workbook(path, read_only=True)["Вопросы"].iter_rows(values_only=True))
    out = []
    for row in rows[1:]:
        choices = [(row[i], row[i + 1], row[i + 2] or 1) for i in range(2, len(row) - 2, 3) if row[i]]
        out.append(choices)
    return out


def _manual(path, pick) -> dict[str, int]:
    totals: Counter[str] = Counter()
    for choices in _key(path):
        _label, scale, value = pick(choices)
        totals[scale] += value
    return dict(totals)


def _first_ii(choices):
    return next((c for c in choices if c[1] == "ii"), choices[0])


def _first(choices):
    return choices[0]


def _last(choices):
    return choices[-1]


def _pass(student, test, pick) -> CareerAttempt:
    """Пройти тест учеником: `pick` выбирает вариант по коду шкалы, как в ключе."""
    client = login(student.user)
    started = client.post(f"/api/career/my/tests/{test.pk}/start/")
    assert started.status_code == 200, started.content
    payload = started.json()
    assert payload["options"] == []
    scale_of = {c.pk: c.scale.code for item in test.items.all() for c in item.choices.select_related("scale")}
    answers = []
    for item in payload["items"]:
        assert "scale" not in item and all(set(c) == {"id", "label"} for c in item["choices"])
        rows = [(c["label"], scale_of[c["id"]], c["id"]) for c in item["choices"]]
        answers.append({"item": item["id"], "choice": pick(rows)[2]})
    saved = client.post(f"/api/career/my/attempts/{payload['id']}/answers/", {"answers": answers}, format="json")
    assert saved.status_code == 200 and saved.json()["answered"] == saved.json()["total"], saved.content
    done = client.post(f"/api/career/my/attempts/{payload['id']}/finish/")
    assert done.status_code == 200, done.content
    return CareerAttempt.objects.get(pk=payload["id"])


def _scores(attempt) -> dict[str, tuple[int, str]]:
    return {row.scale.code: (row.score, row.label) for row in attempt.scores.select_related("scale")}


YOVAISHI_LABEL = {range(0, 4): "склонность не выражена", range(4, 7): "слабо выраженная склонность"}
YOVAISHI_LABEL |= {range(7, 10): "склонность к этому виду деятельности", range(10, 13): "ярко выраженная склонность"}
KLIMOV_LABEL = {range(0, 4): "интерес слабый или не выражен", range(4, 6): "умеренный интерес"}
KLIMOV_LABEL |= {range(6, 9): "выраженный интерес к этому типу профессий"}


def _label(table, score: int) -> str:
    return next(label for span, label in table.items() if score in span)


def test_both_examples_load_clean_with_bounds(teacher):
    yov, kli = _upload(teacher, YOVAISHI), _upload(teacher, KLIMOV)
    assert (yov.items.count(), yov.scales.count(), yov.analysis_min_score) == (24, 6, 7)
    assert (kli.items.count(), kli.scales.count(), kli.analysis_min_score) == (20, 5, 4)
    assert all(item.choices.count() == 3 for item in yov.items.all())
    assert all(item.choices.count() == 2 for item in kli.items.all())
    assert not yov.options.exists() and not kli.options.exists()
    # в вопросе 7 оба первых варианта ведут в «iii» (так в ключе Йовайши): выбирается один,
    # поэтому у «iii» 12 вариантов, но не больше 11 баллов
    limits = {yov.scales.get(pk=pk).code: span for pk, span in bounds(yov).items()}
    assert limits == {"i": (0, 12), "ii": (0, 12), "iii": (0, 11), "iv": (0, 12), "v": (0, 12), "vi": (0, 12)}
    assert set(bounds(kli).values()) == {(0, 8)} and len(bounds(kli)) == 5


def test_scores_match_the_key_and_retake_archives(teacher, boston, pupils):
    yov, kli = _upload(teacher, YOVAISHI), _upload(teacher, KLIMOV)
    for test in (yov, kli):
        activate(test, [{"group": boston.pk, "students": None}], by=teacher)
    a, b = pupils["aliya"], pupils["damir"]
    cases = [
        (a, yov, YOVAISHI, _first_ii, YOVAISHI_LABEL),
        (a, kli, KLIMOV, _first, KLIMOV_LABEL),
        (b, yov, YOVAISHI, _last, YOVAISHI_LABEL),
        (b, kli, KLIMOV, _last, KLIMOV_LABEL),
    ]
    attempts = {}
    for student, test, path, pick, labels in cases:
        attempt = _pass(student, test, pick)
        expected = _manual(path, pick)
        got = _scores(attempt)
        assert set(got) == {s.code for s in test.scales.all()}
        assert {code: score for code, (score, _l) in got.items()} == {c: expected.get(c, 0) for c in got}
        assert all(label == _label(labels, score) for score, label in got.values())
        attempts[(student.pk, test.pk)] = attempt
    # ручной счёт по ключу: ученик А в Йовайши выбирает «ii», где она есть (12 вопросов)
    assert _scores(attempts[(a.pk, yov.pk)])["ii"] == (12, "ярко выраженная склонность")
    assert _scores(attempts[(b.pk, yov.pk)])["vi"] == (12, "ярко выраженная склонность")
    assert _scores(attempts[(a.pk, kli.pk)])["nature"] == (6, "выраженный интерес к этому типу профессий")
    assert _scores(attempts[(b.pk, kli.pk)])["sign"] == (6, "выраженный интерес к этому типу профессий")

    # ученику в результатах — шесть и пять полосок с границами и подписью
    client = login(a.user)
    result = client.get(f"/api/career/my/attempts/{attempts[(a.pk, yov.pk)].pk}/").json()
    assert len(result["scores"]) == 6 and all(
        r["low"] == 0 and r["high"] in (11, 12) and r["label"] for r in result["scores"]
    )
    result = client.get(f"/api/career/my/attempts/{attempts[(a.pk, kli.pk)].pk}/").json()
    assert len(result["scores"]) == 5 and all((r["low"], r["high"]) == (0, 8) for r in result["scores"])

    old = attempts[(a.pk, kli.pk)]
    assert login(teacher).post(f"/api/career/attempts/{old.pk}/retake/").status_code == 200
    old.refresh_from_db()
    assert old.archived_at is not None
    row = next(t for t in client.get("/api/career/my/").json()["tests"] if t["id"] == kli.pk)
    assert row["status"] == "none" and row["attempt"] is None
    fresh = _pass(a, kli, _first)
    assert fresh.pk != old.pk and CareerAttempt.objects.filter(student=a, test=kli).count() == 2


def test_prompt_carries_each_test_above_its_threshold(teacher, boston, pupils):
    yov, kli = _upload(teacher, YOVAISHI), _upload(teacher, KLIMOV)
    for test in (yov, kli):
        activate(test, [{"group": boston.pk, "students": None}], by=teacher)
    student = pupils["aliya"]
    attempts = [_pass(student, yov, _first_ii), _pass(student, kli, _first)]
    text = analysis._prompt(student, attempts, [])
    assert f"Тест «{yov.title}»" in text and f"Тест «{kli.title}»" in text
    # Йовайши, порог 7: только «ii» (12); Климов, порог 4: всё, кроме «sign» (2)
    assert "- Исследовательская (интеллектуальная) работа: 12 — ярко выраженная склонность" in text
    assert "Работа с людьми: 6" not in text
    assert "- Человек — природа: 6" in text and "Человек — знаковая система" not in text
