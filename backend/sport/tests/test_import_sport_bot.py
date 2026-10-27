"""Перенос из бота на вымышленном файле: проверка ничего не пишет, применение — без дублей."""

from __future__ import annotations

import datetime as dt
from io import StringIO

import pytest
from django.core.management import call_command
from openpyxl import Workbook, load_workbook

from sport.models import EnrollmentSource, SectionEnrollment, SectionSlot
from students.models import Student, StudyGroup

pytestmark = pytest.mark.django_db

HEAD = ("Имя", "Класс", "Телефон", "tg_id", "Дата записи")


@pytest.fixture
def school():
    nine = StudyGroup.objects.create(code="NINE", parallel=9)
    ten = StudyGroup.objects.create(code="TEN", parallel=10)
    make = Student.objects.create
    return {
        "aruzhan": make(last_name="Қасымова", first_name="Аружан", group=ten, graduation_year=2028),
        "timur": make(last_name="Сапаров", first_name="Тимур", group=nine, graduation_year=2029),
        "dana1": make(last_name="Касымова", first_name="Дана", group=ten, graduation_year=2028),
        "dana2": make(last_name="Касымова", first_name="Дана", group=ten, graduation_year=2028),
        "twice": make(last_name="Ермеков", first_name="Алихан", group=ten, graduation_year=2028),
    }


def bot_file(tmp_path):
    book = Workbook()
    book.remove(book.active)
    sheets = {
        "Теннис": [
            ("аружан касымова", "10.0", "+77000000000", 111, dt.datetime(2026, 9, 10, 14, 5)),
            ("Сапаров Тимур", 9, "+77000000001", 112, "11.09.2026 10:00"),
            ("Касымова Дана", "10", "", 113, dt.datetime(2026, 9, 12, 9, 0)),
            ("Неизвестный Ученик", "10", "", 114, dt.datetime(2026, 9, 12, 9, 0)),
            ("Ермеков Алихан", "10", "", 115, dt.datetime(2026, 9, 12, 9, 0)),
        ],
        "Футбол": [("Алихан Ермеков", "10", "", 115, dt.datetime(2026, 9, 13, 9, 0))],
    }
    for title, rows in sheets.items():
        sheet = book.create_sheet(title)
        sheet.append(HEAD)
        for row in rows:
            sheet.append(row)
    book.create_sheet("Ученики").append(("tg_id", "Имя"))
    path = tmp_path / "bot.xlsx"
    book.save(path)
    return path


def run(path, *args):
    out = StringIO()
    call_command("import_sport_bot", str(path), *args, stdout=out)
    return out.getvalue()


def test_check_writes_nothing_and_reports(school, tmp_path):
    path = bot_file(tmp_path)
    text = run(path)
    assert "нашёлся: 2" in text and "не нашёлся: 1" in text
    assert "несколько кандидатов: 1" in text and "в двух секциях: 2" in text
    assert not SectionEnrollment.objects.exists()
    report = load_workbook(tmp_path / "bot_отчёт.xlsx")
    assert report.sheetnames == [
        "Нашлись",
        "Не нашлись",
        "Несколько кандидатов",
        "В двух секциях",
        "Уже в LMS",
        "Переполнение",
    ]
    # телефонов и tg_id в отчёте нет
    cells = {str(c.value) for sheet in report for row in sheet.iter_rows() for c in row}
    assert "+77000000000" not in cells and "111" not in cells


def test_apply_sets_hour_only_when_one_slot_and_is_idempotent(school, tmp_path):
    path = bot_file(tmp_path)
    assert "Записано в LMS: 2" in run(path, "--apply")
    aruzhan = SectionEnrollment.objects.get(student=school["aruzhan"])
    timur = SectionEnrollment.objects.get(student=school["timur"])
    # у 10–11 три часа тенниса — без часа; у 8–9 один — час ставится сам
    assert aruzhan.slot is None and aruzhan.source == EnrollmentSource.BOT
    assert aruzhan.enrolled_at.date() == dt.date(2026, 9, 10)
    assert timur.slot == SectionSlot.objects.get(section__title="Теннис", grade_from=8)
    assert not SectionEnrollment.objects.filter(student=school["twice"]).exists()
    text = run(path, "--apply")
    assert "Записано в LMS: 0" in text and "уже перенесён: 2" in text
    assert SectionEnrollment.objects.count() == 2
