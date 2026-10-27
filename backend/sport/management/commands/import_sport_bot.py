"""Перенос записей в секции из Телеграм-бота: сначала проверка, с `--apply` — запись.

Источник правды — листы секций («Теннис», «Волейбол», «ММА», «Футбол»,
«Шахматы»): Имя | Класс | Телефон | tg_id | Дата записи. Сводный лист
«Ученики» не читается. Ученик ищется по ФИО (порядок слов, регистр,
казахские буквы и опечатки — `suggestions.name_matching`) среди учеников
своего класса. Телефоны и tg_id не читаются и не сохраняются.

Без `--apply` — только отчёт: в консоль числа, в xlsx построчно — нашёлся,
не нашёлся, несколько кандидатов, записан в две секции, переполнение,
уже записан в LMS. С `--apply` — записи с датой из файла и источником
«перенос из бота»: у группы классов одно занятие — час ставится сам,
несколько — без часа (ученик видит «Выбери время»). Лимиты при переносе
не проверяются: бот их уже проверял, переполнение уходит в отчёт.
Повторный запуск дублей не создаёт. Ученик из двух секций не переносится
никуда — его решает директор спорта руками.
"""

# i18n-skip-file: команда владельца в терминале и отчёт для него же — не интерфейс

from __future__ import annotations

import datetime as dt
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils import timezone

from core.parallels import parallel_of
from sport import services
from sport.models import EnrollmentSource, Section, SectionEnrollment, SectionSlot
from students.models import Student
from suggestions import name_matching

SHEETS = ("Теннис", "Волейбол", "ММА", "Футбол", "Шахматы")
GRADES = (8, 9, 10, 11)

FOUND = "нашёлся"
MISSING = "не нашёлся"
AMBIGUOUS = "несколько кандидатов"
TWICE = "в двух секциях"
ALREADY = "уже записан в LMS"
DONE = "уже перенесён"


@dataclass
class Row:
    sheet: str
    line: int
    name: str
    grade: int | None
    enrolled_at: dt.datetime | None
    state: str = ""
    student: Student | None = None
    note: str = ""
    slot: SectionSlot | None = None
    candidates: list[str] = field(default_factory=list)


def parse_grade(raw) -> int | None:
    """«9», «9.0», 9.0 — девятый класс; остальное — нет класса."""
    try:
        value = int(float(str(raw).strip().replace(",", ".")))
    except (TypeError, ValueError):
        return None
    return value if value in GRADES else None


def parse_moment(raw) -> dt.datetime | None:
    if isinstance(raw, dt.datetime):
        moment = raw
    elif isinstance(raw, dt.date):
        moment = dt.datetime.combine(raw, dt.time(12, 0))
    else:
        text = str(raw or "").strip()
        moment = None
        for pattern in (
            "%Y-%m-%d %H:%M:%S",
            "%Y-%m-%d %H:%M",
            "%d.%m.%Y %H:%M:%S",
            "%d.%m.%Y %H:%M",
            "%d.%m.%Y",
            "%Y-%m-%d",
        ):
            try:
                moment = dt.datetime.strptime(text, pattern)
                break
            except ValueError:
                continue
        if moment is None:
            return None
    return timezone.make_aware(moment) if timezone.is_naive(moment) else moment


def read_rows(path: Path) -> list[Row]:
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    rows: list[Row] = []
    for title in SHEETS:
        if title not in book.sheetnames:
            continue
        values = list(book[title].iter_rows(values_only=True))
        if not values:
            continue
        head = [str(cell or "").strip() for cell in values[0]]
        try:
            name_at, grade_at, date_at = head.index("Имя"), head.index("Класс"), head.index("Дата записи")
        except ValueError as error:
            raise CommandError(f"Лист «{title}»: нет колонки {error}") from error
        for line, cells in enumerate(values[1:], start=2):
            name = str(cells[name_at] or "").strip() if name_at < len(cells) else ""
            if not name:
                continue
            rows.append(
                Row(
                    sheet=title,
                    line=line,
                    name=" ".join(name.split()),
                    grade=parse_grade(cells[grade_at] if grade_at < len(cells) else None),
                    enrolled_at=parse_moment(cells[date_at] if date_at < len(cells) else None),
                )
            )
    return rows


def match(rows: list[Row]) -> None:
    """Каждой строке — ученик своего класса или причина, почему его нет."""
    everyone = list(Student.objects.filter(is_active=True).select_related("group"))
    by_grade: dict[int, list[Student]] = defaultdict(list)
    for student in everyone:
        by_grade[parallel_of(student)].append(student)
    for row in rows:
        if row.grade is None:
            row.state, row.note = MISSING, "в файле нет класса 8–11"
            continue
        outcome = name_matching.find(row.name, students=by_grade[row.grade])
        if outcome.is_confident and outcome.best:
            row.state = FOUND
            row.student = next(s for s in by_grade[row.grade] if s.pk == outcome.best.student_id)
        elif outcome.is_missing:
            row.state = MISSING
            elsewhere = name_matching.find(row.name, students=everyone)
            if elsewhere.is_confident and elsewhere.best:
                row.note = f"похожий в другом классе: {elsewhere.best.full_name}, {elsewhere.best.group_code or ''}"
        else:
            row.state = AMBIGUOUS
            row.candidates = [f"{c.full_name} ({c.group_code or ''}, {c.confidence:.2f})" for c in outcome.candidates]


def resolve(rows: list[Row], sections: dict[str, Section]) -> list[str]:
    """Два раза в разных секциях, уже в LMS, час, переполнение — и строки отчёта о переполнении."""
    found = [r for r in rows if r.state == FOUND]
    seen: dict[int, set[str]] = defaultdict(set)
    for row in found:
        seen[row.student.pk].add(row.sheet)
    first_line: dict[tuple[int, str], int] = {}
    for row in found:
        if len(seen[row.student.pk]) > 1:
            row.state = TWICE
            row.note = ", ".join(sorted(seen[row.student.pk]))
            continue
        key = (row.student.pk, row.sheet)
        if key in first_line:
            row.state, row.note = DONE, f"повтор строки {first_line[key]} того же листа"
            continue
        first_line[key] = row.line
    active = {
        e.student_id: e
        for e in SectionEnrollment.objects.filter(
            student__in=[r.student for r in rows if r.state == FOUND], ended_at__isnull=True
        ).select_related("section")
    }
    for row in rows:
        if row.state != FOUND:
            continue
        current = active.get(row.student.pk)
        if current is not None:
            if current.source == EnrollmentSource.BOT and current.section.title == row.sheet:
                row.state = DONE
            else:
                row.state, row.note = ALREADY, f"{current.section.title}, {current.get_source_display()}"
            continue
        section = sections.get(row.sheet)
        if section is None:
            row.state, row.note = MISSING, "в LMS нет такой секции"
            continue
        fitting = [s for s in section.slots.all() if s.archived_at is None and s.fits(row.grade)]
        if not fitting:
            row.state, row.note = MISSING, "у секции нет занятия для этого класса"
            continue
        row.slot = fitting[0] if len(fitting) == 1 else None
        row.note = services.slot_words(row.slot) if row.slot else "без часа — выберет сам"
    overflow = []
    for section in sections.values():
        groups: dict[tuple[int, int], list[SectionSlot]] = defaultdict(list)
        for slot in section.slots.all():
            if slot.archived_at is None:
                groups[(slot.parallel_from, slot.parallel_to)].append(slot)
        for (low, high), slots in groups.items():
            coming = sum(
                1 for r in rows if r.state in (FOUND, DONE) and r.sheet == section.title and low <= r.grade <= high
            )
            taken = services.active_enrollments().filter(slot__in=slots).count()
            no_hour = services.active_enrollments().filter(section=section, slot__isnull=True).count()
            capacity = sum(s.capacity for s in slots)
            if coming > capacity:
                overflow.append(
                    f"{section.title} {low}–{high}: из бота {coming}, мест {capacity}, сейчас в LMS {taken + no_hour}"
                )
    return overflow


def write_report(path: Path, rows: list[Row], overflow: list[str]) -> None:
    from openpyxl import Workbook

    book = Workbook()
    book.remove(book.active)
    head = ("Лист", "Строка", "Имя в файле", "Класс", "Дата записи", "Ученик в LMS", "Группа", "Час", "Пометка")
    for title, states in (
        ("Нашлись", (FOUND,)),
        ("Не нашлись", (MISSING,)),
        ("Несколько кандидатов", (AMBIGUOUS,)),
        ("В двух секциях", (TWICE,)),
        ("Уже в LMS", (ALREADY, DONE)),
    ):
        sheet = book.create_sheet(title)
        sheet.append(head)
        for row in rows:
            if row.state not in states:
                continue
            sheet.append(
                (
                    row.sheet,
                    row.line,
                    row.name,
                    row.grade or "",
                    timezone.localtime(row.enrolled_at).strftime("%d.%m.%Y %H:%M") if row.enrolled_at else "",
                    row.student.full_name if row.student else "",
                    row.student.group.code if row.student and row.student.group_id else "",
                    services.slot_words(row.slot) if row.slot else "",
                    (
                        "; ".join(row.candidates)
                        if row.candidates
                        else (row.note if row.state != DONE else row.note or row.state)
                    ),
                )
            )
    sheet = book.create_sheet("Переполнение")
    sheet.append(("Секция и классы",))
    for line in overflow:
        sheet.append((line,))
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)


class Command(BaseCommand):
    help = "Перенести записи в секции из файла бота: без --apply — только отчёт"

    def add_arguments(self, parser):
        parser.add_argument("file", help="xlsx из бота: листы секций")
        parser.add_argument("--apply", action="store_true", help="записать в LMS")
        parser.add_argument("--report", help="куда положить xlsx-отчёт; по умолчанию рядом с файлом")

    def handle(self, *args, **options):
        path = Path(options["file"])
        if not path.is_file():
            raise CommandError(f"Нет файла {path}")
        rows = read_rows(path)
        sections = {s.title: s for s in Section.objects.prefetch_related("slots")}
        match(rows)
        overflow = resolve(rows, sections)
        report = Path(options["report"]) if options.get("report") else path.with_name(f"{path.stem}_отчёт.xlsx")
        counts = Counter(row.state for row in rows)
        created = 0
        if options["apply"]:
            with transaction.atomic():
                for row in rows:
                    if row.state != FOUND:
                        continue
                    _enrollment, made = services.import_enrollment(
                        student=row.student,
                        section=sections[row.sheet],
                        slot=row.slot,
                        enrolled_at=row.enrolled_at or timezone.now(),
                    )
                    created += int(made)
        write_report(report, rows, overflow)
        self.stdout.write(f"Строк в листах секций: {len(rows)}")
        for state in (FOUND, MISSING, AMBIGUOUS, TWICE, ALREADY, DONE):
            self.stdout.write(f"  {state}: {counts.get(state, 0)}")
        no_hour = sum(1 for r in rows if r.state == FOUND and r.slot is None)
        self.stdout.write(f"  из них без часа (выберут сами): {no_hour}")
        for line in overflow:
            self.stdout.write(f"  переполнение — {line}")
        if options["apply"]:
            self.stdout.write(self.style.SUCCESS(f"Записано в LMS: {created}"))
        else:
            self.stdout.write("Проверка без записи. Записать — тот же запуск с --apply")
        self.stdout.write(f"Отчёт: {report}")
